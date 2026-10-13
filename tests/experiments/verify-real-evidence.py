"""Accept retained real-model outputs; whitespace in year units is immaterial."""
import hashlib
import json
import re
import sys
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[2]
evidence = root / 'docs/evidence'
prefix = sys.argv[1] if len(sys.argv) > 1 else 'preparation-fixed-real'
assert '/' not in prefix and '..' not in prefix
report = json.loads((evidence / (prefix + '.json')).read_text())
first = next(r for r in report['results'] if r['name'] == 'real-model-first-turn')
final = next(r for r in report['results'] if r['name'] == 'real-model-final')
assert first['execution']['state'] == 'needs_input'
assert first['preparation']['guide'] in ('ready', 'missing')
assert first['preparation']['promises'] == 'missing'
assert final['state'] == 'completed'
assert final['preparation'] == {'guide': 'ready', 'promises': 'ready'}
assert final['selectedChapter'] == 'ready'
assert final['otherUnchanged'] is True
question_id = first['execution']['questions'][0]['callId']
pending = [json.loads(t['result']) for t in final['tools'] if t['name'] == 'ask_user_question']
assert any(p['callId'] == question_id for p in pending)
validations = [t['result'] for t in final['tools'] if t['name'] == 'validate_body']
assert any(v.startswith('[校验通过] body/写作指引.md') for v in validations)
assert any(v.startswith('[校验通过] body/工程验证章节.docx') for v in validations)

outputs = evidence / (prefix + '-files')
promises = (outputs / '关键事实与承诺.md').read_text()
assert re.search(r'\|[^|\n]*维护[^|\n]*\|\s*(?:3|三)\s*年\s*\|', promises)
document = Document(outputs / '工程验证章节.docx')
text = '\n'.join(p.text for p in document.paragraphs)
text += '\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
assert re.search(r'(?:3|三)\s*年', text)
retained = outputs / '保留章节.docx'
assert '人工保留内容，不允许改单章以外文件。' in '\n'.join(p.text for p in Document(retained).paragraphs)

result = {
    'passed': True,
    'model': report['model'],
    'requests': report['realCount'],
    'elapsedMs': final['elapsedMs'],
    'sessionId': first['execution']['session_id'],
    'checks': {
        'guideGeneratedAndValidated': True,
        'questionAskedBeforePromisesAndBody': True,
        'sameSessionContinuation': True,
        'bothPreparationFilesReady': True,
        'selectedDocxGeneratedAndValidated': True,
        'userThreeYearValueInPromisesAndDocx': True,
        'nonselectedDocxBytesUnchanged': True,
    },
    'otherDocxSha256': hashlib.sha256(retained.read_bytes()).hexdigest(),
    'originalHarnessIssue': None if any(r['name']=='real-model-assertions' and r.get('pass') for r in report['results']) else 'Original assertion rejected 3 年 solely due to whitespace; retained actual outputs were reverified without rewriting them.',
    'limits': 'One synthetic two-chapter sample; no full tender/content-quality or adversarial protection validation.',
}
(evidence / (prefix + '-acceptance.json')).write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False))
