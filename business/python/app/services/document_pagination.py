"""原生Word域与分节；分页计算交给打开文档的Office/WPS。"""
from copy import deepcopy
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt
from .document_format import fonts


def field(p, code, value):
    for kind, text in [('begin',None),('instruction',code),('separate',None),('value',str(value)),('end',None)]:
        r=p.add_run()
        if kind=='instruction':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' '+text+' '
        elif kind=='value':
            e=OxmlElement('w:t');e.text=text
        else:
            e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),kind)
            if kind=='begin':e.set(qn('w:dirty'),'true')
        r._r.append(e)
    return e


def bookmark(p, name, number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));p.append(start);p.append(end)


def insert_toc(doc, anchor, entries):
    # 自定义Word版式可能只带潜在样式，先补齐原生目录样式定义。
    for name in ('TOC Heading', 'toc 1', 'toc 2', 'toc 3'):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            style.base_style = doc.styles['Normal']
    heading=doc.add_paragraph('目录',style='TOC Heading')
    heading.paragraph_format.keep_with_next=True
    heading.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # TOC Heading 在部分模板继承 Heading 1，必须显式退出章节列表。
    num=heading._p.get_or_add_pPr().get_or_add_numPr()
    number=OxmlElement('w:numId');number.set(qn('w:val'),'0');num.append(number)
    ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');heading._p.get_or_add_pPr().append(ol)
    elements=[heading._p]
    rows=[]
    headings=[p for p in doc.paragraphs if p.style.name.startswith('Heading ') and
              p._p.find('./'+qn('w:pPr')+'/'+qn('w:outlineLvl')) is None]
    start=max([int(x.get(qn('w:id'))) for x in doc.element.body.iter(qn('w:bookmarkStart')) if (x.get(qn('w:id')) or '').isdigit()]+[0])+1
    used=set()
    for prefix,title,depth in entries:
        if depth>3:continue
        text=(prefix+title).strip()
        hit=next((p for p in headings if id(p._p) not in used and p.text.strip()==title.strip()),None)
        p=doc.add_paragraph(style=f'toc {depth}')
        p.paragraph_format.left_indent=Cm(.74*(depth-1))
        p.paragraph_format.right_indent=Pt(0)
        p.paragraph_format.keep_with_next=False
        width=doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
        p.paragraph_format.tab_stops.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        p.paragraph_format.line_spacing=1.5
        if not rows:
            for tag,kind,value in [('fldChar','begin',None),('instrText',None,' TOC \\o "1-3" \\h \\z \\u '),('fldChar','separate',None)]:
                e=OxmlElement('w:'+tag)
                if kind:
                    e.set(qn('w:fldCharType'),kind)
                    if kind=='begin':e.set(qn('w:dirty'),'true')
                if value:e.text=value;e.set(qn('xml:space'),'preserve')
                p.add_run()._r.append(e)
        p.add_run(text)
        if hit is not None:
            used.add(id(hit._p));name=f'_TenderHeading{start}';bookmark(hit._p,name,start);start+=1
            p.add_run('\t');field(p,'PAGEREF '+name+' \\h','')
            value=p._p.findall('.//'+qn('w:t'))[-1]
            rows.append({'title':text,'depth':depth,'name':name,'value':value})
        else:
            rows.append({'title':text,'depth':depth,'name':None,'value':None})
        elements.append(p._p)
    if rows:
        end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)
    if anchor is not None:
        for e in reversed(elements):anchor.addnext(e)
    else:
        for e in reversed(elements):doc.element.body.insert(0,e)
    return heading._p,rows


def configure_sections(doc,toc_start,has_cover,selected):
    if has_cover:
        previous=toc_start.getprevious()
        if previous is None:
            previous=OxmlElement('w:p');toc_start.addprevious(previous)
        if previous.tag!=qn('w:p'):
            p=OxmlElement('w:p');previous.addnext(p);previous=p
        section=deepcopy(doc.sections[-1]._sectPr)
        for tag in ('footerReference','headerReference','titlePg','pgNumType'):
            for e in list(section.findall(qn('w:'+tag))):section.remove(e)
        typ=section.find(qn('w:type'))
        if typ is None:typ=OxmlElement('w:type');section.append(typ)
        typ.set(qn('w:val'),'nextPage')
        previous.get_or_add_pPr().append(section)
    # 总页数引用正文最后的书签，跨多节连续编号也能排除任意页数的封面。
    paragraphs = list(doc.element.body.iter(qn('w:p')))
    if not paragraphs:
        raise ValueError('合册没有可编号的正文')
    existing = {b.get(qn('w:name')) for b in doc.element.body.iter(qn('w:bookmarkStart'))}
    number = max([int(b.get(qn('w:id'))) for b in doc.element.body.iter(qn('w:bookmarkStart')) if (b.get(qn('w:id')) or '').isdigit()]+[0])+1
    end_name = '_TenderCountedEnd'+str(number)
    while end_name in existing:
        number+=1;end_name='_TenderCountedEnd'+str(number)
    bookmark(paragraphs[-1],end_name,number)
    counted_start = 1 if has_cover else 0
    for i,sec in enumerate(doc.sections):
        sec.different_first_page_header_footer=False
        for footer in (sec.footer,sec.even_page_footer,sec.first_page_footer):
            footer.is_linked_to_previous=False
            for e in list(footer._element):footer._element.remove(e)
            p=footer.add_paragraph()
            if has_cover and i==0:continue
            spec={'font':'宋体','font_size':10.5,'alignment':'center',**selected.get('footer',{})}
            p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT}[spec['alignment']]
            p.paragraph_format.line_spacing=Pt(spec['line_spacing']) if spec.get('spacing_rule')=='exact' else spec.get('line_spacing',1.0)
            text=spec.get('footer_text','第 {PAGE}页，共{TOTAL}页')
            for part in re.split(r'(\{PAGE\}|\{TOTAL\})',text):
                if part in ('{PAGE}','{TOTAL}'):field(p,'PAGE' if part=='{PAGE}' else 'PAGEREF '+end_name+' \\* Arabic','')
                else:p.add_run(part)
            for run in p.runs:fonts(run._r.get_or_add_rPr(),spec['font'],spec['font_size'])
        if has_cover and i==0:continue
        num=sec._sectPr.find(qn('w:pgNumType'))
        if num is None:num=OxmlElement('w:pgNumType');sec._sectPr.append(num)
        num.attrib.pop(qn('w:start'),None)
        if i==counted_start:num.set(qn('w:start'),'1')
        num.set(qn('w:fmt'),'decimal')
    settings=doc.settings.element
    update=settings.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');settings.append(update)
    update.set(qn('w:val'),'true')


def finalize(doc,path,rows,has_cover,selected):
    """校验可更新的Word结构，不伪造分页缓存，也不运行外部排版器。"""
    from .export_validation import verify
    from . import page_references
    references, unresolved = page_references.prepare(doc,rows)
    if unresolved:
        details = '；'.join('表'+str(r['table']+1)+'第'+str(r['row']+1)+'行：'+r['reason'] for r in unresolved)
        raise ValueError(f'有{len(unresolved)}处资料页码未能建立自动引用，未发布整册；须补齐实际附件及对应关系：'+details)
    doc.save(path)
    verification = verify(path,rows,has_cover,selected)
    return {'field_update':'on_open','pagination_calculated':False,
            'headings':[{'title':r['title'],'bookmark':r['name']} for r in rows],
            'page_references':[{k:v for k,v in ref.items() if k not in ('first','last')} for ref in references],
            'unresolved_page_references':[], 'export_verification':verification}
