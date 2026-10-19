"""将有明确章节目标的资料清单页码换成原生 PAGEREF，保留逐格变更记录。"""
import re
from docx.oxml.ns import qn


def _key(text):
    text = re.sub(r'^\d+(?:\.\d+)*\s*', '', text.strip())
    text = re.sub(r'[（(](?:格式|附件)[^）)]*[）)]', '', text)
    return re.sub(r'\s+', '', text)


def prepare(doc, headings):
    from .document_pagination import bookmark, field
    body = list(doc.element.body)
    targets = {}
    marker = max([int(x.get(qn('w:id'))) for x in doc.element.body.iter(qn('w:bookmarkStart')) if (x.get(qn('w:id')) or '').isdigit()]+[0])+1
    for i, heading in enumerate(headings):
        if not heading['name']: continue
        para = next((p for p in doc.paragraphs if any(b.get(qn('w:name')) == heading['name'] for b in p._p.iter(qn('w:bookmarkStart')))), None)
        if para is None: continue
        following = next((h for h in headings[i+1:] if h['name'] and h['depth'] <= heading['depth']), None)
        end_index = len(body)-1
        if following:
            end = next(p for p in doc.paragraphs if any(b.get(qn('w:name')) == following['name'] for b in p._p.iter(qn('w:bookmarkStart'))))
            end_index = body.index(end._p)-1
        last = para._p
        for block in body[body.index(para._p):end_index+1]:
            for p in ([block] if block.tag == qn('w:p') else block.iter(qn('w:p'))):
                if next(p.iter(qn('w:t'),qn('w:drawing'),qn('w:pict'),qn('w:object')),None) is not None:
                    last = p
        start_name, end_name = f'TenderRefStart{marker}', f'TenderRefEnd{marker}'
        bookmark(para._p,start_name,marker);marker+=1
        bookmark(last,end_name,marker);marker+=1
        targets.setdefault(_key(heading['title']),[]).append((start_name,end_name))
    references, unresolved = [], []
    for table_index, table in enumerate(doc.tables):
        if not table.rows: continue
        headers = [re.sub(r'\s+','',c.text) for c in table.rows[0].cells]
        if '材料所在页码' in headers:
            column = headers.index('材料所在页码')
            for row_index,row in enumerate(table.rows[1:],1):
                if column < len(row.cells) and re.fullmatch(r'\s*\d+(?:\s*[-—至]\s*\d+)?\s*',row.cells[column].text):
                    unresolved.append({'table':table_index,'row':row_index,'column':column,'original':row.cells[column].text,
                                       'reason':'人员材料没有可唯一定位的附件书签，保留原值，需补齐实际证明材料及对应关系'})
        if '页码范围' not in headers or '资料名称' not in headers: continue
        page_column, title_column = headers.index('页码范围'), headers.index('资料名称')
        for row_index, row in enumerate(table.rows[1:],1):
            if len(row.cells) <= max(page_column,title_column): continue
            title = row.cells[title_column].text
            exact = targets.get(_key(title),[])
            parts = [_key(title)] if len(exact) == 1 else re.split(r'[、，,]',_key(title))
            matches = [targets.get(part,[]) for part in parts]
            if not matches or any(len(m) != 1 for m in matches):
                unresolved.append({'table':table_index,'row':row_index,'title':title,'reason':'没有唯一对应的目录章节，保留原值'})
                continue
            cell = row.cells[page_column]; original = cell.text
            if not re.fullmatch(r'\s*\d+(?:\s*[-—至]\s*\d+)?\s*',original) or next(cell._tc.iter(qn('w:fldChar'),qn('w:fldSimple'),qn('w:commentReference')),None) is not None:
                unresolved.append({'table':table_index,'row':row_index,'title':title,'reason':'页码格含说明或复杂内容，保留原文'})
                continue
            anchors = [e for e in cell._tc.iter() if e.tag in {qn('w:bookmarkStart'),qn('w:bookmarkEnd'),qn('w:commentRangeStart'),qn('w:commentRangeEnd')}]
            p = cell.paragraphs[0]
            for other in list(cell._tc):
                if other.tag == qn('w:p') and other is not p._p: cell._tc.remove(other)
            for child in list(p._p):
                if child.tag != qn('w:pPr'):p._p.remove(child)
            start,end = matches[0][0][0],matches[-1][0][1]
            field(p,'PAGEREF '+start+' \\h','')
            first = p._p.findall('.//'+qn('w:t'))[-1]
            p.add_run('-');field(p,'PAGEREF '+end+' \\h','')
            last = p._p.findall('.//'+qn('w:t'))[-1]
            for anchor in anchors:p._p.append(anchor)
            references.append({'table':table_index,'row':row_index,'column':page_column,'title':title,'original':original,
                               'start':start,'end':end,'first':first,'last':last})
    return references, unresolved
