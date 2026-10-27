"""Read Excel values without changing files, calculating formulas or executing macros."""
from datetime import date, datetime, time
from pathlib import Path
import re

from . import ParseResult, register

MAX_CELLS = 1_000_000
TEXT_VIEW = 'excel-values-2'


def column_name(index):
    result = ''
    while index:
        index, digit = divmod(index - 1, 26)
        result = chr(65 + digit) + result
    return result


def cell_text(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (datetime, date, time)):
        if isinstance(value, datetime) and value.time() == time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, (int, float)):
        fmt = number_format.split(';')[0]
        if re.fullmatch(r'0{2,}', fmt) and float(value).is_integer():
            return str(int(value)).zfill(len(fmt))
        percent = re.fullmatch(r'0(?:\.(0+))?%', fmt)
        if percent:
            return f'{value * 100:.{len(percent[1] or "")}f}%'
        fixed = re.fullmatch(r'(#,##0|0)(?:\.(0+))?', fmt)
        if fixed:
            return format(value, f'{"," if "," in fmt else ""}.{len(fixed[2] or "")}f')
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def _xlsx(path):
    from openpyxl import load_workbook
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    values = None
    try:
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = []
        cells = 0
        for sheet in formulas:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                raise ValueError(f'工作表“{sheet.title}”范围超过100万单元格，请移除多余空白行列或拆分文件')
            rows, warnings = [], []
            for number, (source_row, value_row) in enumerate(zip(sheet.iter_rows(), values[sheet.title].iter_rows()), 1):
                cells += len(source_row)
                if cells > MAX_CELLS:
                    raise ValueError('工作簿超过100万单元格，请拆分文件后上传')
                texts = []
                for source, cached in zip(source_row, value_row):
                    value = cached.value if source.data_type == 'f' else source.value
                    if source.data_type == 'f' and value is None:
                        value = '【公式结果缺失】'
                        warnings.append({'cell': source.coordinate, 'message': '公式没有已保存结果，请在Excel中计算并保存后替换文件；本系统不重新计算公式。'})
                    elif cached.data_type == 'e' or source.data_type == 'e':
                        warnings.append({'cell': source.coordinate, 'message': f'单元格包含错误值：{value}，请核对原件。'})
                    texts.append(cell_text(value, source.number_format))
                if any(texts):
                    rows.append({'number': number, 'cells': texts})
            sheets.append({'name': sheet.title, 'hidden': sheet.sheet_state != 'visible', 'columns': sheet.max_column or 0, 'rows': rows, 'warnings': warnings})
        return sheets
    finally:
        formulas.close()
        if values is not None:
            values.close()


def _xls(path):
    import xlrd
    from openpyxl.styles.numbers import STRIP_RE, is_datetime, is_timedelta_format
    from openpyxl.utils.datetime import from_excel
    book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    try:
        result, cells = [], 0
        for sheet in book.sheets():
            cells += sheet.nrows * sheet.ncols
            if cells > MAX_CELLS:
                raise ValueError('工作簿超过100万单元格，请拆分文件后上传')
            rows, warnings = [], []
            for number in range(sheet.nrows):
                texts = []
                for column in range(sheet.ncols):
                    cell = sheet.cell(number, column)
                    value = cell.value
                    fmt = book.format_map.get(book.xf_list[cell.xf_index].format_key)
                    number_format = fmt.format_str if fmt else 'General'
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        # Literal text, escaped/spacing characters and colour/locale
                        # tags are presentation, not date/time fields. Strip literals
                        # before splitting so a quoted semicolon stays a literal.
                        value_format = re.sub(r'"[^"]*"|[\\_*].', '', number_format)
                        value_format = STRIP_RE.sub('', value_format).split(';')[0].lower()
                        if is_timedelta_format(value_format):
                            value = from_excel(value, timedelta=True)
                        else:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                            # A bare "mm" is a calendar month, not a clock value.
                            if is_datetime(value_format) == 'time' and re.search('[hs]', value_format):
                                value = value.time()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#ERROR')
                        warnings.append({'cell': f'{column_name(column + 1)}{number + 1}', 'message': f'单元格包含错误值：{value}，请核对原件。'})
                    texts.append(cell_text(value, number_format))
                if any(texts):
                    rows.append({'number': number + 1, 'cells': texts})
            result.append({'name': sheet.name, 'hidden': bool(sheet.visibility), 'columns': sheet.ncols, 'rows': rows, 'warnings': warnings})
        return result
    finally:
        book.release_resources()


def read_workbook(path):
    try:
        sheets = _xlsx(path) if Path(path).suffix.lower() == '.xlsx' else _xls(path)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Excel无法读取，请检查文件是否损坏、加密或扩展名不符；可在Excel中重新保存后替换。') from error
    if not any(sheet['rows'] for sheet in sheets):
        raise ValueError('Excel没有可读取的单元格内容，请检查文件后替换')
    return sheets


def _markdown(value):
    return value.replace('\\', '\\\\').replace('|', '\\|').replace('\r\n', '\n').replace('\r', '\n').replace('\n', '<br>')


@register(['.xls', '.xlsx'])
def convert(path):
    sheets = read_workbook(path)
    lines = ['公式仅使用原件已保存结果，不重新计算；合并单元格的内容位于其左上角。', '']
    for sheet in sheets:
        lines.extend([f'# 工作表：{sheet["name"]}', ''])
        if sheet['hidden']:
            lines.extend(['此工作表在原件中隐藏，内容仍纳入解读。', ''])
        columns = sheet['columns']
        if sheet['rows']:
            lines.append('| 原表行号 | ' + ' | '.join(column_name(i + 1) for i in range(columns)) + ' |')
            lines.append('| --- | ' + ' | '.join('---' for _ in range(columns)) + ' |')
            lines.extend('| ' + str(row['number']) + ' | ' + ' | '.join(_markdown(cell) for cell in row['cells']) + ' |' for row in sheet['rows'])
        else:
            lines.append('（空工作表）')
        lines.extend(['', *[f'待核对 {warning["cell"]}：{warning["message"]}' for warning in sheet['warnings']], ''])
    return ParseResult('\n'.join(lines), {'conversion': 'excel-values', 'text_view': TEXT_VIEW, 'tables': sum(bool(s['rows']) for s in sheets), 'sheets': [s['name'] for s in sheets], 'warnings': [f'{s["name"]}!{w["cell"]}：{w["message"]}' for s in sheets for w in s['warnings']]})
