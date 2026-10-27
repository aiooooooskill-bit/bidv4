"""tender-body 正文产物机器校验工具（提示不是门禁：逐条给修复写法，不阻塞流程）。

校验契约（真值 = tender-body SKILL.md + references/guide-format.md）：
- 写作指引（body/写作指引.md）：需正文的目录叶子每节有行（多册键=「册名/标题」，
  与目录标题逐字一致）、模式列取值合法（素材修订/格式跟随/推理撰写，「+」组合；
  非正文节点可「—」）、素材列块 id 可解析（blk_ 存在于素材库）、同块多派=弱级
  提示（2026-09-15 契约修正：块=拷贝授权范围非注入原子，同一块的不同区间派
  不同节合法——各写手只注入自己需要的区间，同一内容区间不进两节由节间查重
  兜底；不同块同素材文件区间重叠=弱级提示，治历史重复建块）
- 正文节（body/<册>/<节>.docx）：待办批注清点（docx_comment_add 落的
  Word 批注——收尾汇报逐条点名）、内联占位【待补/待澄清】**判不过**（占位文字会进
  交付稿，正规落点是批注）、疑似残留（选用素材块所在文件名主干入扫描表——弱级来源，
  通用产品词可能误报，提示核对而非必须清零）、素材使用率（正文与选用
  块的字符级 shingle 重叠率——「查到素材却凭空写」的机械防线，提示不是门禁；
  block_ids 支持 blk_…:L起-L止 后缀，只注入了块内区间时按区间算）；
  docx 节按**接受全部修订后的终稿视角**取文本（评委最终看到的），批注/占位定位段落
  为 P 段号、表格行为 T{t}R{r}（与 docx_section_read 视图互查）
- 全局（section="body"）：承诺清单（关键事实与承诺.md 的「事项｜值」表）每个值
  是否落进任意正文节——去空白归一化子串匹配，未落正文点名（该节未写/值被改写）；
  节间两两查重（归一化 shingle 交集/较小方 ≥50% 点名——同一素材进多节或原文
  照搬，独立于承诺清单跑）；「整本-」合册产物不参与（派生产物）；反向「正文
  出现与清单不一致的值」属语义审阅，不在此判
"""
from __future__ import annotations
import re
from pathlib import Path
from ..work_paths import normalize_work_path, resolve_work_path
from docx import Document
from docx.oxml.ns import qn
from .. import db, runctx
from ..promises import promise_rows
from ..artifact_store import work_dir
from ..knowledge import materials_lib
from . import body_contract
from .check_residue import re_split_name, scan_residue
from .docx_ops import material_usage, _accepted_text, _heading_style_ids, section_comments_labeled, section_lines_labeled, section_text_lines
from ..markdown import iter_tables, markdown_text
_BLOCK_ID_RE = re.compile('blk_[0-9a-f]{12}')
_PLACEHOLDER_MARKS = body_contract.PLACEHOLDER_MARKS
_OVERLAP_HINT = 0.1
_DUP_RATIO = 0.5
_SHINGLE = 10
_TOOL_NAME_RE = re.compile('\\b(?:docx_\\w+|check_name_residue|parse_document|search_\\w+|read_artifact|publish_artifact|assemble_tender|list_templates)\\b')
_NOTE_ANAPHORA_RE = re.compile('【缺[：:]\\s*(?:上述|以上|前述|前面|前面所|以上所)')
_MD_HEADING_RE = re.compile('^\\s{0,3}#{1,6}\\s')
_MD_TABLE_ROW_RE = re.compile('^\\|.*\\|\\s*$')
_FIGURE_KINDS = ('表', '分层', '辐射', '甘特', '原型', '流程')
_FIGURE_ITEM_RE = re.compile('^(表|分层|辐射|甘特|原型|流程)\\s*[:：]\\s*(\\S.*)$')

def _normalize(text: str) -> str:
    """以成熟解析器取 Markdown 语义文本，再消除空白及 DOCX 拼格分隔符。

    素材的标题、表格、强调、链接与纯文本终稿按同一内容口径比对。
    """
    # 同时支持素材 Markdown 和 DOCX 纯文本；表格管线是 DOCX 拼格格式。
    return re.sub(r'\s+', '', markdown_text(text).replace('|', ''))


def _shingles(text: str) -> set[str]:
    return {text[i:i + _SHINGLE] for i in range(max(0, len(text) - _SHINGLE + 1))}

def _overlap_ratio(section_text: str, block_text: str) -> float | None:
    """块 shingle 在正文中的出现比例（None=块太短，比率无意义）。"""
    bs = _shingles(_normalize(block_text))
    if len(bs) < 5:
        return None
    return len(bs & _shingles(_normalize(section_text))) / len(bs)

def _parse_id_ranges(tail: str) -> list[list[int]] | None:
    """block_ids 后缀区间「L133-L318,L400-L450」→ [[s,e],…]；空/解析失败返回
    None（回落整块——校验侧宽松，授权围栏在注入工具侧硬拦）。"""
    out: list[list[int]] = []
    for seg in re.split('[，,、]', (tail or '').strip()):
        seg = seg.strip()
        if not seg:
            continue
        a, _, b = seg.partition('-')
        try:
            lo, hi = (int(a.strip().lstrip('Ll')), int(b.strip().lstrip('Ll')))
        except ValueError:
            return None
        if lo > hi:
            lo, hi = (hi, lo)
        out.append([lo, hi])
    return out or None

def _block_text(bid: str) -> tuple[dict | None, str]:
    """按块 id 取块行与切片全文（块失效返回 (None, "")）。

    bid 支持「blk_xxx:L133-L318」后缀写法（2026-09-15 注入粒度自选批）：写手
    只注入了块内区间时，使用率按该子区间算——整块算会把子区间注入误报
    「未实质使用」；后缀解析失败回落整块。"""
    base, sep, tail = bid.partition(':')
    sub = _parse_id_ranges(tail) if sep else None
    b = db.mt_get_block(base)
    if not b:
        return (None, '')
    f = db.mt_get_file(b['file_id'])
    if not f:
        return (b, '')
    md_path, _, _ = materials_lib.mt_parse_paths(f['file_name'])
    return (b, materials_lib._slice(md_path, sub or (b.get('ranges') or [])))

def _residue_names(block_ids: list[str]) -> list[str]:
    """选用素材块所在文件名的主干词（≥4 字）——旧项目名最常见的藏身处。
    容 blk 后缀写法（取「:」前的裸 id）。"""
    names: list[str] = []
    seen: set[str] = set()
    for bid in block_ids:
        b = db.mt_get_block(bid.partition(':')[0])
        if not b:
            continue
        f = db.mt_get_file(b['file_id'])
        if not f:
            continue
        for part in re_split_name(Path(f['file_name']).stem):
            if len(part) >= 4 and part not in seen:
                seen.add(part)
                names.append(part)
    return names

def _parse_guide_rows(content: str) -> list[tuple[int, list[str], dict[str, int]]]:
    """指引表的 (行号, cells, 列名→下标)；无合法表返回空。"""
    rows: list[tuple[int, list[str], dict[str, int]]] = []
    for _, header, body in iter_tables(content):
        idx = {c: i for i, c in enumerate(header)}
        if '节' not in idx or '模式' not in idx:
            continue
        col = idx['节']
        for lineno, cells in body:
            if col < len(cells) and cells[col].strip():
                rows.append((lineno, cells, idx))
    return rows

def _ranges_overlap(a: list, b: list) -> bool:
    """两组行号区间是否有交集（与前端建块重叠提示 anyOverlap 同语义）。"""
    return any((s1 <= e2 and s2 <= e1 for s1, e1 in a for s2, e2 in b))

def _validate_guide(content_text: str, task_id: str) -> tuple[list[str], list[str], list[str]]:
    issues: list[str] = []
    warnings: list[str] = []
    _, content, _ = body_contract.load_directory(task_id)
    if content is None:
        warnings.append('无目录产物或内容不可读——叶子对账跳过，仅校验模式取值与素材块 id')
    multi = body_contract.multi_volume(content) if content else False
    leaf_keys: set[str] = set()
    prose_keys: set[str] = set()
    delivery: dict[str, str] = {}
    if content:
        for vol, title, mode in body_contract.body_nodes_for_task(task_id, content):
            key = f'{vol}/{title}' if multi else title
            leaf_keys.add(key)
            delivery[key] = mode
            if mode not in body_contract.NON_PROSE_DELIVERY:
                prose_keys.add(key)
    rows = _parse_guide_rows(content_text)
    if not rows:
        issues.append('指引缺少表格——须为「| 节 | 模式 | 依据 | 素材 | 缺口/备注 |」的 Markdown 表')
        return (issues, warnings, [])
    row_keys: set[str] = set()
    row_blocks: dict[str, list[str]] = {}
    block_cache: dict[str, dict | None] = {}
    for lineno, cells, idx in rows:
        key = cells[idx['节']].strip()
        row_keys.add(key)
        mode = cells[idx['模式']].strip() if idx.get('模式', 1) < len(cells) else ''
        material = cells[idx['素材']].strip() if idx.get('素材', 3) < len(cells) else ''
        if content and key not in leaf_keys:
            issues.append(f'写作指引.md:{lineno} 行「{key}」不在目录正文目标中——节名须与目录逐字一致（含有自有正文的父节点，多册写「册名/标题」），目录改版后须重对指引')
        tokens = [t.strip() for t in mode.split('+') if t.strip()]
        if tokens == ['—']:
            if key in prose_keys:
                warnings.append(f'写作指引.md:{lineno} 「{key}」被目录标为需正文但模式写了「—」——若实为表格/索引/填表类节点，指引行注明「按招标格式填表」即可；确需正文则补写模式')
        elif not tokens:
            issues.append(f'写作指引.md:{lineno} 模式列为空——素材修订/格式跟随/推理撰写（可「+」组合）')
        else:
            bad = [t for t in tokens if t not in body_contract.MODES]
            if bad:
                issues.append(f"写作指引.md:{lineno} 模式「{'、'.join(bad)}」不合法——只能取 {'/'.join(body_contract.MODES)}（「+」组合；非正文节点写「—」）")
            if '素材修订' in tokens and (not _BLOCK_ID_RE.findall(material)) and ('【缺】' not in material):
                warnings.append(f'写作指引.md:{lineno} 「{key}」素材修订模式但素材列无块 id 也未标【缺】——先 search_references 检索，确实没有再标【缺】（缺料清单要向用户点名）')
        for bid in dict.fromkeys(_BLOCK_ID_RE.findall(material)):
            if bid not in block_cache:
                block_cache[bid] = db.mt_get_block(bid)
            b = block_cache[bid]
            if b is None:
                issues.append(f'写作指引.md:{lineno} 素材块 {bid} 失效（块已删除）——重检索并更新指引')
                continue
            keys = row_blocks.setdefault(bid, [])
            if key not in keys:
                keys.append(key)
        dep = cells[idx['依据']].strip() if idx.get('依据', 2) < len(cells) else ''
        if not dep and content and (key in prose_keys):
            warnings.append(f'写作指引.md:{lineno} 「{key}」依据列为空——纯过渡章允许；应写节须挂 REQ/SCORE/MAND ID')
        fig = cells[idx['图示']].strip() if idx.get('图示', 5) < len(cells) else ''
        if fig and fig != '—':
            for item in (x.strip() for x in re.split('[、;；,，]', fig)):
                if item and (not _FIGURE_ITEM_RE.match(item)):
                    warnings.append(f"写作指引.md:{lineno} 「{key}」图示项「{item[:30]}」记法须为「类型:主题」（类型∈{'/'.join(_FIGURE_KINDS)}；无图示写「—」）")
        note = cells[idx['缺口/备注']].strip() if idx.get('缺口/备注', 4) < len(cells) else ''
        if note and note != '—':
            tool_hit = _TOOL_NAME_RE.search(note)
            if tool_hit:
                warnings.append(f'写作指引.md:{lineno} 「{key}」缺口列混入工具名 {tool_hit.group(0)}——执行链路由模式列与派发说明自动补（依据列含 TPL 编号时程序自动加「本节含格式件」与「原件定位」行），该列是给用户看的缺料清单，删工具名；行号仅在程序定位不到时以「原件：<格式名> L起-L止」保留')
            if _NOTE_ANAPHORA_RE.search(note):
                warnings.append(f'写作指引.md:{lineno} 「{key}」缺口项用了代词回指（上述/以上/前面）——该列会被摘成「需要你提供」独立展示，代词摘出来后指代丢失，请列具体字段名（如「公司全称、注册地址、成立时间、法定代表人姓名」）')
    for bid, keys in row_blocks.items():
        if len(keys) > 1:
            title = (block_cache.get(bid) or {}).get('title') or bid
            warnings.append(f"素材块 {bid}《{title}》同时派给 {len(keys)} 个节（{'、'.join(keys)}）——同一块可多节共用，但各写手只注入自己需要的行号区间（docx_material_inject 传 lines），同一内容区间不得进两节（节间查重兜底）；派发说明/备注列写明各节取哪些章节")
    by_file: dict[str, list[str]] = {}
    for bid, keys in row_blocks.items():
        b = block_cache.get(bid) or {}
        if len(keys) == 1 and b.get('file_id'):
            by_file.setdefault(b['file_id'], []).append(bid)
    for fid, bids in by_file.items():
        if len(bids) < 2:
            continue
        f = db.mt_get_file(fid)
        fname = (f or {}).get('file_name') or fid
        for i in range(len(bids)):
            for j in range(i + 1, len(bids)):
                ba, bb = (block_cache[bids[i]], block_cache[bids[j]])
                if not _ranges_overlap(ba.get('ranges') or [], bb.get('ranges') or []):
                    continue
                warnings.append(f"素材块《{ba.get('title') or bids[i]}》与《{bb.get('title') or bids[j]}》同属素材文件 {fname} 且勾选区间重叠，分别派给「{row_blocks[bids[i]][0]}」「{row_blocks[bids[j]][0]}」——内容大面积相同，两节正文可能重复：确认是否重复建块（保留其一），或改拆成不重叠区间")
    if content:
        required_keys = {key for key in leaf_keys if delivery[key] != '目录容器'}
        for key in sorted(required_keys - row_keys):
            issues.append(f"指引缺行：「{key}」（交付形态={delivery.get(key) or '未标注'}）——需正文的每节都要一行")
    return (issues, warnings, [])

def _validate_section(lines: list[str], block_ids: list[str] | None, labels: list[str] | None=None) -> tuple[list[str], list[str], list[str]]:
    issues: list[str] = []
    warnings: list[str] = []
    notes: list[str] = []
    if not any((ln.strip() for ln in lines)):
        hint = '缺料用 docx_comment_add 加批注记待办'
        return ([f'正文为空——每节须有实质内容（{hint}），不得留空文件'], warnings, notes)
    for i, ln in enumerate(lines, 1):
        if any((mark in ln for mark in _PLACEHOLDER_MARKS)):
            loc = labels[i - 1] if labels and i <= len(labels) else f'P{i}'
            snippet = ln.strip()[:30]
            issues.append(f'{loc}：正文含内联占位「{snippet}」——占位文字会随交付稿出去，删除该文字、改用 docx_comment_add 加批注记待办（缺料/待澄清/待核验都落批注，不落正文）')
    for i, ln in enumerate(lines, 1):
        loc = labels[i - 1] if labels and i <= len(labels) else f'P{i}'
        if _MD_HEADING_RE.match(ln):
            issues.append(f'{loc}：行首 markdown 标题残字「{ln.strip()[:24]}」——docx 通道不解析 #，层级用 Heading 样式（建节后 docx_section_revise 的 insert_after 带 style）')
        elif '**' in ln:
            issues.append(f'{loc}：正文含「**」加粗残字——docx 通道不解析 markdown，星号会原样印进交付稿；删除星号用普通文本')
        elif not loc.startswith('T') and _MD_TABLE_ROW_RE.match(ln.strip()):
            issues.append(f'{loc}：整行为竖线拼的 markdown 表格残字——表格用建节 body 的 table 块或 docx_diagram_insert 生成，不在正文写竖线表')
        elif '`' in ln:
            warnings.append(f'{loc}：正文含反引号——确认不是 markdown 代码残字（代码类内容用普通文字描述）')
    if block_ids is None:
        notes.append('未传 block_ids——旧名残留与素材使用率检查跳过（使用计划里选了哪些块就传哪些）')
        return (issues, warnings, notes)
    for hit in scan_residue(lines, _residue_names(block_ids), labels=labels):
        warnings.append(f'疑似残留（来自素材文件名，可能为通用术语）：{hit}——逐条核对，旧项目用语→替换')
    for bid in dict.fromkeys(block_ids):
        b, text = _block_text(bid)
        if b is None:
            issues.append(f'素材块 {bid} 失效（块已删除）——重检索换块后重写本节或更新使用计划')
            continue
        if not text:
            warnings.append(f"素材块「{b.get('title') or bid}」取不到内容（解析产物缺失）——跳过重叠率")
            continue
        ratio = _overlap_ratio('\n'.join(lines), text)
        if ratio is None:
            continue
        if ratio < _OVERLAP_HINT:
            warnings.append(f"素材块「{b.get('title') or bid}」重叠率 {ratio:.0%}——未实质使用素材（素材修订=先把选用的区间贴进底稿再改写适配，不是看着参考另写一篇；只注入了块内区间时 block_ids 写 blk_…:L起-L止 按区间算）")
    return (issues, warnings, notes)

def _figure_plan_reconcile(doc, p: Path, body_root: Path) -> list[str]:
    """图示对账（提示级，2026-09-14 批）：指引「图示」列计划 vs 节内实收表格数。
    「不进计划的能力=不存在」的对账面——计划了没产出=能力闲置、计划外产出=
    注水嫌疑；计划外只对推理撰写节提示（素材修订/格式跟随的表格来自注入与拷
    原件，不背计划的账）；节名与指引行对不上（文件名清洗差异）静默跳过。"""
    guide = body_root / body_contract.GUIDE_NAME
    if not guide.is_file():
        return []
    try:
        rel = p.relative_to(body_root)
    except ValueError:
        return []
    key = '/'.join((*rel.parts[:-1], p.stem))
    rows = _parse_guide_rows(guide.read_text(encoding='utf-8', errors='replace'))
    row = next((r for r in rows if r[1][r[2]['节']].strip() == key), None)
    if row is None:
        return []
    _, cells, idx = row
    fig = cells[idx['图示']].strip() if idx.get('图示', 5) < len(cells) else ''
    mode = cells[idx['模式']].strip() if idx.get('模式', 1) < len(cells) else ''
    plan = [x.strip() for x in re.split('[、;；,，]', fig) if x.strip() and x.strip() != '—']
    actual = len(doc.tables) + len(doc.element.body.findall('.//' + qn('a:blip')))
    out: list[str] = []
    if plan and actual < len(plan):
        out.append(f'〔图示〕指引计划 {len(plan)} 项、节内表格/图片 {actual} 项——未产出项收尾前补齐或改指引：' + '；'.join(plan))
    if not plan and actual and ('推理撰写' in mode):
        out.append(f'〔图示〕指引未计划图示、节内有 {actual} 项表格/图片——推理撰写节的补进指引图示列（计划先行），或在收尾汇报说明理由')
    return out

def _section_dup_warnings(body_dir: Path, sections: list[Path], finals: list[str]) -> list[str]:
    """节间两两查重：同一素材块进多节、或一节原文照搬另一节（2026-09-08 实测：
    同一 625 段素材块全文进了「总体建设方案」与「系统集成方案」两节）。
    重叠率 = 归一化 shingle 交集 / 较小方；短节（<5 shingle，承诺行级别）跳过；
    提示不是门禁——收尾前逐对裁决（拆块或改写）。"""
    sigs: list[tuple[str, set[str]]] = []
    for q, text in zip(sections, finals, strict=True):
        sh = _shingles(_normalize(text))
        if len(sh) < 5:
            continue
        sigs.append((q.relative_to(body_dir).with_suffix('').as_posix(), sh))
    warns: list[str] = []
    for i in range(len(sigs)):
        for j in range(i + 1, len(sigs)):
            (na, sa), (nb, sb) = (sigs[i], sigs[j])
            ratio = len(sa & sb) / min(len(sa), len(sb))
            if ratio >= _DUP_RATIO:
                warns.append(f'节间疑似重复：「{na}」与「{nb}」内容重叠约 {ratio:.0%}——同一素材进多节或原文照搬；对照写作指引素材列，拆块或改写')
    return warns

def body_global_checks(body_dir: Path, files: list[Path] | None = None, check_promises: bool = True) -> dict:
    """全局收尾比对：承诺清单的每个值是否落进了正文（单向：清单→正文）+ 节间查重。

    匹配 = 去空白归一化后的子串（「90 天」↔「90天」视为一致）；未在任何正文节
    出现的值点名提示（该节未写或值被改写，核对后补落正文或修正清单）。节间查重
    独立于承诺清单（清单缺失/格式不对也照跑——两节大面积重叠才是它要抓的事）。
    全部为提示级——提示不是门禁；反向「正文出现与清单不一致的值」属语义审阅，
    不在此判。
    """
    sections = files if files is not None else [q for q in sorted(body_dir.rglob('*.docx')) if not (q.parent == body_dir and q.name.startswith(body_contract.VOLUME_PREFIX))]
    if not sections:
        return {'summary': '[校验通过] body 全局：暂无正文节——跳过承诺比对', 'warnings': [], 'notes': []}

    def _final_text(q: Path) -> str:
        from ..parse.docx_text import load_text_document
        return '\n'.join(section_text_lines(load_text_document(q)[0]))
    finals: list[str] = []
    broken: list[str] = []
    readable: list[Path] = []
    for q in sections:
        try:
            finals.append(_final_text(q))
            readable.append(q)
        except Exception:
            broken.append(f'{q.parent.name}/{q.name}')
    if broken and (not finals):
        return {'summary': '[校验失败] body 全局：正文节全部损坏', 'warnings': ['正文节全部损坏（' + '、'.join(broken) + '）——删除重建或人工核查'], 'notes': []}
    sections = readable
    dup_warns = _section_dup_warnings(body_dir, sections, finals)
    skip_note: str | None = None
    promise_summary = ''
    warn_lines: list[str] = []
    promise_path = body_dir / body_contract.PROMISE_NAME
    if not check_promises:
        skip_note = '承诺清单未就绪或已过期，未进行当前承诺比对'
    elif not promise_path.is_file():
        skip_note = '无关键事实与承诺清单——跳过承诺比对（开工时应生成 body/关键事实与承诺.md）'
    else:
        parsed, header_seen = promise_rows(promise_path.read_text(encoding='utf-8', errors='replace'))
        rows = [(r.line, r.item, r.value) for r in parsed if r.item and r.value]
        if not header_seen:
            skip_note = '承诺清单不是「事项｜值｜说明」表——跳过比对（格式契约见技能 references/guide-format.md）'
        else:
            texts = [re.sub('\\s+', '', t) for t in finals]
            placed = 0
            for lineno, item, val in rows:
                key = re.sub('\\s+', '', val)
                if len(key) < 2:
                    continue
                if any((key in t for t in texts)):
                    placed += 1
                else:
                    warn_lines.append(f'关键事实与承诺.md:{lineno} 「{item}={val}」未在任何正文出现——该节未写或值被改写，核对后补落正文或修正清单')
            promise_summary = f'，承诺 {len(rows)} 项比对，{placed} 项已落正文' + (f'，{len(warn_lines)} 项未落正文' if warn_lines else '')
    warn_lines.extend(dup_warns)
    if broken:
        warn_lines.append('损坏节（无法按终稿视角读取，已跳过、未参与比对/查重）：' + '、'.join(broken) + '——删除重建或人工核查')
    head = f'[校验通过] body 全局：{len(sections)} 节正文{promise_summary}' if skip_note is None else f'[校验通过] body 全局：{skip_note}'
    if dup_warns:
        head += f'，节间查重 {len(dup_warns)} 对疑似重复'
    return {'summary': head, 'warnings': warn_lines, 'notes': [skip_note] if skip_note else []}

def _validate_body_global(body_dir: Path) -> str:
    result = body_global_checks(body_dir)
    return '\n'.join([result['summary'], *[f'⚠️ {m}' for m in result['warnings']]])

def validate_body(section: str, block_ids: list[str] | None=None) -> str:
    """机器校验正文产物：写作指引表 / 单节正文 / body 全局承诺比对+节间查重。

    写完指引后传 section="body/写作指引.md" 校验表完整性（每节有行/模式合法/
    素材块可解析/同块多派弱提示——块是拷贝授权范围，同一块的不同区间可派不同节，
    各写手只注入自己需要的区间、同一内容区间不进两节，节间查重兜底）；写完一节
    正文后传该节路径与使用计划选用的素材块 id
    （docx 节按接受全部修订后的终稿视角校验，docx_comment_add 落的待办批注清点进
    收尾点名、内联占位判不过；定位段落为 P 段号、表格行为 T 行号）；
    整本收尾传 section="body"（目录）跑全局承诺比对与节间查重——承诺清单每个值
    未在任何正文出现、两节内容大面积重叠（同一素材进多节/原文照搬）都会被点名。
    [校验通过] 才算该节完成；[校验未通过] 逐条给出
    「文件:行号 现状→修复」。本轮写作范围内的就绪稿件复核失败会退回待校验，
    已有正文与旧合稿保留；所需章节全部通过才算阶段完成。内容检查不重写DOCX。
    明确选定的局部修订范围内，核对当前依据后通过校验可登记核验保留，
    无需为完成状态虚构文字修改；普通检查不刷新过期依据。
    Args:
        section: work/ 下的相对路径（如 body/写作指引.md、body/技术部分/3.1
                 项目理解.docx）；传 "body" 触发全局承诺比对
        block_ids: 本节使用计划选用的素材块 id（blk_…，search_references 的
                   evidence 里带；只注入了块内区间时写 blk_…:L起-L止 按区间算
                   使用率；校验指引与全局模式时忽略）
    """
    try:
        ctx = runctx.current_run()
        task_id = ctx.task_id if ctx else None
        if not task_id:
            return '[校验失败] 缺少任务上下文：正文产物在当前任务的 work/body/ 下'
        raw = normalize_work_path(section)
        wroot = work_dir(task_id).resolve()
        p = resolve_work_path(task_id, raw)
        if p.is_dir():
            return _validate_body_global(p)
        if not p.exists():
            return f'[校验失败] {raw} 不存在'
        if p.suffix == '.docx':
            issues, warnings, notes = validate_docx_file(p, block_ids, wroot / 'body')
        elif (guide := wroot / 'body' / body_contract.GUIDE_NAME).is_file() and p.samefile(guide):
            content_text = p.read_text(encoding='utf-8', errors='replace')
            issues, warnings, notes = _validate_guide(content_text, task_id)
        else:
            return '[校验失败] 正文仅支持 DOCX；写作指引使用原生 MD（或传 body 跑全局比对）'
        head = f'[校验通过] {raw}' if not issues else f'[校验未通过] {raw} 共 {len(issues)} 处问题（按提示修复后重新运行 validate_body，全部通过才算完成）：'
        parts = [head]
        parts.extend((f'- {i}' for i in issues))
        if issues and p.suffix == '.docx':
            point = next((match[1] for issue in issues if (match := re.match(r'^P(\d+)：', issue))), None)
            if point:
                parts.append(f'[定点回读] docx_section_read(path="{raw}", para={point})取得该P段完整原文；其他问题按实际P段号回读，不猜段号或表格位置。')
        parts.append(f'[待办批注/占位] 共 {len(notes)} 处（收尾汇报逐条点名，不得只字不提）：' if notes else '[待办批注/占位] 无')
        parts.extend((f'- {n}' for n in notes))
        parts.extend((f'⚠️ {w}' for w in warnings))
        return '\n'.join(parts)
    except Exception as e:
        return f'[校验失败] {type(e).__name__}: {e}'


def validate_docx_file(p: Path, block_ids: list[str] | None = None, body_root: Path | None = None):
    """Original DOCX terminal-view checks, also available for an external review input."""
    from ..parse.docx_text import load_text_document
    doc, _ = load_text_document(p)
    labeled = section_lines_labeled(doc)
    effective_blocks = list(dict.fromkeys(material_usage(doc) + (block_ids or [])))
    issues, warnings, notes = _validate_section([t for _, t in labeled], effective_blocks, labels=[lbl for lbl, _ in labeled])
    if (labeled and not any(t.strip() and t.strip() != p.stem for _, t in labeled)
            and not doc.tables and not list(doc.element.iter(qn('w:drawing')))):
        issues.append('正文仅有节标题，不能代替声明/函件全文；按本节原格式补齐正文，未知字段保留批注')
    # Joined native cell text contains synthetic pipes. Inspect actual cell
    # paragraphs when detecting literal Markdown tables instead.
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            if any(_MD_TABLE_ROW_RE.match(_accepted_text(p._p).strip()) for cell in row.cells for p in cell.paragraphs):
                issues.append(f'T{ti}R{ri}：单元格含竖线拼的 markdown 表格残字——请改为原生 Word 表格，不在单元格中填写 Markdown 表格文本')
    if body_root is not None:
        warnings.extend(_figure_plan_reconcile(doc, p, body_root))
    notes.extend((f"{loc}〔批注〕：{t[:50]}{('…' if len(t) > 50 else '')}" for loc, t in section_comments_labeled(doc)))
    outline_paras = [str(i) for i, para in enumerate(doc.paragraphs, 1) if (ppr := para._p.find(qn('w:pPr'))) is not None and ppr.find(qn('w:outlineLvl')) is not None]
    if outline_paras:
        shown = '、'.join((f'P{x}' for x in outline_paras[:6])) + ('…' if len(outline_paras) > 6 else '')
        notes.append(f'〔大纲级别〕{len(outline_paras)} 段直挂大纲级别（{shown}）——历史拷贝残留，Word 导航窗格会出现树外条目（合册会自动摘出）；如需正文层级请改用标题样式')
    heading_ids = _heading_style_ids(doc)
    numbered_style_ids: set[str] = set()
    for sid in heading_ids:
        style_el = doc.styles.element.find(f'''{qn('w:style')}[@{qn('w:styleId')}="{sid}"]''')
        sppr = style_el.find(qn('w:pPr')) if style_el is not None else None
        if sppr is not None and sppr.find(qn('w:numPr')) is not None:
            numbered_style_ids.add(sid)

    def _para_numbered(para) -> bool:
        ppr = para._p.find(qn('w:pPr'))
        if ppr is None:
            return False
        pstyle = ppr.find(qn('w:pStyle'))
        sid = pstyle.get(qn('w:val')) if pstyle is not None else None
        if ppr.find(qn('w:numPr')) is not None:
            return sid in heading_ids
        return sid in numbered_style_ids
    numbered_paras = [str(i) for i, para in enumerate(doc.paragraphs, 1) if _para_numbered(para)]
    if numbered_paras:
        shown = '、'.join((f'P{x}' for x in numbered_paras[:6])) + ('…' if len(numbered_paras) > 6 else '')
        notes.append(f'〔自动编号〕{len(numbered_paras)} 个标题段带自动编号（{shown}）——素材拷入残留，整本会渲染素材内部层级号（合册会自动剥除）；无需处理，继续写作即可')
    return issues, warnings, notes
