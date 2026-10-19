import json
import pytest

from app.services import outline_contract as oc
from app.tools.assemble_tender import assemble_tender, _parse_response_docs
from tests.test_assemble_tender import env


def aligned_structure(env):
    from pathlib import Path
    tid = env['task']['id']
    oc.plan_batch(tid, oc.PlanBatch())
    name = oc._load(tid)['baseline'][0]['name']
    fragment = Path(oc.dispatch(tid, name)['fragment_path'])
    fragment.parent.mkdir(parents=True, exist_ok=True)
    fragment.write_text((env['out'] / 'outline/tender-response-docs.md').read_text())
    oc.check(tid, name, child_completed=True)
    return fragment


def prepare(env):
    tid = env['task']['id']
    oc.begin(tid, 'action-1')
    aligned_structure(env)
    doc = _parse_response_docs((env['out'] / 'outline/tender-response-docs.md').read_text())[0]
    bindings = []
    for path, node in oc._nodes(doc['directory']):
        for rid in set([*node['来源位置'], *node['理由来源']]):
            bindings.append({'source_id': rid, 'path': list(path)})
    bindings.append({'source_id': 'TPL-01', 'path': ['技术方案（附件N）']})
    plan = oc.VolumePlan(name=doc['name'], root_titles=[n['目录名称'] for n in doc['directory']],
        headings=[{'source_id': 'SCORE-01', 'quote': '完整性', 'path': ['技术方案（附件N）', '完整性论述']},
                  {'source_id': 'SCORE-01', 'quote': '可行性', 'path': ['技术方案（附件N）', '可行性论述']}], bindings=bindings)
    excluded = [oc.Exclusion(source_id='SCORE-02', reason='报价评分不适用于本技术册，按本册范围排除')]
    oc.freeze_plan(tid, [plan], excluded, [oc.UnitDecision(unit_id="SCORE-01:1",kind="heading",reason="完整性与可行性属于需分别论述的实质分类",targets=[{"volume":plan.name,"path":h.path} for h in plan.headings]),oc.UnitDecision(unit_id="SCORE-02:1",kind="out_of_scope",reason="报价评审属于其他响应文件，不在技术册中重复提交",targets=[])])
    return tid, plan, excluded


def complete(env):
    tid, plan, excluded = prepare(env)
    p = env['out'] / 'outline/tender-response-docs.md'
    text = p.read_text().replace('  - 总体设计', '  - 总体设计\n  - 完整性论述\n  - 可行性论述')
    text = text.replace('## 来源标注\n', '## 来源标注\n- 完整性论述 :: SCORE-01\n- 可行性论述 :: SCORE-01\n')
    text = text.replace('技术方案（附件N） :: MAND-01, REQ-01', '技术方案（附件N） :: MAND-01, REQ-01, TPL-01')
    p.write_text(text)
    (env['out'] / 'outline/fragments' / (plan.name + '.md')).write_text(text)
    return tid, plan, excluded


def test_missing_categories_block_publish_and_do_not_replace_old_artifact(env):
    assert assemble_tender().startswith('[组装发布成功]')
    from app.services.page_flow import directory
    before = directory(env['task']['id'])
    tid, _, _ = prepare(env)
    with pytest.raises(ValueError, match='缺少实际标题路径'):
        oc.check(tid)
    assert assemble_tender().startswith('[组装失败]')
    assert directory(tid) == before


def test_main_merge_preserves_project_metadata_and_original_global_notes(env):
    file = env['out'] / 'outline/tender-response-docs.md'
    prefix = '## 项目信息\n- 项目名称：原项目名\n- 原有说明：仅为验收演示，缺证明不能当作真实资质\n\n'
    file.write_text(prefix + file.read_text())
    tid, _, _ = complete(env)
    assert oc.check(tid)['ok']
    file.write_text(file.read_text().replace(prefix, ''))
    assert assemble_tender().startswith('[组装失败]')
    with pytest.raises(ValueError, match='项目信息或全局说明缺失'):
        oc.require_publish(tid, {})


@pytest.mark.parametrize('mutation, message', [
    ('- 运维服务方案\n', '一级标题'),
    ('  - 需求理解\n', '原有节点'),
    ('概述=对招标需求的理解与响应', '原节点概述'),
])
def test_original_structure_and_content_are_hard_requirements(env, mutation, message):
    tid, _, _ = complete(env)
    p = env['out'] / 'outline/tender-response-docs.md'
    p.write_text(p.read_text().replace(mutation, '' if mutation.endswith('\n') else '概述=删改原说明'))
    with pytest.raises(ValueError):
        oc.check(tid)
    assert assemble_tender().startswith('[组装失败]')


def test_attachment_demotion_and_new_container_cannot_change_frozen_root(env):
    tid, _, _ = complete(env)
    p = env['out'] / 'outline/tender-response-docs.md'
    p.write_text(p.read_text().replace('- 技术方案（附件N）\n', '- 新建容器\n  - 技术方案（附件N）\n'))
    with pytest.raises(ValueError, match='一级标题'):
        oc.check(tid)


def test_plan_cannot_be_weakened_after_failure(env):
    tid, plan, excluded = prepare(env)
    plan.headings.pop()
    with pytest.raises(ValueError, match='不能因缺项'):
        oc.freeze_plan(tid, [plan], excluded, [oc.UnitDecision(unit_id="SCORE-01:1",kind="heading",reason="完整性与可行性属于需分别论述的实质分类",targets=[{"volume":plan.name,"path":h.path} for h in plan.headings]),oc.UnitDecision(unit_id="SCORE-02:1",kind="out_of_scope",reason="报价评审属于其他响应文件，不在技术册中重复提交",targets=[])])


def test_full_points_preserved_and_direct_json_publish_cannot_bypass(env):
    tid, _, _ = complete(env)
    assert oc.check(tid)['ok']
    result = assemble_tender()
    assert result.startswith('[组装发布成功]'), result
    from app.services.page_flow import directory
    from app import publish
    content = directory(tid)['content']
    assert content['registry']['SCORE-01']['评分要点'] == '完整性、可行性'
    content['response_documents'][0]['directory'].pop()
    with pytest.raises(publish.PublishError):
        publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid, source={'skill': 'direct'})


def test_mutation_after_pass_and_stale_input_are_rechecked(env):
    tid, _, _ = complete(env)
    assert oc.check(tid)['ok']
    p = env['out'] / 'analysis/evaluation.md'
    p.write_text(p.read_text() + '\n新评分依据\n')
    with pytest.raises(ValueError, match='依据已变化'):
        oc.check(tid)


def test_generic_project_heading_does_not_create_fake_volume(env):
    p = env['out'] / 'outline/tender-response-docs.md'
    docs = _parse_response_docs('# 投标响应文件分解\n\n' + p.read_text())
    assert [d['name'] for d in docs] == ['技术部分']


def test_empty_decisions_cannot_silently_skip_scoring_subitems(env):
    tid, plan, excluded = prepare(env)
    oc.begin(tid, 'new-action')
    aligned_structure(env)
    with pytest.raises(ValueError, match='全部片段逐项分类'):
        oc.freeze_plan(tid, [plan], excluded, [])


def test_multi_source_citations_are_not_cross_joined(tmp_path):
    a, b = '主文件.docx', '补充.docx'
    for name, lines in [(a, ['主一', '主二']), (b, ['补一', '补二'])]:
        p = tmp_path / 'parse' / name
        p.mkdir(parents=True)
        (p / (name + '.md')).write_text('\n'.join(lines))
    reg = {'SCORE-01': {'评分要点': '方案完整性', '出处': f'{a} L1；{b} L2'}}
    oc.enrich_registry(reg, tmp_path, [a,b])
    assert reg['SCORE-01']['原文依据'] == f'{a} L1: 主一\n{b} L2: 补二'


def test_body_dispatch_keeps_score_details_and_raw_obligations():
    from app.dispatch_enrich import _requirement_lines
    lines, _ = _requirement_lines(['SCORE-01'], {'依据': 0}, {'SCORE-01': {'type': '评分点', 'text': '方案', '评分要点': '四类方案', '原文依据': '原文中全部十五个具体内容项', '出处': '采购文件L2'}})
    assert '四类方案' in lines[0] and '全部十五个具体内容项' in lines[0]
    assert 'SCORE-01' not in lines[0]


def test_small_batches_preserve_originals_and_require_final_complete(env):
    tid, plan, excluded = prepare(env)
    previous = oc._load(tid)['plan']
    oc.begin(tid, 'batched-action')
    aligned_structure(env)
    patch = {'name':plan.name,
             'bindings':[{'source_id':'TPL-01','path':['技术方案（附件N）']}]}
    accepted = oc.plan_batch(tid, oc.PlanBatch.model_validate({'volumes':[patch], 'exclusions':[e.model_dump() for e in excluded], 'decisions':previous['decisions']}))
    assert accepted['ok'] and not accepted['frozen']
    assert accepted['pending_sources']==[] and accepted['pending_units']==[]
    assert oc.dispatch(tid,plan.name)['phase'] == 'structure_locked'
    assert assemble_tender().startswith('[组装失败]')
    assert oc.plan_batch(tid, oc.PlanBatch(complete=True))['frozen']
    assert oc.dispatch(tid,plan.name)['volume']['root_titles']==plan.root_titles


def test_batch_size_is_an_actual_schema_constraint():
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        oc.PlanBatch.model_validate({'volumes':[{'name':'x','bindings':[{'source_id':'REQ-01','path':['x']}]*9}]})


def test_invalid_plan_is_a_real_tool_error_with_short_actionable_unit_ids(env):
    from app.tools.outline import outline_plan
    prepare(env)
    with pytest.raises(ValueError, match='next_decision_unit_ids') as error:
        outline_plan(operation='decide',unit_id='missing')
    assert 'next_action' in str(error.value)
    assert 'input_value=' not in str(error.value)


def test_scalar_tool_appends_multiple_targets_without_reclassifying_or_dropping_evidence(env):
    from app.tools.outline import outline_plan
    tid, plan, _ = prepare(env)
    oc.begin(tid, 'scalar-tool')
    aligned_structure(env)
    reason = '完整性和可行性是本项目两个独立实质类别，分别设置对应标题'
    for name in ['完整性论述', '可行性论述']:
        result = outline_plan(operation='decide', unit_id='SCORE-01:1', kind='heading', reason=reason,
                              volume=plan.name, level1='技术方案（附件N）', level2=name)
        assert result['ok']
    current = oc._load(tid)['draft_plan']
    assert [t['path'][-1] for t in current['decisions'][0]['targets']] == ['完整性论述', '可行性论述']
    before = oc.path(tid).read_bytes()
    with pytest.raises(ValueError, match='kind及reason须保留'):
        outline_plan(operation='decide', unit_id='SCORE-01:1', kind='detail', reason=reason,
                     volume=plan.name, level1='技术方案（附件N）')
    assert oc.path(tid).read_bytes() == before


def test_scalar_tool_rejects_skipped_title_levels_and_encoded_arrays(env):
    from app.tools.outline import outline_plan
    tid, plan, _ = prepare(env)
    before = oc.path(tid).read_bytes()
    with pytest.raises(ValueError, match='不能跳级'):
        outline_plan(operation='bind', source_id='REQ-01', volume=plan.name, level3='孤立三级')
    with pytest.raises(ValueError, match='不传数组'):
        outline_plan(root_titles=['错误数组'])
    assert oc.path(tid).read_bytes() == before


def test_removed_original_source_is_reviewed_before_freeze_not_after_it(env):
    tid, plan, excluded = prepare(env)
    decisions = [oc.UnitDecision.model_validate(d) for d in oc._load(tid)['plan']['decisions']]
    oc.begin(tid, 'review-changed-sources')
    fragment = aligned_structure(env)
    original = fragment.read_text()
    fragment.write_text(original.replace('REQ-01', ''))
    before = oc.path(tid).read_bytes()
    with pytest.raises(ValueError, match='原来源移除尚未复核'):
        oc.freeze_plan(tid, [plan], excluded, decisions)
    assert oc.path(tid).read_bytes() == before
    fragment.write_text(original)
    # 新增的有效TPL绑定可在冻结后落实；本检查只针对旧来源移除，不禁止合法补充。
    assert oc.freeze_plan(tid, [plan], excluded, decisions)['frozen']


def test_model_context_retains_semantics_without_private_duplicate_raw_index(env):
    tid = env['task']['id']
    begun = oc.begin(tid, 'compact-context')
    context = json.loads(oc.context_path(tid).read_text())
    private = oc._load(tid)
    assert set(context['registry']) == set(private['registry'])
    assert context['baseline'] == private['baseline']
    assert context['units'] == {u:v for u,v in private['units'].items() if v['review_required']}
    assert set(context['optional_heading_unit_ids']) == {u for u,v in private['units'].items() if not v['review_required']}
    assert 'inputs' not in context and 'baseline_markdown' not in context
    assert '完整性、可行性' in oc.path(tid).with_suffix('.txt').read_text()
    assert begun['context_path'] == str(oc.context_path(tid))


def test_existing_body_headings_can_be_classified_with_verified_original_blocks(env):
    from docx import Document
    title = '技术方案（附件N）'
    body = env['out'] / 'body' / (title + '.docx')
    body.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); doc.add_heading(title, 1); doc.add_paragraph('已有正文不能缺失。')
    doc.add_heading('部署方式', 2); doc.add_paragraph('本地部署及接口说明。'); doc.save(body)
    tid, plan, excluded = prepare(env)
    context = json.loads(oc.context_path(tid).read_text())
    item = next(x for x in context['body_inventory'] if x['title'] == title)
    assert item['headings'][-1] == {'block': 3, 'title': '部署方式'}
    oc.begin(tid, 'body-heading-review')
    aligned_structure(env)
    previous = oc._load(tid)
    heading = {'path': [title, '部署方式'], 'body_path': 'body/' + title + '.docx', 'block': 3,
               'reason': '原技术方案已有部署方式说明，按实际内容归集到本章'}
    patch = {'name': plan.name, 'content_headings': [heading]}
    assert oc.plan_batch(tid, oc.PlanBatch.model_validate({'volumes': [patch]}))['ok']
    assert oc._load(tid)['draft_plan']['volumes'][0]['content_headings'] == [heading]
    assert oc._load(tid)['registry'] == previous['registry']


@pytest.mark.parametrize('mutation', ['wrong_block', 'invented_title', 'outside_parent', 'changed_source'])
def test_body_heading_evidence_cannot_be_fabricated_or_moved_to_other_chapters(env, mutation):
    from docx import Document
    title = '技术方案（附件N）'
    body = env['out'] / 'body' / (title + '.docx'); body.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); doc.add_heading(title, 1); doc.add_heading('部署方式', 2); doc.save(body)
    tid = env['task']['id']; oc.begin(tid, 'body-heading-evidence'); aligned_structure(env)
    heading = {'path': [title, '部署方式'], 'body_path': 'body/' + title + '.docx', 'block': 2,
               'reason': '原技术方案已有部署方式说明，按实际内容归集到本章'}
    if mutation == 'wrong_block': heading['block'] = 1
    if mutation == 'invented_title': heading['path'][-1] = '原文没有的标题'
    if mutation == 'outside_parent': heading['path'][0] = '运维服务方案'
    if mutation == 'changed_source':
        doc.add_paragraph('并发变更。'); doc.save(body)
    before = oc.path(tid).read_bytes()
    with pytest.raises(ValueError):
        oc.plan_batch(tid, oc.PlanBatch.model_validate({'volumes': [{'name': oc._load(tid)['baseline'][0]['name'], 'content_headings': [heading]}]}))
    assert oc.path(tid).read_bytes() == before


def test_bindings_cannot_sneak_in_parameter_headings(env):
    tid, plan, _ = prepare(env)
    oc.begin(tid, 'parameter-plan')
    with pytest.raises(ValueError, match='新增标题须先作heading决策'):
        oc.plan_batch(tid, oc.PlanBatch.model_validate({'volumes':[{'name':plan.name,'bindings':[{'source_id':'REQ-01','path':['技术方案（附件N）','服务期限三个月']}]}]}))


def test_unplanned_titles_block_publication_even_without_source_ids(env):
    tid, _, _ = complete(env)
    p = env['out'] / 'outline/tender-response-docs.md'
    p.write_text(p.read_text().replace('  - 总体设计', '  - 无依据附加标题\n  - 总体设计'))
    with pytest.raises(ValueError, match='新增标题不在冻结分类计划'):
        oc.check(tid)


def test_optional_general_requirement_can_plan_a_real_new_heading(env):
    tid, plan, _ = prepare(env)
    oc.begin(tid, 'new-requirement-heading')
    aligned_structure(env)
    result = oc.plan_batch(tid, oc.PlanBatch.model_validate({'decisions':[{'unit_id':'REQ-01:1','kind':'heading','reason':'依据本项目普通需求中具有独立实质内容的分类新增对应章节','targets':[{'volume':plan.name,'path':['技术方案（附件N）','独立需求响应']}]}]}))
    assert result['ok']
    assert oc._load(tid)['draft_plan']['volumes'][0]['headings'][0]['source_id'] == 'REQ-01'


def test_literal_double_quotes_in_classification_are_preserved(env):
    tid = env['task']['id']
    p = env['out'] / 'analysis/evaluation.md'
    p.write_text(p.read_text().replace('完整性、可行性', '完整性与"可行性"'))
    oc.begin(tid, 'quoted-plan')
    aligned_structure(env)
    doc = _parse_response_docs((env['out'] / 'outline/tender-response-docs.md').read_text())[0]
    result = oc.plan_batch(tid, oc.PlanBatch.model_validate({'decisions':[{'unit_id':'SCORE-01:1','kind':'heading','reason':'按本项目独立评分分类建立实际标题且保留原文字面引号','targets':[{'volume':doc['name'],'path':['技术方案（附件N）','完整性与可行性']}]}]}))
    assert result['ok']
    assert '"可行性"' in oc._load(tid)['draft_plan']['volumes'][0]['headings'][0]['quote']


def test_child_can_align_actual_titles_before_semantic_freeze_but_cannot_publish(env):
    tid = env['task']['id']
    oc.begin(tid, 'structure-first')
    oc.plan_batch(tid, oc.PlanBatch())
    name = oc._load(tid)['baseline'][0]['name']
    assignment = oc.dispatch(tid, name)
    assert assignment['phase'] == 'structure_locked'
    from pathlib import Path
    fragment = Path(assignment['fragment_path'])
    fragment.parent.mkdir(parents=True, exist_ok=True)
    fragment.write_text((env['out'] / 'outline/tender-response-docs.md').read_text().replace('  - 总体设计', '  - 新增独立评分标题\n  - 总体设计'))
    assert oc.check(tid, name)['phase'] == 'structure_only'
    assert assemble_tender().startswith('[组装失败]')
    fragment.write_text(fragment.read_text().replace('- 技术方案（附件N）\n','- 新容器\n  - 技术方案（附件N）\n'))
    with pytest.raises(ValueError, match='一级标题'):
        oc.check(tid, name)


def test_scoring_decisions_require_real_child_completion_not_main_structure_check(env):
    tid = env['task']['id']
    oc.begin(tid, 'enforce-child-order')
    oc.plan_batch(tid, oc.PlanBatch())
    name = oc._load(tid)['baseline'][0]['name']
    from pathlib import Path
    fragment = Path(oc.dispatch(tid, name)['fragment_path'])
    fragment.parent.mkdir(parents=True, exist_ok=True)
    fragment.write_text((env['out'] / 'outline/tender-response-docs.md').read_text())
    assert oc.check(tid, name)['phase'] == 'structure_only'
    assert not oc._load(tid)['aligned_volumes']
    with pytest.raises(ValueError, match='真实子代理'):
        oc.plan_batch(tid, oc.PlanBatch(complete=True))
    assert oc.check(tid, name, child_completed=True)['ok']
    assert oc._load(tid)['aligned_volumes'] == [name]


def test_unknown_volume_cannot_fake_child_completion(env):
    tid = env['task']['id']
    oc.begin(tid, 'invalid-volume')
    oc.plan_batch(tid, oc.PlanBatch())
    with pytest.raises(ValueError, match='不属于'):
        oc.check(tid, '不存在的册', child_completed=True)
    assert not oc._load(tid)['aligned_volumes']


@pytest.mark.parametrize('operation', ['decide', 'complete'])
def test_another_volume_child_cannot_satisfy_this_volume_scoring_review(env, operation):
    from pathlib import Path
    tid = env['task']['id']
    file = env['out'] / 'outline/tender-response-docs.md'
    text = file.read_text()
    first = _parse_response_docs(text)[0]['name']
    second = first + '补充册'
    file.write_text(text + '\n\n' + text.replace('响应文件：' + first, '响应文件：' + second).replace('响应文件: ' + first, '响应文件: ' + second))
    oc.begin(tid, 'per-volume-child-review')
    oc.plan_batch(tid, oc.PlanBatch())
    fragment = Path(oc.dispatch(tid, first)['fragment_path'])
    assert oc.check(tid, first, child_completed=True)['ok']
    assert oc._load(tid)['aligned_volumes'] == [first]
    before = oc.path(tid).read_bytes()
    if operation == 'decide':
        batch = oc.PlanBatch.model_validate({'decisions':[{'unit_id':'SCORE-01:1','kind':'heading',
            'reason':'本册也需独立核对完整性与可行性评分分类',
            'targets':[{'volume':second,'path':['技术方案（附件N）','完整性论述']}]}]})
        with pytest.raises(ValueError, match='对应册尚未完成真实子代理'):
            oc.plan_batch(tid, batch)
    else:
        draft = oc._load(tid)['draft_plan']
        draft['volumes'][1]['headings'].append({'source_id':'SCORE-01','quote':'完整性','path':['技术方案（附件N）','完整性论述']})
        with pytest.raises(ValueError, match='对应册尚未完成真实子代理'):
            oc.freeze_plan(tid, [oc.VolumePlan.model_validate(v) for v in draft['volumes']], [], [])
    assert oc.path(tid).read_bytes() == before
    assert fragment.is_file()


def test_unchanged_price_rubric_does_not_require_rewriting_another_volume(env):
    tid = env['task']['id']
    file = env['out'] / 'outline/tender-response-docs.md'
    text = file.read_text()
    first = _parse_response_docs(text)[0]['name']
    second = first + '补充册'
    file.write_text(text + '\n\n' + text.replace('响应文件：' + first, '响应文件：' + second).replace('响应文件: ' + first, '响应文件: ' + second))
    oc.begin(tid, 'unchanged-price-volume')
    oc.plan_batch(tid, oc.PlanBatch())
    oc.dispatch(tid, first)
    oc.check(tid, first, child_completed=True)
    result = oc.plan_batch(tid, oc.PlanBatch.model_validate({'decisions':[{'unit_id':'SCORE-02:1','kind':'rubric',
        'reason':'价格评分公式作为原报价材料的评审规则，不新增实质分类标题',
        'targets':[{'volume':second,'path':['技术方案（附件N）']}]}]}))
    assert not result['frozen']
    assert oc._load(tid)['aligned_volumes'] == [first]
    assert not oc.fragment_path(tid, second).exists()


@pytest.mark.parametrize('tail,valid', [(['完整性论述'],False), (['组织方案'],True), (['组织方案','完整性论述'],True)])
def test_explicit_composite_scheme_categories_control_score_dimension_depth(env, tail, valid):
    root = '综合材料（组织方案、应急方案）'
    p = env['out'] / 'outline/tender-response-docs.md'
    p.write_text(p.read_text().replace('技术方案（附件N）', root))
    tid = env['task']['id']
    oc.begin(tid, 'composite-category-depth')
    aligned_structure(env)
    doc = oc._load(tid)['baseline'][0]
    plan = oc.VolumePlan(name=doc['name'],root_titles=[n['目录名称'] for n in doc['directory']],headings=[{'source_id':'SCORE-01','quote':'完整性','path':[root,*tail]}],bindings=[])
    if valid:
        assert oc.freeze_plan(tid,[plan],[],[],complete=False)['ok']
    else:
        with pytest.raises(ValueError,match='多个方案类别'):
            oc.freeze_plan(tid,[plan],[],[],complete=False)


def test_frozen_plan_can_add_verified_sources_without_changing_requirements(env):
    tid, plan, _ = complete(env)
    before = oc._load(tid)['plan']
    result = oc.plan_batch(tid, oc.PlanBatch.model_validate({'volumes':[{'name':plan.name,'bindings':[{'source_id':'REQ-01','path':['技术方案（附件N）','完整性论述']}]}]}))
    assert result['frozen']
    after = oc._load(tid)['plan']
    assert before['decisions'] == after['decisions'] and before['exclusions'] == after['exclusions']
    assert before['volumes'][0]['headings'] == after['volumes'][0]['headings']
    assert all(b in after['volumes'][0]['bindings'] for b in before['volumes'][0]['bindings'])
    assert oc._load(tid)['binding_additions_audit'] == [{'volume':plan.name,'source_id':'REQ-01','path':['技术方案（附件N）','完整性论述']}]
    with pytest.raises(ValueError,match='不能减少或修改'):
        oc.plan_batch(tid, oc.PlanBatch.model_validate({'decisions':[before['decisions'][0]]}))


def test_frozen_source_append_cannot_create_an_unreviewed_title(env):
    tid, plan, _ = complete(env)
    with pytest.raises(ValueError,match='新增标题须先作heading决策'):
        oc.plan_batch(tid,oc.PlanBatch.model_validate({'volumes':[{'name':plan.name,'bindings':[{'source_id':'REQ-01','path':['技术方案（附件N）','参数标题']}]}]}))


def test_frozen_source_append_cannot_drop_previously_approved_associations(env):
    tid, _, _ = complete(env)
    before = oc._load(tid)['plan']
    volumes = [oc.VolumePlan.model_validate(v) for v in before['volumes']]
    volumes[0].bindings.pop()
    with pytest.raises(ValueError,match='删除或迁移'):
        oc.freeze_plan(tid,volumes,[oc.Exclusion.model_validate(e) for e in before['exclusions']],[oc.UnitDecision.model_validate(d) for d in before['decisions']])


def test_dispatch_seeds_existing_directory_before_model_edits_without_touching_original(env):
    tid = env['task']['id']
    original = env['out'] / 'outline/tender-response-docs.md'
    text, data = original.read_text(), original.read_bytes()
    oc.begin(tid,'seed-working-copy')
    oc.plan_batch(tid,oc.PlanBatch())
    name = oc._load(tid)['baseline'][0]['name']
    assignment = oc.dispatch(tid,name)
    from pathlib import Path
    assert Path(assignment['fragment_path']).read_text() == text
    assert original.read_bytes() == data


def test_fragment_names_reuse_safe_naming_and_do_not_escape_task(env):
    tid = env['task']['id']
    expected = env['out'] / 'outline/fragments'
    paths = [oc.fragment_path(tid,n) for n in ['../../正式项目','技术/资格册','技术:资格册']]
    assert all(p.parent == expected and p.suffix == '.md' for p in paths)
    assert len(set(paths)) == len(paths)
    assert oc.fragment_path(tid,'技术册').name == '技术册.md'
