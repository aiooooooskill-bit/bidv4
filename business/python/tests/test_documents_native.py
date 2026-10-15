import json
from pathlib import Path
import pytest
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas
from app.document import prepare,finish,render_ocr_page

TEXT='本项目采购运行维护服务。供应商应提供项目团队、实施计划与应急响应机制，服务期限三年。招标要求与投标承诺均应逐项回应。'*4

def test_docx_cache_and_changed_source(tmp_path):
    src=tmp_path/'采购.docx';out=tmp_path/'parse';d=Document();d.add_heading('采购说明',1);d.add_paragraph(TEXT);d.save(src)
    result=prepare(src,out);assert len(result['paths'])==3
    assert prepare(src,out)['cached'] is True
    d.add_paragraph('补充来源说明');d.save(src);assert prepare(src,out)['meta']['sha256']!=result['meta']['sha256']


def test_short_native_word_chapter_and_empty_document(tmp_path):
    source=tmp_path/'封面.docx';d=Document();d.add_paragraph('采购项目投标文件（验收演示）');d.save(source)
    result=prepare(source,tmp_path/'cover-parse')
    assert Path(result['paths'][0]).read_text()=='采购项目投标文件（验收演示）'
    empty=tmp_path/'空文件.docx';Document().save(empty)
    with pytest.raises(ValueError,match='内容为空'):prepare(empty,tmp_path/'empty-parse')


def test_word_parser_reads_current_revision_text_without_changing_source_indexes(tmp_path):
    from docx.oxml import OxmlElement
    source=tmp_path/'修订标书.docx';d=Document();p=d.add_paragraph('供应商：')
    def change(p,kind,text):
        node=OxmlElement('w:'+kind);run=OxmlElement('w:r');value=OxmlElement('w:delText' if kind=='del' else 'w:t');value.text=text;run.append(value);node.append(run);p._p.append(node)
    change(p,'del','旧公司');change(p,'ins','新公司（验收演示）')
    table=d.add_table(rows=1,cols=1);p=table.cell(0,0).paragraphs[0]
    change(p,'del','旧承诺');change(p,'ins','三年维护')
    d.add_paragraph('签署日期：2026年8月19日');d.save(source);original=source.read_bytes()
    result=prepare(source,tmp_path/'parse');md=Path(result['paths'][0]).read_text()
    assert '新公司（验收演示）' in md and '三年维护' in md
    assert '旧公司' not in md and '旧承诺' not in md
    assert [row[0] for row in result['meta']['element_lines']]==[0,1,2]
    assert source.read_bytes()==original


def test_native_word_cache_requires_current_text_view(tmp_path):
    source=tmp_path/'声明.docx';d=Document();d.add_paragraph('当前有效声明文字');d.save(source)
    out=tmp_path/'parse';result=prepare(source,out)
    meta_path=Path(result['paths'][2]);meta=json.loads(meta_path.read_text());meta.pop('text_view');meta_path.write_text(json.dumps(meta))
    Path(result['paths'][0]).write_text('遗漏新增文字的旧缓存')
    refreshed=prepare(source,out)
    assert not refreshed.get('cached')
    assert Path(refreshed['paths'][0]).read_text()=='当前有效声明文字'
    assert prepare(source,out)['cached']

def test_docx_text_and_terminal_checks_do_not_read_binary_parts(tmp_path,monkeypatch):
    from zipfile import ZipFile
    from app.tools.validate_body import validate_docx_file
    src=tmp_path/'带图带批注.docx';image=tmp_path/'图.png'
    Image.new('RGB',(100,100),'white').save(image)
    d=Document();d.add_heading('服务说明',1);p=d.add_paragraph(TEXT)
    d.add_comment(p.runs[0],text='容量检查备注',author='工程测试')
    d.add_picture(str(image));d.add_table(rows=1,cols=1).cell(0,0).text='表格承诺'
    d.save(src)
    read=ZipFile.read
    def read_metadata(self,name,*args,**kw):
        assert str(name).endswith('.rels') or not str(name).startswith('word/media/')
        return read(self,name,*args,**kw)
    monkeypatch.setattr(ZipFile,'read',read_metadata)
    result=prepare(src,tmp_path/'parse')
    assert result['meta']['image_count']==1
    md=Path(result['paths'][0]).read_text()
    assert '# 服务说明' in md and '表格承诺' in md and '![](图片)' in md
    assert any('容量检查备注' in n for n in validate_docx_file(src)[2])

@pytest.mark.parametrize('ext',['.png','.jpg','.jpeg','.gif','.webp'])
def test_image_uses_dsh_ocr_plan_and_keeps_original(tmp_path,ext):
    src=tmp_path/('采购'+ext);Image.new('RGB',(100,100),'white').save(src);original=src.read_bytes()
    result=prepare(src,tmp_path/'parse');assert result['needs_ocr'] and len(result['pages'])==1
    render_ocr_page(result['plan'],1)
    assert Path(result['pages'][0]['path']).suffix=='.png'
    with pytest.raises(ValueError,match='OCR页面未完整'):finish(result['plan'],[])
    completed=finish(result['plan'],[{'page':1,'text':TEXT}]);assert len(completed['paths'])==3;assert src.read_bytes()==original
    assert prepare(src,tmp_path/'parse')['cached']

def test_scanned_pdf_requires_every_page_in_order(tmp_path):
    src=tmp_path/'扫描.pdf';c=canvas.Canvas(str(src));c.showPage();c.showPage();c.save()
    result=prepare(src,tmp_path/'parse');assert [p['page'] for p in result['pages']]==[1,2]
    with pytest.raises(ValueError,match='OCR页面未完整'):finish(result['plan'],[{'page':2,'text':TEXT},{'page':1,'text':TEXT}])
    done=finish(result['plan'],[{'page':1,'text':TEXT},{'page':2,'text':TEXT}]);md=Path(done['paths'][0]).read_text();assert '<!-- p:1 -->' in md and '<!-- p:2 -->' in md
    assert done['meta']['ocr_pages']==[1,2]

def test_source_change_during_ocr_never_publishes(tmp_path):
    src=tmp_path/'采购.png';Image.new('RGB',(100,100)).save(src);out=tmp_path/'parse';p=prepare(src,out);src.write_bytes(b'changed')
    with pytest.raises(ValueError,match='来源文件已变化'):finish(p['plan'],[{'page':1,'text':TEXT}])
    assert not (out/(src.name+'.meta.json')).exists()

def test_mixed_pdf_keeps_digital_text_and_ocr_page(tmp_path):
    src=tmp_path/'混合.pdf';c=canvas.Canvas(str(src))
    for i in range(10):c.drawString(50,750-i*20,'Original digital procurement requirements '+str(i))
    c.showPage();c.showPage();c.save()
    result=prepare(src,tmp_path/'parse');assert [p['page'] for p in result['pages']]==[2]
    done=finish(result['plan'],[{'page':2,'text':TEXT}]);md=Path(done['paths'][0]).read_text()
    assert 'Original digital procurement requirements' in md and TEXT in md

def test_scan_over_128_pages_is_lazy_and_requires_complete_outputs(tmp_path):
    src=tmp_path/'长扫描.pdf';c=canvas.Canvas(str(src))
    for _ in range(129):c.showPage()
    c.save();out=tmp_path/'parse'
    result=prepare(src,out)
    assert len(result['pages'])==129
    assert not list((out/'images').glob('*.png'))
    render_ocr_page(result['plan'],129)
    assert list((out/'images').glob('*.png'))==[out/'images/page-129.png']
    outputs=[{'page':p,'text':TEXT} for p in range(1,130)]
    with pytest.raises(ValueError,match='OCR页面未完整'):finish(result['plan'],outputs[:-1])
    assert not (out/(src.name+'.meta.json')).exists()
    completed=finish(result['plan'],outputs)
    assert completed['meta']['ocr_pages']==list(range(1,130))
    assert '<!-- p:129 -->' in Path(completed['paths'][0]).read_text()

def test_image_budget_checked_before_conversion(tmp_path):
    src=tmp_path/'大图.png';Image.new('RGB',(5000,4001)).save(src);out=tmp_path/'parse'
    result=prepare(src,out)
    with pytest.raises(ValueError,match='2000万像素'):render_ocr_page(result['plan'],1)
    assert not (out/'images/page-1.png').exists()


def test_native_tables_are_not_markdown_but_literal_cell_residue_is(tmp_path):
    from app.tools.validate_body import validate_docx_file
    source = tmp_path / '原生表格.docx'
    d = Document(); table = d.add_table(rows=1, cols=4)
    table.cell(0, 1).text = '产品名称'; table.cell(0, 2).text = '实际承诺'
    d.save(source)
    assert not any('markdown 表格残字' in s for s in validate_docx_file(source)[0])
    table.cell(0, 2).text = '| 名称 | 内容 |'
    d.save(source)
    assert any('T1R1：单元格含竖线拼' in s for s in validate_docx_file(source)[0])
