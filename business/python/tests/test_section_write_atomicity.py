import json,subprocess,sys,threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import pytest
from docx import Document
from app import artifact_store,publish,runctx
from app.services import build_state,page_flow,source_usage
from app.tools.docx_ops import docx_section_revise,docx_section_create,section_text_lines
from tests.util import init_env

@pytest.fixture
def env(tmp_path,monkeypatch):
 task,conv=init_env(tmp_path,monkeypatch);tid=task['id'];artifact_store.ensure_task_skeleton(tid);root=artifact_store.work_dir(tid)
 publish.publish_artifact('tender.directory/tender-response-docs@1',{'response_documents':[{'name':'技术册','directory':[{'目录名称':'实施方案','交付形态':'正文编写'}]}]},task_id=tid)
 path=root/'body/实施方案.docx';doc=Document();doc.add_paragraph('原始服务方案');doc.save(path)
 source=artifact_store.sources_dir(tid)/'本次采购.docx';source.write_bytes(b'real-source-version-one')
 source_usage.record_read(tid,'parent',{'name':source.name,'sha256':page_flow.file_hash(source)});build_state.record_write(tid,'body/实施方案.docx','parent')
 runctx.set_run('parent','atomic-test',tid)
 yield tid,root,source
 runctx.clear_run()

def request(env,name='docx_section_revise',parameters=None):
 tid,root,_=env
 if parameters is None:parameters={'path':'body/实施方案.docx','edits':[{'para':1,'action':'replace','find':'原始服务方案','text':'修订后服务方案'}]}
 payload={'operation':'tool','args':{'name':name,'parameters':parameters},'data_root':str(root.parents[2]),'task_id':tid,'session_id':'parent','call_id':'atomic-rpc'}
 x=subprocess.run([sys.executable,str(Path(__file__).parents[1]/'runner.py')],input=json.dumps(payload),text=True,capture_output=True,timeout=15)
 return x.returncode,json.loads(x.stdout)

@pytest.mark.parametrize('change',['missing','modified'])
def test_rpc_invalid_reference_never_publishes_docx_or_state(env,change):
 tid,root,source=env;path=root/'body/实施方案.docx';before=path.read_bytes();state=(root/'body/.build-state.json').read_bytes()
 if change=='missing':source.unlink()
 else:source.write_bytes(b'version-two')
 code,result=request(env);assert code!=0 and '参考来源发生变化' in result['error'];assert path.read_bytes()==before;assert (root/'body/.build-state.json').read_bytes()==state;assert not list(root.glob('.section-write-*'))

def test_rpc_valid_revision_records_draft_and_then_validation(env):
 tid,root,_=env;code,result=request(env);assert code==0,result;assert '[已修订]' in result['result'];assert '修订后服务方案' in ''.join(section_text_lines(Document(root/'body/实施方案.docx')))
 value=build_state.read(tid);assert value['pending_sections']['body/实施方案.docx']['hash']==page_flow.file_hash(root/'body/实施方案.docx');assert value['source_dependencies']['body/实施方案.docx']

def test_source_changes_during_draft_generation_leaves_original_unchanged(env):
 from app.services.section_write import execute
 tid,root,source=env;path=root/'body/实施方案.docx';original=path.read_bytes();state=(root/'body/.build-state.json').read_bytes()
 def operation():
  out=docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'原始服务方案','text':'新方案'}]);source.write_bytes(b'changed-during-generation');return out
 with pytest.raises(ValueError,match='参考来源发生变化'):execute(tid,'body/实施方案.docx','parent',None,False,operation)
 assert path.read_bytes()==original;assert (root/'body/.build-state.json').read_bytes()==state

def test_state_save_failure_after_file_replace_restores_original(env,monkeypatch):
 from app.services.section_write import execute
 tid,root,_=env;path=root/'body/实施方案.docx';original=path.read_bytes();state=(root/'body/.build-state.json').read_bytes()
 def fail(*a,**k):raise OSError('injected metadata write failure')
 monkeypatch.setattr(build_state,'save',fail)
 with pytest.raises(OSError,match='metadata write failure'):execute(tid,'body/实施方案.docx','parent',None,False,lambda:docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'原始服务方案','text':'新方案'}]))
 assert path.read_bytes()==original;assert (root/'body/.build-state.json').read_bytes()==state

def test_failed_rebuild_does_not_rotate_original_restorepoints(env):
 tid,root,source=env;path=root/'body/实施方案.docx';original=path.read_bytes();source.unlink()
 code,result=request(env,'docx_section_create',{'path':'body/实施方案.docx','title':'实施方案','paragraphs':'重新编写','replace':True});assert code!=0;assert path.read_bytes()==original;assert not path.with_name(path.name+'.restorepoints').exists()

def test_successful_rebuild_keeps_actual_original_restorepoint(env):
 tid,root,_=env;path=root/'body/实施方案.docx';original=path.read_bytes()
 code,result=request(env,'docx_section_create',{'path':'body/实施方案.docx','title':'实施方案','paragraphs':'重新编写','replace':True});assert code==0,result
 points=list(path.with_name(path.name+'.restorepoints').glob('*.bak'));assert len(points)==1 and points[0].read_bytes()==original

def test_new_nested_docx_is_published_and_tracked(env):
 tid,root,_=env;code,result=request(env,'docx_section_create',{'path':'body/新册/新章节.docx','title':'新章节','paragraphs':'新内容'});assert code==0,result;assert (root/'body/新册/新章节.docx').is_file()

def test_concurrent_same_chapter_rejects_stale_draft_instead_of_losing_update(env):
 from app.services.section_write import execute
 tid,root,_=env;barrier=threading.Barrier(2);results=[]
 def worker(text):
  runctx.set_run('parent',text,tid)
  def operation():
   out=docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'原始服务方案','text':text}]);barrier.wait(timeout=5);return out
  try:return execute(tid,'body/实施方案.docx','parent',None,False,operation)
  except ValueError as e:return str(e)
  finally:runctx.clear_run()
 with ThreadPoolExecutor(2) as pool:results=list(pool.map(worker,['写手甲方案','写手乙方案']))
 assert sum('[已修订]' in x for x in results)==1;assert sum('正文已由其他写入更新' in x for x in results)==1
 assert not list(root.glob('.section-write-*'))

def test_invalid_edit_is_not_published_and_leaves_no_staging(env):
 tid,root,_=env;path=root/'body/实施方案.docx';original=path.read_bytes();state=(root/'body/.build-state.json').read_bytes()
 code,result=request(env,parameters={'path':'body/实施方案.docx','edits':[{'para':999,'action':'delete'}]});assert code==0 and '[修订失败]' in result['result'];assert path.read_bytes()==original;assert (root/'body/.build-state.json').read_bytes()==state;assert not list(root.glob('.section-write-*'))


def test_reference_changes_after_preflight_rolls_back_docx_and_ledger(env,monkeypatch):
 from app.services.section_write import execute
 from app.services import asset_usage
 tid,root,source=env;path=root/'body/实施方案.docx';original=path.read_bytes();state=(root/'body/.build-state.json').read_bytes();attach=asset_usage.attach_reads;calls=0
 def change_at_second_check(*args,**kwargs):
  nonlocal calls
  calls+=1
  if calls==2:source.write_bytes(b'changed-after-preflight')
  return attach(*args,**kwargs)
 monkeypatch.setattr(asset_usage,'attach_reads',change_at_second_check)
 with pytest.raises(ValueError,match='参考来源发生变化'):execute(tid,'body/实施方案.docx','parent',None,False,lambda:docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'原始服务方案','text':'新方案'}]))
 assert calls==2;assert path.read_bytes()==original;assert (root/'body/.build-state.json').read_bytes()==state
