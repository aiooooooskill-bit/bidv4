import json
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from app import artifact_store,runctx
from app.parse import convert
from app.tools.docx_ops import docx_section_create,docx_source_inject,docx_section_read
from tests.util import init_env

@pytest.mark.parametrize('styled',[True,False])
def test_soft_breaks_keep_markdown_and_source_injection_coordinates(tmp_path,monkeypatch,styled):
    task,_=init_env(tmp_path,monkeypatch);tid=task['id'];artifact_store.ensure_task_skeleton(tid)
    source=artifact_store.sources_dir(tid)/'格式.docx';doc=Document()
    doc.add_paragraph('前言第一行\n前言第二行\n前言第三行')
    if styled:doc.add_heading('投标函',1)
    else:doc.add_paragraph('一、投标函\n补充说明')
    doc.add_paragraph('致：招标代理机构');doc.add_paragraph('投标供应商地址：')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='带换行的表格\n内容';table.cell(0,1).text='已填写'
    doc.add_heading('下一份附件',1);doc.add_paragraph('不能被拷入本章的正文')
    doc.save(source);parsed=convert(source);lines=parsed.md.splitlines()
    start=next(n for n,line in enumerate(lines,1) if line=='致：招标代理机构')
    end=start+1
    runctx.set_run('line-map','inject',tid)
    try:
        assert docx_section_create(path='body/投标函.docx',title='投标函').startswith('[已创建]')
        result=docx_source_inject(source=source.name,dest='body/投标函.docx',lines=f'L{start}-L{end}')
        assert result.startswith('[已注入]'),result
        view=docx_section_read(path='body/投标函.docx')
        assert '致：招标代理机构' in view and '投标供应商地址：' in view
        assert '不能被拷入本章' not in view and '带换行的表格' not in view
    finally:runctx.clear_run()


@pytest.mark.parametrize('soft_break',[False,True])
def test_numbered_heading_discards_list_prefix_and_preserves_line_map(tmp_path,soft_break):
    from app.parse import outline_with_lines
    source=tmp_path/'编号标题.docx';doc=Document()
    p=doc.add_paragraph('一、投标函'+('\n补充说明' if soft_break else ''))
    p._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'))
    doc.add_paragraph('致：招标代理机构');doc.save(source)
    parsed=convert(source)
    assert parsed.info['conversion']=='docx-numbered'
    assert parsed.md.splitlines()[0]=='## 一、投标函'
    assert outline_with_lines(parsed.md)[0]['标题']=='一、投标函'
    expected_start=3 if soft_break else 2
    assert [1,expected_start,expected_start] in parsed.info['element_lines']
    if soft_break:assert parsed.md.splitlines()[1]=='补充说明'
