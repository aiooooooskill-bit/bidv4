import asyncio
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from app import artifact_store, db
from app.services import build_state, outline_contract, page_flow
from app.services.artifacts import ContentUpdate, update_artifact_content
from app.services.directory_edits import SectionRename
from app.services.errors import BusinessError
from app.tools.assemble_tender import assemble_tender
from tests.test_outline_contract import complete, env


def published(env):
    artifact_store.ensure_task_skeleton(env['task']['id'])
    tid, _, _ = complete(env)
    assert assemble_tender().startswith('[组装发布成功]')
    assert outline_contract.status(tid, 'action-1')['ok']
    return tid, page_flow.directory(tid)


def save(current, content, renames=(), **kwargs):
    return asyncio.run(update_artifact_content(current['artifact_id'], ContentUpdate(
        content=content, base_content_seq=current['content_seq'], renamed_sections=renames, **kwargs)))


def body(tid, title, pending=False):
    path = 'body/' + title + '.docx'
    file = artifact_store.work_dir(tid) / path
    file.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(title, 1)
    document.add_paragraph('人工已填写的正文、事实12345及承诺完整保留。')
    document.add_table(rows=1, cols=2).cell(0, 0).text = '原始表格98765'
    document.add_comment(document.paragraphs[1].runs[0], text='原始核对批注', author='人工')
    document.save(file)
    if pending:
        build_state.record_write(tid, path)
    else:
        build_state.record_section(tid, path)
    return file, file.read_bytes()


def test_saved_directory_is_effective_without_rewriting_ai_working_draft(env):
    tid, current = published(env)
    file = env['out'] / 'outline/tender-response-docs.md'
    before = file.read_bytes()
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'].reverse()
    save(current, content)
    accepted = page_flow.directory(tid)
    assert accepted['content'] == content
    assert accepted['user_input_version'] == page_flow.input_versions(tid)['outline']
    assert accepted['user_saved_at'] > 0
    assert file.read_bytes() == before
    # AI 发布自身的校验仍然严格，页面采用独立的人工保存依据。
    assert not outline_contract.status(tid, 'action-1')['ok']
    (env['out'] / 'analysis/evaluation.md').write_text('招标评分依据发生变化')
    assert accepted['user_input_version'] != page_flow.input_versions(tid)['outline']


@pytest.mark.parametrize('pending', [False, True])
def test_rename_and_move_preserve_exact_word_bytes_and_chapter_records(env, pending):
    tid, current = published(env)
    old, original = body(tid, '需求理解', pending)
    state_before = build_state.read(tid)
    content = deepcopy(current['content'])
    roots = content['response_documents'][0]['directory']
    section = roots[0]['children'].pop(0)
    section['目录名称'] = '人工确认的需求分析'
    roots[1].setdefault('children', []).append(section)
    save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='人工确认的需求分析')])
    path = 'body/人工确认的需求分析.docx'
    assert not old.exists()
    assert (env['out'] / path).read_bytes() == original
    state = build_state.read(tid)
    key = 'pending_sections' if pending else 'sections'
    assert state[key][path] == state_before[key]['body/需求理解.docx']
    assert 'body/需求理解.docx' not in state[key]
    snapshot = build_state.snapshot(tid)
    assert next(s for s in snapshot['body_sections'] if s['title'] == '人工确认的需求分析')['status'] == ('draft' if pending else 'stale')
    assert snapshot['orphan_body_files'] == []
    assert state['section_titles'][path] == ['需求理解']


def test_move_without_rename_keeps_parent_and_child_bodies(env):
    tid, current = published(env)
    parent, parent_bytes = body(tid, '技术方案（附件N）')
    child, child_bytes = body(tid, '需求理解')
    content = deepcopy(current['content'])
    roots = content['response_documents'][0]['directory']
    roots[1].setdefault('children', []).append(roots[0]['children'].pop(0))
    save(current, content)
    snapshot = build_state.snapshot(tid)
    assert snapshot['orphan_body_files'] == []
    assert {'技术方案（附件N）', '需求理解'} <= {s['title'] for s in snapshot['body_sections'] if s['status'] != 'missing'}
    assert parent.read_bytes() == parent_bytes and child.read_bytes() == child_bytes


def test_rename_swap_preserves_each_chapter_and_rejects_existing_unassigned_file(env):
    tid, current = published(env)
    first, first_bytes = body(tid, '需求理解')
    second, second_bytes = body(tid, '总体设计')
    content = deepcopy(current['content'])
    nodes = content['response_documents'][0]['directory'][0]['children']
    nodes[0]['目录名称'], nodes[1]['目录名称'] = nodes[1]['目录名称'], nodes[0]['目录名称']
    save(current, content, [SectionRename(volume='技术部分', from_title=a, to_title=b) for a, b in [('需求理解', '总体设计'), ('总体设计', '需求理解')]])
    assert first.read_bytes() == second_bytes and second.read_bytes() == first_bytes
    current = page_flow.directory(tid)
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = '已有独立正文'
    orphan = env['out'] / 'body/已有独立正文.docx'
    orphan.write_bytes(first_bytes)
    with pytest.raises(BusinessError, match='未覆盖'):
        save(current, content, [SectionRename(volume='技术部分', from_title='总体设计', to_title='已有独立正文')])
    assert page_flow.directory(tid) == current
    assert first.read_bytes() == second_bytes and second.read_bytes() == first_bytes
    assert orphan.read_bytes() == first_bytes


def test_failed_save_rolls_back_directory_word_records_and_acceptance(env, monkeypatch):
    tid, current = published(env)
    old, original = body(tid, '需求理解')
    records = build_state.read(tid)
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = '新版需求理解'
    upsert = db.upsert_artifact_index
    def fail_once(row):
        if row['content_seq'] > current['content_seq']:
            raise OSError('模拟索引提交失败')
        return upsert(row)
    monkeypatch.setattr(db, 'upsert_artifact_index', fail_once)
    with pytest.raises(OSError, match='模拟索引提交失败'):
        save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='新版需求理解')])
    assert page_flow.directory(tid) == current
    assert build_state.read(tid) == records
    assert old.read_bytes() == original
    assert not (env['out'] / 'body/新版需求理解.docx').exists()


def test_later_ai_publication_supersedes_manual_save_marker(env):
    tid, current = published(env)
    save(current, current['content'])
    accepted = page_flow.directory(tid)
    assert accepted['user_input_version']
    row = db.get_artifact_index(current['artifact_id'])
    db.upsert_artifact_index({**row, 'content_seq': row['content_seq'] + 1})
    assert 'user_input_version' not in page_flow.directory(tid)


def test_multi_volume_same_title_rename_keeps_other_volume_and_single_to_multi_paths(env):
    tid, current = published(env)
    _, original = body(tid, '需求理解')
    content = deepcopy(current['content'])
    content['response_documents'].append({'name': '商务部分', 'directory': [{'目录名称': '需求理解', 'level': 1, 'children': []}]})
    save(current, content)
    assert (env['out'] / 'body/技术部分/需求理解.docx').read_bytes() == original
    business = env['out'] / 'body/商务部分/需求理解.docx'
    business.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); doc.add_paragraph('另一册正文不得串接或覆盖'); doc.save(business)
    unrelated = business.read_bytes()
    current = page_flow.directory(tid)
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = '技术需求分析'
    save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='技术需求分析')])
    assert (env['out'] / 'body/技术部分/技术需求分析.docx').read_bytes() == original
    assert business.read_bytes() == unrelated
    assert build_state.snapshot(tid)['orphan_body_files'] == []


def test_filename_collision_and_stale_rename_are_rejected_without_writes(env):
    tid, current = published(env)
    old, original = body(tid, '需求理解')
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'].extend([
        {'目录名称': 'A/B', 'level': 2}, {'目录名称': 'A:B', 'level': 2}])
    with pytest.raises(BusinessError, match='文件名重复'):
        save(current, content)
    with pytest.raises(BusinessError, match='改名依据已变化'):
        save(current, current['content'], [SectionRename(volume='技术部分', from_title='已不存在的旧标题', to_title='需求理解')], force=True)
    assert page_flow.directory(tid) == current
    assert old.read_bytes() == original


def test_failed_rollback_keeps_word_and_preservation_directory(env, monkeypatch):
    tid, current = published(env)
    _, original = body(tid, '需求理解')
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = '新版需求理解'
    rename, upsert = Path.rename, db.upsert_artifact_index
    def reject_recovery(path, target):
        if path.name == '新版需求理解.docx' and Path(target).parent.name.startswith('.directory-edit-'):
            raise PermissionError('模拟文件恢复被拒绝')
        return rename(path, target)
    def reject_commit(row):
        if row['content_seq'] > current['content_seq']:
            raise OSError('模拟索引提交失败')
        return upsert(row)
    monkeypatch.setattr(Path, 'rename', reject_recovery)
    monkeypatch.setattr(db, 'upsert_artifact_index', reject_commit)
    with pytest.raises(BusinessError, match='保全目录已保留'):
        save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='新版需求理解')])
    assert (env['out'] / 'body/新版需求理解.docx').read_bytes() == original
    assert list(env['out'].glob('.directory-edit-*'))


@pytest.mark.parametrize('titles', [('API方案', 'api方案'), ('Café方案', 'Cafe\u0301方案')])
def test_equivalent_destination_names_are_rejected_before_any_write(env, titles):
    tid, current = published(env)
    first, first_bytes = body(tid, '需求理解')
    second, second_bytes = body(tid, '总体设计')
    state = build_state.read(tid)
    content = deepcopy(current['content'])
    nodes = content['response_documents'][0]['directory'][0]['children']
    nodes[0]['目录名称'], nodes[1]['目录名称'] = titles
    renames = [SectionRename(volume='技术部分', from_title=old, to_title=new)
               for old, new in zip(('需求理解', '总体设计'), titles)]
    with pytest.raises(BusinessError, match='文件名重复'):
        save(current, content, renames)
    assert page_flow.directory(tid) == current
    assert build_state.read(tid) == state
    assert first.read_bytes() == first_bytes and second.read_bytes() == second_bytes
    assert not list(env['out'].glob('.directory-edit-*'))


def test_case_only_rename_preserves_body_and_records(env):
    tid, current = published(env)
    content = deepcopy(current['content'])
    node = content['response_documents'][0]['directory'][0]['children'][0]
    node['目录名称'] = 'API方案'
    save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='API方案')])
    _, original = body(tid, 'API方案')
    current = page_flow.directory(tid)
    state = build_state.read(tid)
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = 'api方案'
    save(current, content, [SectionRename(volume='技术部分', from_title='API方案', to_title='api方案')])
    assert (env['out'] / 'body/api方案.docx').read_bytes() == original
    assert build_state.read(tid)['sections']['body/api方案.docx'] == state['sections']['body/API方案.docx']
    assert 'body/API方案.docx' not in build_state.read(tid)['sections']
    assert build_state.snapshot(tid)['orphan_body_files'] == []


@pytest.mark.parametrize('missing_file', [False, True])
def test_equivalent_unassigned_destination_is_never_overwritten(env, missing_file):
    tid, current = published(env)
    source, original = body(tid, '需求理解')
    destination, unrelated = body(tid, 'API方案')
    if missing_file:
        state = build_state.read(tid)
        state['sections']['body/API方案.docx'] = deepcopy(state['sections']['body/需求理解.docx'])
        build_state.save(tid, state)
        destination.unlink()
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][0]['children'][0]['目录名称'] = 'api方案'
    with pytest.raises(BusinessError, match='未覆盖'):
        save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='api方案')])
    assert page_flow.directory(tid) == current
    assert source.read_bytes() == original
    if not missing_file:
        assert destination.read_bytes() == unrelated


def test_duplicate_structural_parents_allow_rename_move_and_later_edits(env):
    tid, current = published(env)
    first, first_bytes = body(tid, '需求理解')
    second, second_bytes = body(tid, '总体设计')
    content = deepcopy(current['content'])
    roots = content['response_documents'][0]['directory']
    first_node, second_node = roots[0]['children'][:2]
    del roots[0]['children'][:2]
    first_node['目录名称'] = '需求分析'
    roots.extend([{'目录名称': title, 'level': 1, 'children': [
        {'目录名称': '组织安排', 'level': 2, 'children': [{**node, 'level': 3}]}]}
        for title, node in [('开发服务', first_node), ('运维服务', second_node)]])
    save(current, content, [SectionRename(volume='技术部分', from_title='需求理解', to_title='需求分析')])
    current = page_flow.directory(tid)
    assert (env['out'] / 'body/需求分析.docx').read_bytes() == first_bytes
    assert second.read_bytes() == second_bytes
    snapshot = build_state.snapshot(tid)
    assert snapshot['orphan_body_files'] == []
    assert '组织安排' not in {s['title'] for s in snapshot['body_sections']}
    # 同名结构父节点无需正文身份，也应能单独改名；另一个同名父节点不受影响。
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'][-2]['children'][0]['目录名称'] = '开发组织'
    save(current, content, [SectionRename(volume='技术部分', from_title='组织安排', to_title='开发组织')])
    assert (env['out'] / 'body/需求分析.docx').read_bytes() == first_bytes
    assert second.read_bytes() == second_bytes
    assert build_state.snapshot(tid)['orphan_body_files'] == []


def test_parent_that_owns_body_cannot_share_a_filename(env):
    tid, current = published(env)
    parent, original = body(tid, '技术方案（附件N）')
    content = deepcopy(current['content'])
    content['response_documents'][0]['directory'].append({
        '目录名称': '技术方案（附件N）', 'level': 1,
        'children': [{'目录名称': '独立叶子', 'level': 2}]})
    with pytest.raises(BusinessError, match='文件名重复'):
        save(current, content)
    assert page_flow.directory(tid) == current
    assert parent.read_bytes() == original
