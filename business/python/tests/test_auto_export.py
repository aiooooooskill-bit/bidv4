"""Word唯一导出入口：原生更新域、源内容保全、失败不覆盖旧发布稿。"""
import asyncio
import json
import os
import subprocess
import sys
from pathlib import Path
import pytest
from docx import Document
from docx.oxml.ns import qn
from app import artifact_store,db
from app.services import artifacts,export_validation,build_state
from app.services.errors import BusinessError
from app.tools.docx_ops import docx_section_create,docx_assemble_volume
from tests.test_document_pagination import setup,seed,node


def export(tid):
    request={'operation':'export','data_root':os.environ['DATA_DIR'],'task_id':tid}
    environment=dict(os.environ);environment['PYTHONPATH']=os.pathsep.join(sys.path)
    result=subprocess.run([sys.executable,str(Path(__file__).parents[1]/'runner.py')],input=json.dumps(request),text=True,capture_output=True,env=environment,timeout=120)
    return json.loads(result.stdout)


def inputs(tmp_path,monkeypatch):
    tid,root=setup(tmp_path,monkeypatch)
    nodes=[node('封面'),node('目录'),node('技术方案',[{'目录名称':'管理方案','level':2,'children':[{'目录名称':'服务流程','level':3,'children':[]}]}]),node('附件证明')]
    seed(tid,nodes)
    for name in ['封面','服务流程','附件证明']:
        docx_section_create('body/'+name,name,paragraphs='原有正文与附件内容保持完整。')
    return tid,root,nodes


def package(tid):
    row=next(r for r in db.list_artifact_index(task_id=tid) if r['kind']=='tender.volume')
    meta=json.loads(artifact_store.read_content_resolved(row['artifact_id'],row))
    word=Path(asyncio.run(artifacts.get_artifact_file(row['artifact_id'])).path)
    return row,meta,word


def verify_downloads(tid):
    row,meta,word=package(tid)
    assert meta['field_update']=='on_open' and not any(k.startswith('pdf_') for k in meta)
    assert not artifact_store.package_files(row['artifact_id'],row,'*.pdf')
    d=Document(word)
    fs=export_validation.fields(d.element.body)
    assert sum(f['code'].strip().startswith('TOC ') for f in fs)==1
    names={b.get(qn('w:name')) for b in d.element.body.iter(qn('w:bookmarkStart'))}
    for f in fs:
        if f['code'].strip().startswith('PAGEREF '):
            assert f['code'].split()[1] in names
            assert not f['value'].strip() # 不能把估算页码当实际分页。
    assert d.settings.element.find(qn('w:updateFields')).get(qn('w:val'))=='true'
    assert not export_validation.fields(d.sections[0].footer._element)
    for section in d.sections[1:]:
        codes=[f['code'].strip() for f in export_validation.fields(section.footer._element)]
        assert codes[0]=='PAGE' and codes[1].startswith('PAGEREF _TenderCountedEnd')
        assert section.footer.paragraphs[0].alignment==1
    assert d.sections[1]._sectPr.find(qn('w:pgNumType')).get(qn('w:start'))=='1'
    with pytest.raises(BusinessError,match='仅支持Word'):
        asyncio.run(artifacts.get_artifact_file(row['artifact_id'],format='pdf'))
    return row,meta,word


def test_word_export_rebuilds_numbering_and_toc_after_growth_and_removal(tmp_path,monkeypatch):
    tid,root,nodes=inputs(tmp_path,monkeypatch)
    originals={p.name:p.read_bytes() for p in (root/'body').glob('*.docx')}
    assert 'result' in export(tid)
    row,_,_=verify_downloads(tid)
    assert 'result' in export(tid)
    assert package(tid)[0]['content_seq']==row['content_seq']
    docx_section_create('body/新增方案','新增方案',paragraphs='新增方案正文。')
    added=root/'body/新增方案.docx';d=Document(added)
    for _ in range(2):d.add_page_break();d.add_paragraph('实际增加的正文页。')
    d.save(added);seed(tid,nodes[:2]+[node('新增方案')]+nodes[2:])
    assert 'result' in export(tid)
    _,_,word=verify_downloads(tid)
    assert any(p.text.startswith('2 技术方案\t') for p in Document(word).paragraphs)
    assert sum(p.text=='实际增加的正文页。' for p in Document(word).paragraphs)==2
    added.rename(tmp_path/'新增正文原件保留.docx');seed(tid,nodes)
    assert 'result' in export(tid)
    _,_,word=verify_downloads(tid)
    assert any(p.text.startswith('1 技术方案\t') for p in Document(word).paragraphs)
    assert all((root/'body'/name).read_bytes()==blob for name,blob in originals.items())


def test_broken_page_target_rejected_without_replacing_published_word(tmp_path,monkeypatch):
    tid,root,_=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    before,_,word=package(tid);previous=word.read_bytes()
    original=export_validation.verify
    def corrupt(path,rows,has_cover,selected):
        d=Document(path)
        target=next(t for t in d.element.body.iter(qn('w:instrText')) if 'PAGEREF _TenderHeading' in (t.text or ''))
        target.text=' PAGEREF _TenderHeading999999 '
        d.save(path);return original(path,rows,has_cover,selected)
    monkeypatch.setattr(export_validation,'verify',corrupt)
    result=docx_assemble_volume();assert '未知或重复的章节引用' in result,result
    after,_,word=package(tid)
    assert after['content_seq']==before['content_seq'] and word.read_bytes()==previous


def test_word_publication_failure_preserves_previous_complete_version(tmp_path,monkeypatch):
    tid,root,_=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    before,_,word=package(tid);previous=word.read_bytes()
    source=root/'body/服务流程.docx';d=Document(source);d.add_paragraph('这是用户新补充的实际正文。');d.save(source)
    def fail_word(*a,**k):raise OSError('test: Word写入失败')
    monkeypatch.setattr(artifact_store,'write_package_file',fail_word)
    result=docx_assemble_volume();assert 'Word写入失败' in result,result
    after,_,word=package(tid)
    assert after['content_seq']==before['content_seq'] and word.read_bytes()==previous


def test_download_rejects_word_modified_after_validation(tmp_path,monkeypatch):
    tid,_,_=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    row,_,word=package(tid);word.write_bytes(word.read_bytes()+b'test tampering')
    with pytest.raises(BusinessError,match='重新合册'):
        asyncio.run(artifacts.get_artifact_file(row['artifact_id']))
    assert build_state.snapshot(tid)['volume_states'][0]['status']!='ready'


def test_unresolved_material_references_block_export_and_preserve_previous_word(tmp_path,monkeypatch):
    tid,root,_=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    before,_,word=package(tid);previous=word.read_bytes()
    source=root/'body/附件证明.docx';d=Document(source);table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='人员';table.cell(0,1).text='材料所在页码'
    table.cell(1,0).text='测试人员';table.cell(1,1).text='50-51';d.save(source);original=source.read_bytes()
    result=export(tid)
    assert 'error' in result and '1处资料页码未能建立自动引用' in result['error'],result
    after,_,word=package(tid)
    assert after['content_seq']==before['content_seq'] and word.read_bytes()==previous and source.read_bytes()==original
    assert build_state.snapshot(tid)['volume_states'][0]['status']!='ready'
    from app.tools.docx_ops import _publish_volume
    pagination=json.loads((root/'body/整本-格式验证册.pagination.json').read_text())
    pagination['unresolved_page_references']=[{'reason':'缺少实际附件'}]
    assert '不能发布整册' in _publish_volume(tid,root/'body/整本-格式验证册.docx','格式验证册',1,0,0,pagination)
    assert package(tid)[0]['content_seq']==before['content_seq']


@pytest.mark.parametrize('problem',['missing','pending'])
def test_export_rejects_known_incomplete_sources_before_rewriting_word(tmp_path,monkeypatch,problem):
    tid,root,nodes=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    before,_,word=package(tid);published=word.read_bytes()
    working=root/'body/整本-格式验证册.docx';original=working.read_bytes()
    if problem=='missing':seed(tid,nodes+[node('未提供的实际附件')])
    else:
        state=build_state.read(tid);state.setdefault('pending_sections',{})['body/服务流程.docx']={'basis':None,'hash':'pending'};build_state.save(tid,state)
    result=export(tid);assert result['status']==422 and '整册未完成生成或发布' in result['error'],result
    assert working.read_bytes()==original
    after,_,word=package(tid);assert after['content_seq']==before['content_seq'] and word.read_bytes()==published


def test_reexport_repairs_damaged_previous_download_without_modifying_sources(tmp_path,monkeypatch):
    tid,root,_=inputs(tmp_path,monkeypatch);assert 'result' in export(tid)
    previous,_,word=package(tid)
    source_bytes={p:p.read_bytes() for p in (root/'body').glob('*.docx') if not p.name.startswith('整本-')}
    word.write_bytes(word.read_bytes()+b'tampered archive trailer')
    assert 'result' in export(tid)
    current,_,_=verify_downloads(tid)
    assert current['content_seq']==previous['content_seq']+1
    assert all(p.read_bytes()==value for p,value in source_bytes.items())
