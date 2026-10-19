import asyncio,json
import pytest
from app import db,company_store as cs
from app.knowledge import ingest,store,fts
from app.services import knowledge,company_library as lib,materials
from app.services.errors import BusinessError

def setup(tmp_path,monkeypatch):
    monkeypatch.setenv('DATA_DIR',str(tmp_path));db.init_db();store.ensure_dirs()
    c=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'phone':'123'})))
    i=db.kb_insert_item(file_name='许可证.png',file_hash='real',title='开户许可证',ext='.png')
    md,outline,_,_=store.kb_parse_paths(i['file_name']);md.parent.mkdir(parents=True);md.write_text('历史银行旧账号');outline.write_text('[]')
    metadata={'doc_type':'qualification_certificate','extra':{'银行账号':{'value':'旧账号','source':'第1页'}},'statement':'开户银行许可证由图片识别，包含企业名和开户银行信息，请核对有效性。'}
    ingest.finish(i['id'],metadata)
    return c,db.kb_get_item(i['id']),metadata

def test_manual_review_survives_reingest_and_same_second_writes(tmp_path,monkeypatch):
    c,i,m=setup(tmp_path,monkeypatch);state=cs.asset_state(i['id']);m['extra']['银行账号']['value']='人工正确账号'
    body=lib.ReviewBody(revision=state['revision'],company_id=c['id'],title=i['title'],business=m,text='经核对的最新银行账号')
    asyncio.run(lib.review_asset(i['id'],body))
    with pytest.raises(BusinessError,match='其他操作更新'):asyncio.run(lib.review_asset(i['id'],body))
    ingest.finish(i['id'],{'doc_type':'other','extra':{'银行账号':{'value':'新识别错误值'}},'statement':'重新识别的机器数据，与人工核对结果可能存在差异，不作为当前公司事实。'})
    current=asyncio.run(knowledge.get_item(i['id']))
    assert current['business']['extra']['银行账号']['value']=='人工正确账号'
    assert current['suggested']['extra']['银行账号']['value']=='新识别错误值'
    assert current['review_status']=='confirmed' and current['company_id']==c['id']
    content=asyncio.run(knowledge.get_item_content(i['id']))
    assert content['content']=='经核对的最新银行账号' and content['original_content']=='历史银行旧账号'
    hits=db.kb_search_segments(fts.build_match_expr('人工正确账号'),domain='kb',company_id=c['id'])
    assert any(h['item_id']==i['id'] for h in hits)

def test_scope_precedes_limit_and_unassigned_is_explicit(tmp_path,monkeypatch):
    a,item,m=setup(tmp_path,monkeypatch)
    b=asyncio.run(lib.save_company(lib.CompanyBody(name='乙公司')))
    # More high ranking foreign-company hits than the tool's limit.
    for n in range(40):
        k=db.kb_insert_item(file_name=f'其他{n}.png',file_hash=str(n),title='资质',ext='.png')
        cs.attach_asset(k['id'],b['id']);db.kb_replace_segments(k['id'],[{'body':fts.segment_for_fts('银行账号')}])
    cs.attach_asset(item['id'],a['id'])
    hits=db.kb_search_segments(fts.build_match_expr('银行账号'),limit=1,domain='kb',company_id=a['id'])
    assert len(hits)==1 and hits[0]['item_id']==item['id']
    assert asyncio.run(knowledge.list_items(company_id=a['id']))['total']==1
    unknown=db.kb_insert_item(file_name='未知.png',file_hash='unknown',title='未知',ext='.png')
    assert asyncio.run(knowledge.list_items(company_id='unassigned'))['items'][0]['id']==unknown['id']
    assert cs.asset_state(unknown['id'])['company_id'] is None

def test_entity_revision_and_independent_representative(tmp_path,monkeypatch):
    c,i,m=setup(tmp_path,monkeypatch)
    r=asyncio.run(lib.save_representative(lib.RepresentativeBody(name='张三')))
    assert r['revision']==1
    update=lib.RepresentativeBody(**{**r,'phone':'888'})
    asyncio.run(lib.save_representative(update))
    with pytest.raises(BusinessError):asyncio.run(lib.save_representative(update))
    assert 'company_id' not in r
    assert asyncio.run(lib.overview())['representatives'][0]['id']==r['id']
    db.kb_update_item(i['id'],extract_status='running')
    with pytest.raises(BusinessError,match='正在识别'):asyncio.run(lib.review_asset(i['id'],lib.ReviewBody(revision=cs.asset_state(i['id'])['revision'],company_id=c['id'],title=i['title'],business=m)))

def test_materials_are_global_with_source_and_pagination(tmp_path,monkeypatch):
    setup(tmp_path,monkeypatch)
    f=db.mt_insert_file(file_name='技术方案.docx',file_hash='mt')
    db.mt_replace_blocks(f['id'],[{'id':f'mb_{n}','title':f'质量方案{n}','note':'数据治理','ranges':json.dumps([[1,2]]),'chars':100} for n in range(12)])
    page=asyncio.run(materials.list_blocks(file_id=f['id'],page=2,page_size=5))
    assert page['total']==12 and len(page['blocks'])==5 and page['pages']==3
    assert 'company_id' not in page['blocks'][0]

def test_project_snapshot_scope_source_versions_and_idle(tmp_path,monkeypatch):
    from app.services import bidder,page_flow
    from app.tools.company_asset import read_company_asset
    from app.tools.search_knowledge import search_company_assets,search_references
    from app import runctx
    c,i,m=setup(tmp_path,monkeypatch)
    reviewed=lib.ReviewBody(revision=cs.asset_state(i['id'])['revision'],company_id=c['id'],title=i['title'],business=m)
    asyncio.run(lib.review_asset(i['id'],reviewed))
    c=asyncio.run(lib.save_company(lib.CompanyBody(**{**c,'profile':{'legal_name':'当前法人','source-id:bank_account':i['id'],'source-revision:bank_account':str(cs.asset_state(i['id'])['revision']),'bank_account':'核对账号'}})))
    tid=db.create_task('真实生成验证')['id']
    old=page_flow.input_versions(tid)
    binding=bidder.BidderBody(company_id=c['id'],signer_id='legal',sign_date='2026-10-12')
    saved=asyncio.run(bidder.save_bidder(tid,binding))
    assert saved['snapshot']['company']['profile']['bank_account']=='核对账号'
    assert page_flow.input_versions(tid)['analyse']==old['analyse']
    assert page_flow.input_versions(tid)['guide']!=old['guide']
    runctx.set_run('test','test',tid)
    assert '许可证.png' in search_company_assets('开户银行')
    assert json.loads(read_company_asset(i['id']))['business']==m
    # Changes to a document which was adopted to profile require reconfirmation.
    ingest.finish(i['id'],m)
    assert bidder.status(tid)['stale']
    with pytest.raises(BusinessError,match='重新核对'):bidder.require_confirmed(tid)
    c=asyncio.run(lib.save_company(lib.CompanyBody(**{**c,'profile':{**c['profile'],'source-revision:bank_account':str(cs.asset_state(i['id'])['revision'])}})))
    binding.revision=saved['revision'];asyncio.run(bidder.save_bidder(tid,binding))
    b=asyncio.run(lib.save_company(lib.CompanyBody(name='其他公司',profile={'legal_name':'乙法人'})))
    rep=asyncio.run(lib.save_representative(lib.RepresentativeBody(name='通用代表')))
    shared=asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=c['id'],signer_id=rep['id'],sign_date='2026-10-12',revision=binding.revision+1)))
    assert shared['snapshot']['signer']['id']==rep['id']
    assert 'company_id' not in shared['snapshot']['signer']
    binding.company_id=b['id'];binding.signer_id=rep['id'];binding.revision=shared['revision']
    asyncio.run(bidder.save_bidder(tid,binding))
    assert '许可证.png' not in search_company_assets('开户银行')
    assert '不属于' in read_company_asset(i['id'])
    assert '写作素材无命中' in search_references('暂不存在的素材')
    path=store.kb_files_dir()/i['file_name']
    with pytest.raises(BusinessError,match='不属于'):cs.require_asset_path(tid,path)
    from app.config import data_dir
    (data_dir()/'locks').mkdir();(data_dir()/'locks'/tid).write_text('123')
    with pytest.raises(BusinessError,match='正在执行'):asyncio.run(bidder.save_bidder(tid,binding))
    runctx.clear_run()

def test_obsolete_index_writer_cannot_replace_new_review(tmp_path,monkeypatch):
    c,i,m=setup(tmp_path,monkeypatch)
    old=cs.asset_state(i['id'])['revision']
    asyncio.run(lib.review_asset(i['id'],lib.ReviewBody(revision=old,company_id=c['id'],title=i['title'],business=m,text='正确核对全文')))
    db.kb_replace_segments(i['id'],[{'body':fts.segment_for_fts('过时全文')}],expected_revision=old)
    assert not db.kb_search_segments(fts.build_match_expr('过时'),domain='kb')
    assert db.kb_search_segments(fts.build_match_expr('正确核对全文'),domain='kb')

def test_confirmed_identity_fills_actual_tender_letter_docx(tmp_path,monkeypatch):
    from docx import Document
    from app import artifact_store,runctx,document
    from app.services import bidder
    from app.tools.docx_ops import docx_section_create,docx_source_inject,docx_section_revise,docx_section_read
    c,i,m=setup(tmp_path,monkeypatch)
    c=asyncio.run(lib.save_company(lib.CompanyBody(**{**c,'profile':{'legal_name':'核对法人','legal_title':'总经理','address':'核对通讯地址','phone':'123456','fax':'无','email':'bid@example.test','bank_name':'核对开户银行','bank_account':'0001'}})))
    tid=db.create_task('投标函格式验证')['id'];artifact_store.ensure_task_skeleton(tid);runctx.set_run('test','test',tid)
    state=asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=c['id'],signer_id='legal',sign_date='2026-10-12')))
    source=artifact_store.sources_dir(tid)/'投标函格式.docx';source.parent.mkdir(parents=True,exist_ok=True)
    template=Document();table=template.add_table(rows=1,cols=2);table.cell(0,0).text='事项';table.cell(0,1).text='投标人填写'
    identity={'供应商':c['name'],'签字代表':state['snapshot']['signer']['name'],'职务':c['profile']['legal_title'],'地址':c['profile']['address'],'电话':c['profile']['phone'],'传真':c['profile']['fax'],'邮箱':c['profile']['email'],'签署日期':state['snapshot']['sign_date']}
    for label in identity:
        cells=table.add_row().cells;cells[0].text=label;cells[1].text=''
    template.save(source);document.prepare(source,artifact_store.work_dir(tid)/'parse'/source.name)
    target='body/投标函.docx';assert docx_section_create(path=target,title='投标函').startswith('[已创建]')
    assert docx_source_inject(source=source.name,dest=target).startswith('[已注入]')
    edits=[{'table':1,'row':n,'col':2,'action':'fill','text':value} for n,value in enumerate(identity.values(),2)]
    result=docx_section_revise(path=target,edits=json.dumps(edits,ensure_ascii=False));assert result.startswith('[已修订]'),result
    generated=Document(artifact_store.work_dir(tid)/target)
    accepted=docx_section_read(path=target)
    assert all(value in accepted for value in identity.values())
    assert len(generated.tables)==1 and len(generated.tables[0].rows)==len(identity)+1
    runctx.clear_run()


def test_representative_is_independent_and_shared_only_through_project_choices(tmp_path,monkeypatch):
    from app.services import bidder
    from pydantic import ValidationError
    monkeypatch.setenv('DATA_DIR',str(tmp_path));db.init_db()
    rep=asyncio.run(lib.save_representative(lib.RepresentativeBody(name='共用授权代表',phone='123')))
    assert cs.companies()==[]
    assert 'company_id' not in rep
    with pytest.raises(ValidationError):lib.RepresentativeBody(name='旧归属字段',company_id='obsolete')
    companies=[asyncio.run(lib.save_company(lib.CompanyBody(name=name))) for name in ['甲公司','乙公司']]
    tasks=[db.create_task('独立授权选择')['id'] for _ in companies]
    for tid,company in zip(tasks,companies):
        saved=asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=company['id'],signer_id=rep['id'],sign_date='2026-10-15')))
        assert saved['snapshot']['signer']['id']==rep['id'] and not saved['stale']
        with pytest.raises(BusinessError,match='不存在'):
            asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=company['id'],signer_id='missing',sign_date='2026-10-15',revision=saved['revision'])))
    asyncio.run(lib.save_representative(lib.RepresentativeBody(**{**rep,'phone':'456'})))
    assert all(bidder.status(tid)['stale'] for tid in tasks)


def test_material_inventory_unifies_pending_files_and_blocks_without_duplicates(tmp_path,monkeypatch):
    setup(tmp_path,monkeypatch)
    pending=db.mt_insert_file(file_name='售后服务.docx',file_hash='pending')
    ready=db.mt_insert_file(file_name='实施方案.docx',file_hash='ready')
    db.mt_update_file(pending['id'],parse_status='ready')
    db.mt_update_file(ready['id'],parse_status='ready')
    db.mt_replace_blocks(ready['id'],[{'id':f'mb_{n}','title':f'实施素材{n}','note':'交付','ranges':json.dumps([[1,2]]),'chars':100} for n in range(6)])
    page=asyncio.run(materials.list_items(page=1,page_size=5))
    assert page['total']==7 and len(page['items'])==5 and page['pages']==2
    assert page['items'][0]['kind']=='file' and page['items'][0]['id']==pending['id']
    assert all(i['kind']=='block' for i in page['items'][1:])
    tail=asyncio.run(materials.list_items(page=2,page_size=5))
    assert len(tail['items'])==2 and all(i['parse_status']=='ready' for i in tail['items'])
    assert asyncio.run(materials.list_items(q='售后',page_size=5))['total']==1
    assert asyncio.run(materials.list_items(file_id=ready['id'],page_size=5))['total']==6
    db.mt_update_file(ready['id'],parse_status='failed',error='目录解析失败')
    rows=asyncio.run(materials.list_items(file_id=ready['id'],page_size=5))['items']
    assert all(i['source_error']=='目录解析失败' for i in rows)
    db.mt_replace_blocks(pending['id'],[{'id':'mb_new','title':'售后方案','note':'','ranges':json.dumps([[1,2]]),'chars':20}])
    fresh=asyncio.run(materials.list_items(file_id=pending['id'],page_size=5))
    assert fresh['total']==1 and fresh['items'][0]['kind']=='block'
