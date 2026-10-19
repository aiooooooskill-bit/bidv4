"""Word原生域验收：原生域、封面排除、目录页码、招标格式覆盖及原件保全。"""
import hashlib,json
from pathlib import Path
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from app import runctx,publish,artifact_store
from app.tools.docx_ops import docx_section_create,docx_assemble_volume
from app.tools.document_format import docx_format_rule
from app.services import document_format
from tests.util import init_env


def setup(tmp_path,monkeypatch):
    task,conv=init_env(tmp_path,monkeypatch)
    tid=task['id'];artifact_store.ensure_task_skeleton(tid)
    runctx.set_run(conv['id'],'pagination-test',tid)
    return tid,artifact_store.work_dir(tid)


def seed(tid,nodes):
    publish.publish_artifact('tender.directory/tender-response-docs@1',{'numbering':'decimal','response_documents':[{'name':'格式验证册','directory':nodes}]},task_id=tid)


def node(name,children=None):return {'目录名称':name,'level':1,'children':children or []}


def test_real_toc_footer_default_and_source_conservation(tmp_path,monkeypatch):
    tid,root=setup(tmp_path,monkeypatch)
    seed(tid,[node('封面'),node('目录'),node('技术方案',[{'目录名称':'服务流程','level':2,'children':[]}]),node('附件证明')])
    for name in ['封面','服务流程','附件证明']:
        assert docx_section_create('body/'+name,name,paragraphs='原正文保持完整。\n附件及合同数据12345。').startswith('[已创建]')
    body=root/'body/服务流程.docx';src=Document(body);src.paragraphs[1].runs[0].font.name='Arial';src.paragraphs[1].runs[0].font.size=Pt(20);src.paragraphs[1].paragraph_format.line_spacing=1.0;src.save(body)
    before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (root/'body').glob('*.docx')}
    result=docx_assemble_volume();assert '已发布为整本' in result,result
    out=root/'body/整本-格式验证册.docx';d=Document(out);layout=json.loads(out.with_suffix('.pagination.json').read_text())
    assert len(d.sections)==2 and layout['field_update']=='on_open'
    assert layout['pagination_calculated'] is False
    assert layout['headings'][0]['bookmark']
    assert not d.sections[0].footer.paragraphs[0].text
    assert d.sections[1].footer.paragraphs[0].alignment==1
    assert 'PAGEREF _TenderCountedEnd' in d.sections[1].footer._element.xml
    assert 'TOC ' in d.element.body.xml and 'PAGEREF _TenderHeading' in d.element.body.xml
    prose=[p for p in d.paragraphs if p.text=='原正文保持完整。'];assert len(prose)==3
    for p in prose[1:]:
        assert p.paragraph_format.line_spacing==1.5
        assert p.runs[0].font.size==Pt(12)
        assert p.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia'))=='宋体'
    assert all(hashlib.sha256((root/'body'/name).read_bytes()).hexdigest()==h for name,h in before.items())


def test_tender_override_is_evidence_bound_and_stale_rejected(tmp_path,monkeypatch):
    tid,root=setup(tmp_path,monkeypatch);p=root/'parse/source/source.docx.md';p.parent.mkdir(parents=True);quote='正文采用仿宋，三号字，固定28磅行距。一级标题采用黑体，二号字。';p.write_text(quote)
    docx_format_rule('body',font='仿宋',font_size=16,line_spacing=28,spacing_rule='exact',evidence_path='parse/source/source.docx.md',evidence_quote=quote)
    docx_format_rule('heading1',font='黑体',font_size=22,evidence_path='parse/source/source.docx.md',evidence_quote=quote)
    seed(tid,[node('封面'),node('目录'),node('技术方案')])
    for name in ['封面','技术方案']:docx_section_create('body/'+name,name,paragraphs='招标明确格式优先。')
    result=docx_assemble_volume();assert '已发布为整本' in result,result
    d=Document(root/'body/整本-格式验证册.docx');p1=[p for p in d.paragraphs if p.text=='招标明确格式优先。' and p.paragraph_format.line_spacing==Pt(28)][-1]
    assert p1.runs[0].font.size==Pt(16)
    assert p1.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia'))=='仿宋'
    assert d.styles['Heading 1'].font.size==Pt(22)
    p.write_text('招标要求已经修改')
    with pytest.raises(ValueError,match='依据已变化'):document_format.rules(tid)


def test_export_without_any_external_application_or_pdf(tmp_path,monkeypatch):
    import subprocess
    tid,root=setup(tmp_path,monkeypatch);seed(tid,[node('目录'),node('方案')])
    docx_section_create('body/方案','方案',paragraphs='正文不可丢失')
    monkeypatch.setattr(subprocess,'run',lambda *a,**k:(_ for _ in ()).throw(AssertionError('不得调用外部程序')))
    result=docx_assemble_volume();assert '已发布为整本' in result,result
    assert not list(root.rglob('*.pdf'))


def test_multi_page_cover_multi_page_toc_and_footer_override(tmp_path,monkeypatch):
    tid,root=setup(tmp_path,monkeypatch)
    names=['第'+str(i)+'项接口对接及异常处置验证内容' for i in range(1,39)]
    seed(tid,[node('封面'),node('目录'),node('实施方案',[{'目录名称':name,'level':2,'children':[]} for name in names])])
    docx_section_create('body/封面','封面',paragraphs='第一张封面。')
    cover=root/'body/封面.docx';d=Document(cover);d.add_page_break();d.add_paragraph('第二张封面原内容。');d.save(cover)
    for name in names:docx_section_create('body/'+name,name,paragraphs='本项原正文完整保留。')
    src=root/'parse/source/requirements.md';src.parent.mkdir(parents=True);quote='页脚采用宋体五号字，右对齐，格式为第X页/共Y页。';src.write_text(quote)
    docx_format_rule('footer',font='宋体',font_size=10.5,alignment='right',footer_text='第{PAGE}页/共{TOTAL}页',evidence_path='parse/source/requirements.md',evidence_quote=quote)
    result=docx_assemble_volume();assert '已发布为整本' in result,result
    out=root/'body/整本-格式验证册.docx';layout=json.loads(out.with_suffix('.pagination.json').read_text());d=Document(out)
    assert len(d.sections)==2 and layout['pagination_calculated'] is False
    assert len(layout['headings'])==len(names)+1
    assert len(d.sections[0]._sectPr.xpath('./w:pgNumType'))==0
    assert d.sections[1].footer.paragraphs[0].alignment==2
    assert 'PAGEREF _TenderCountedEnd' in d.sections[1].footer._element.xml
    assert sum(p.text=='本项原正文完整保留。' for p in d.paragraphs)==len(names)


def test_native_heading_numbering_survives_insertion_without_rewriting_titles(tmp_path):
    """插入、删除上级时全部下级编号同步，包括素材小标题；标题原文不重写。"""
    from copy import deepcopy
    import re
    from app.services.heading_numbering import HeadingNumbering
    d=Document();document_format.apply(d,{})
    numbering=HeadingNumbering(d,'decimal')
    for level,title in [(1,'投标函'),(1,'技术方案'),(2,'项目管理'),(3,'服务流程')]:
        p=d.add_heading(title,level);numbering.heading(p._p,level)
    internal=d.add_heading('质量内标题',4)
    numbering.subheads([internal._p],3,re.compile(r'^已经编号'))
    business=d.add_heading('业务方案',2);numbering.heading(business._p,2)
    p=d.add_heading('技术数据',3);numbering.heading(p._p,3)
    p=d.add_heading('附件',1);numbering.heading(p._p,1)
    path=tmp_path/'native.docx';d.save(path)
    inserted=deepcopy(d.paragraphs[0]._p)
    for t in inserted.iter(qn('w:t')):t.text='新增说明'
    d.paragraphs[0]._p.addprevious(inserted)
    inserted_child=deepcopy(business._p)
    for t in inserted_child.iter(qn('w:t')):t.text='新增二级方案'
    business._p.addprevious(inserted_child)
    assert internal._p.pPr.numPr.numId.val==d.styles['Heading 1'].element.pPr.numPr.numId.val
    path=tmp_path/'native.docx';d.save(path)
    numbering_xml=d.part.numbering_part.element.xml
    assert '%1.%2.%3' in numbering_xml
    assert [p.text for p in d.paragraphs]==['新增说明','投标函','技术方案','项目管理','服务流程','质量内标题','新增二级方案','业务方案','技术数据','附件']
    assert all(p._p.pPr.numPr.numId.val==numbering.num_id for p in d.paragraphs)
    assert all(not re.match(r'^(?:第.+章|\d+\s)',p.text) for p in d.paragraphs)
    assert d.styles['Heading 1'].element.pPr.numPr.numId.val>0

    inserted.getparent().remove(inserted)
    inserted_child.getparent().remove(inserted_child)
    d.save(path)
    assert [p.text for p in Document(path).paragraphs]==['投标函','技术方案','项目管理','服务流程','质量内标题','业务方案','技术数据','附件']


def test_page_fields_and_source_breaks_survive_growth_and_shrinkage(tmp_path,monkeypatch):
    """三次合册保留原生页数引用和原始分页符，目录、正文和附件在内，封面排除且节文件不被改写。"""
    tid,root=setup(tmp_path,monkeypatch)
    nodes=[node('封面'),node('目录'),node('技术方案'),node('附件证明')]
    for name in ['封面','技术方案','附件证明']:
        docx_section_create('body/'+name,name,paragraphs='原有正文及附件资料完整保留。')
    original={p.name:p.read_bytes() for p in (root/'body').glob('*.docx')}
    def assemble(tree):
        seed(tid,tree)
        before={p.name:p.read_bytes() for p in (root/'body').glob('*.docx') if not p.name.startswith('整本-')}
        result=docx_assemble_volume();assert '已发布为整本' in result or '整本内容未变' in result,result
        path=root/'body/整本-格式验证册.docx';d=Document(path)
        layout=json.loads(path.with_suffix('.pagination.json').read_text())
        assert layout['pagination_calculated'] is False
        assert d.sections[-1].footer.paragraphs[0].alignment==1
        codes=[x.text.strip() for x in d.sections[-1].footer._element.iter(qn('w:instrText'))]
        assert codes[0]=='PAGE' and codes[1].startswith('PAGEREF _TenderCountedEnd')
        assert 'SECTIONPAGES' not in d.sections[-1].footer._element.xml
        total=sum(len(p._p.xpath('.//w:br[@w:type="page"]')) for p in d.paragraphs)
        assert all((root/'body'/name).read_bytes()==blob for name,blob in before.items())
        return total
    first=assemble(nodes)
    docx_section_create('body/新增章节','新增章节',paragraphs='验证新增原文。')
    added=root/'body/新增章节.docx';d=Document(added)
    for _ in range(2):d.add_page_break();d.add_paragraph('本页是测试添加的正文。')
    d.save(added)
    grown=assemble(nodes[:2]+[node('新增章节')]+nodes[2:])
    assert grown==first+2 # 原稿明确增加的两个分页符，章节边界使用pageBreakBefore。
    added.rename(tmp_path/'仅测试新增章节保留.docx')
    final=assemble(nodes)
    assert final==first
    assert all((root/'body'/name).read_bytes()==blob for name,blob in original.items())


def test_small_chinese_font_size_cannot_approve_larger_size(tmp_path,monkeypatch):
    tid,root=setup(tmp_path,monkeypatch)
    evidence=root/'parse/tender/format.md';evidence.parent.mkdir(parents=True);quote='正文采用宋体，小四号字，小三号标题。';evidence.write_text(quote)
    for size in [14,16]:
        with pytest.raises(ValueError,match='字号须与引用原文对应'):
            docx_format_rule('body',font_size=size,evidence_path='parse/tender/format.md',evidence_quote=quote)
    docx_format_rule('body',font_size=12,evidence_path='parse/tender/format.md',evidence_quote=quote)
    assert document_format.rules(tid)['body']['font_size']==12


def test_multiple_sections_share_continuous_page_numbering_and_end_target(tmp_path):
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from app.services.document_pagination import configure_sections,insert_toc,finalize
    from app.services.heading_numbering import HeadingNumbering
    d=Document();document_format.apply(d,{})
    d.settings.element.append(OxmlElement('w:evenAndOddHeaders'))
    d.add_paragraph('封面原文');d.add_page_break();d.add_paragraph('第二张封面原文')
    anchor=d.paragraphs[-1]._p
    number=HeadingNumbering(d,'decimal');h=d.add_heading('方案',1);prefix=number.heading(h._p,1)
    d.add_paragraph('方案正文');d.add_section(WD_SECTION.NEW_PAGE)
    t=d.add_table(rows=1,cols=1);t.cell(0,0).text='横向附件的实际内容'
    toc,rows=insert_toc(d,anchor,[(prefix,'方案',1)])
    configure_sections(d,toc,True,{})
    record=finalize(d,tmp_path/'sections.docx',rows,True,{})
    assert record['field_update']=='on_open' and len(d.sections)==3
    assert d.sections[0].footer.paragraphs[0].text==''
    assert d.sections[1]._sectPr.find(qn('w:pgNumType')).get(qn('w:start'))=='1'
    assert d.sections[2]._sectPr.find(qn('w:pgNumType')).get(qn('w:start')) is None
    assert d.settings.element.findall(qn('w:evenAndOddHeaders')) # 保留自定义奇偶页页眉设置。
    assert all('PAGEREF _TenderCountedEnd' in s.even_page_footer._element.xml for s in d.sections[1:])
    assert not d.sections[0].even_page_footer.paragraphs[0].text
    target=next(b for b in t.cell(0,0).paragraphs[0]._p.iter(qn('w:bookmarkStart')) if b.get(qn('w:name')).startswith('_TenderCountedEnd'))
    assert all(target.get(qn('w:name')) in s.footer._element.xml for s in d.sections[1:])
