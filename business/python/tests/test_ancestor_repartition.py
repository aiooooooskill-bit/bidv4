import hashlib,json
from collections import Counter
from io import BytesIO
import pytest
from docx import Document
from lxml import etree
from app import artifact_store
from app.services import build_state
from app.tools.body_repartition import docx_section_repartition, _blocks
from tests.test_body_repartition import env, assignments
from tests.test_docx_ops import _tiny_png


def child_source(env):
 tid,root,_,parent=env
 build_state.record_section(tid,'body/实施方案.docx');parent.unlink()
 path=root/'body/接口对接.docx';doc=Document();doc.add_heading('接口对接',1);doc.add_paragraph('原父声明及法律承诺完整保留。');doc.add_paragraph('原子章节附件说明。');doc.add_table(rows=1,cols=1).cell(0,0).text='原表格'
 image=root/'original.png';_tiny_png(image);doc.add_picture(str(image));doc.save(path)
 return path


def test_read_view_offers_missing_owned_ancestor_without_opening_siblings(env):
 source=child_source(env);view=json.loads(docx_section_repartition('body/接口对接.docx'))
 assert '实施方案' in view['allowed_titles'];assert '故障处置' not in view['allowed_titles']


def test_move_to_ancestor_conserves_all_ooxml_blocks_and_actual_original_backup(env):
 source=child_source(env);before=source.read_bytes();view=json.loads(docx_section_repartition('body/接口对接.docx'))
 route=[{'title':'实施方案','ranges':[[2,2]]},{'title':'接口对接','ranges':[[1,1],[3,5]]}]
 result=json.loads(docx_section_repartition('body/接口对接.docx',json.dumps(route),view['sha256']))
 assert Counter(etree.tostring(e) for e in _blocks(Document(BytesIO(before))))==Counter(etree.tostring(e) for item in result['outputs'] for e in _blocks(Document(env[1]/item['path'])))
 assert (env[1]/result['restore_point']).read_bytes()==before
 assert len(Document(source).inline_shapes)==1 and len(Document(source).tables)==1
 assert '原父声明' in Document(env[1]/'body/实施方案.docx').paragraphs[0].text


def test_ancestor_permission_does_not_allow_sibling_or_existing_ancestor_overwrite(env):
 source=child_source(env);before=source.read_bytes()
 for title in ['故障处置','其他章节']:
  with pytest.raises(ValueError):docx_section_repartition('body/接口对接.docx',json.dumps([{'title':title,'ranges':[[1,5]]}]),hashlib.sha256(before).hexdigest())
 assert source.read_bytes()==before
 parent=env[1]/'body/实施方案.docx';parent.write_bytes(before)
 with pytest.raises(ValueError,match='已有正文'):docx_section_repartition('body/接口对接.docx',json.dumps([{'title':'实施方案','ranges':[[1,5]]}]),hashlib.sha256(before).hexdigest())
 assert source.read_bytes()==before and parent.read_bytes()==before


def test_all_partition_outputs_inherit_actual_source_dependencies(env):
 tid,root,_,source=env;original=source.read_bytes();reference=artifact_store.sources_dir(tid)/'实际依据.docx';reference.write_bytes(b'original-reference')
 dependency={reference.name:hashlib.sha256(reference.read_bytes()).hexdigest()};value=build_state.read(tid);value.setdefault('source_dependencies',{})['body/实施方案.docx']=dependency;build_state.save(tid,value)
 out=json.loads(docx_section_repartition('body/实施方案.docx',json.dumps(assignments()),hashlib.sha256(original).hexdigest()));value=build_state.read(tid)
 assert all(value['source_dependencies'].get(x['path'])==dependency for x in out['outputs'])
