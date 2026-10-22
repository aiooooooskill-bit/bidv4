import json
from copy import deepcopy
from pathlib import Path
import shutil
import os
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app import db, artifact_store, task_stage
from app.services import page_flow, review, review_method, review_batches
from app.tools.validate_analysis import _REQUIRED_HEADERS
from app.services.errors import BusinessError
from tests.util import init_env

def tender_citations(tid):
    return [{k: e[k] for k in ('document_id', 'start', 'end')} for e in review.review_basis(tid, True)['catalog'][0]['tender']]


def fixture_candidate(tid, session, batch_id, results):
    """Construct explicit literal facts for synthetic engineering scenarios."""
    rows = deepcopy(results)
    for row in rows:
        row.setdefault('applicability', 'established')
        row.setdefault('applicability_reason', '合成材料中双方均明确描述同一维护服务，特殊情况另行核对。')
        row.setdefault('bid_facts', [{'evidence': e, 'text': review_batches.quote(tid, e)['quote']} for e in row['bid']])
        row.setdefault('conflicts', [{'parameter':'维护期限', 'tender':{'fact':{'evidence':row['tender'][0],'text':review_batches.quote(tid,row['tender'][0])['quote']},'dimension':'period'},'bid':{'fact':row['bid_facts'][0],'dimension':'period'}}] if row['status']=='conflict' and row['bid_facts'] and row.get('tender') and row['tender'][0]['document_id'] in {d['id'] for d in review.documents(tid)} else [])
    return review_batches.candidate(tid, session, batch_id, rows)


@pytest.fixture
def setup(tmp_path, monkeypatch, request):
    task, _ = init_env(tmp_path, monkeypatch)
    tid = task['id']; artifact_store.ensure_task_skeleton(tid)
    for name, text in [('招标.docx', '维护服务期限为三年，须按季度提交维护记录，响应时间不超过四小时。'), ('投标.docx', '维护服务期限为一年，按季度提交维护记录，响应时间不超过四小时。')]:
        d = Document(); d.add_heading('服务要求', 1); d.add_paragraph(text * 4)
        if name == '招标.docx':
            d.add_heading('服务边界', 2); d.add_paragraph('特殊情况下，双方另行协商响应安排。')
        d.save(artifact_store.sources_dir(tid) / name)
    supplements = []
    if getattr(request, 'param', False):
        d = Document(); d.add_heading('补充说明', 1); d.add_paragraph('异地支持期限为五年，与主文件三年约定的适用范围待澄清。' * 5); d.save(artifact_store.sources_dir(tid) / '补遗.docx'); supplements.append('补遗.docx')
    page_flow.set_sources(tid, '招标.docx', supplements)
    review.set_input(tid, ['投标.docx'])
    review.prepare(tid)
    docs = review.documents(tid)
    review.review_basis(tid)
    def citation(doc):
        lines = Path(doc['path']).read_text().splitlines(); i = next(i for i, t in enumerate(lines) if '维护服务期限' in t)
        return {'document_id': doc['id'], 'start': i + 1, 'end': i + 1}
    notes = artifact_store.work_dir(tid) / 'review/analysis'; notes.mkdir(parents=True, exist_ok=True)
    for key, tables in _REQUIRED_HEADERS.items():
        text = ''
        for header in tables:
            text += '| ' + ' | '.join(header) + ' |\n| ' + ' | '.join('---' for _ in header) + ' |\n'
            if key == 'requirements-business':
                text += '| 维护服务期限为三年 | 招标.docx L' + str(citation(docs[0])['start']) + ' |\n'
            text += '\n'
        if key != 'clarifications':
            text += '## coverage 声明\n已检查：服务要求；未检查：无\n'
        if key == 'disqualification': text += '反查词表：废标、无效、否决\n'
        (notes / (key + '.md')).write_text(text)
    basis_id = review.review_basis(tid, True)['catalog'][0]['id']
    value = {'requirements': [{'basis_id': basis_id, 'tender': tender_citations(tid), 'status': 'conflict', 'bid': [citation(next(d for d in docs if d['kind'] == 'bid'))]}], 'issues': [{'id': basis_id + ':issue:1', 'requirement_id': basis_id, 'risk_type': 'suggestion', 'title': '维护年限冲突', 'explanation': '投标承诺一年，要求三年。', 'suggestion': '核对实际维护年限并修订。'}], 'limitations': ['签字、盖章及证件真实性须人工确认。']}
    actions = tmp_path / 'data' / 'actions' / tid; actions.mkdir(parents=True)
    record = {'kind': 'review', 'session_id': 'review-test', 'input_version': review.basis(tid), 'created_at': 1}
    (actions / 'one.json').write_text(json.dumps(record))
    return tid, value, actions / 'one.json', record

def finish(setup):
    tid, value, path, record = setup
    batch = review_batches.next_batch(tid, 'review-test')
    p = {k: v for k, v in value['issues'][0].items() if k not in ('id', 'requirement_id')}
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], [{'slot': 1, 'tender': tender_citations(tid), 'status': 'conflict', 'assessment': '双方维护期限不一致。', 'bid': value['requirements'][0]['bid'], 'problems': [p]}])
    review_batches.accept(tid, 'review-test', proposed)
    review_batches.finalize(tid, 'review-test', value['limitations']); review.finish(tid, 'review-test')
    path.write_text(json.dumps({**record, 'finished_state': 'completed'}))
    return tid

def test_report_version_and_review_decision_survive_reload(setup):
    tid = finish(setup)
    assert review.snapshot(tid)['status'] == 'reviewed'
    review.decide(tid, 'review-test', setup[1]['issues'][0]['id'], 'fixed', '已在 Word 修改，待重审')
    assert review.snapshot(tid)['report']['decisions'][setup[1]['issues'][0]['id']]['status'] == 'fixed'
    assert task_stage.enrich([db.get_task(tid)])[0]['review_status'] == 'reviewed'
    p = artifact_store.sources_dir(tid) / '投标.docx'; d = Document(p); d.add_paragraph('版本更新'); d.save(p)
    assert review.snapshot(tid)['status'] == 'stale'
    with pytest.raises(BusinessError): review.decide(tid, 'review-test', setup[1]['issues'][0]['id'], 'confirmed', '')

def test_bid_never_becomes_tender_supplement(setup):
    tid, *_ = setup
    assert page_flow.source_selection(tid)['excluded'] == ['投标.docx']
    with pytest.raises(BusinessError, match='投标文件不能'): page_flow.set_sources(tid, '招标.docx', ['投标.docx'])
    with pytest.raises(BusinessError): review.set_input(tid, ['招标.docx'])

@pytest.mark.parametrize('invalid', ['quote', 'lines', 'role', 'requirement', 'no-bid', 'incomplete-catalog', 'no-issue'])
def test_invalid_evidence_or_incomplete_report_rejected(setup, invalid):
    tid, value, *_ = setup; value = deepcopy(value); req = value['requirements'][0]
    if invalid == 'quote': req['bid'][0]['quote'] = '这段内容原文不存在'
    if invalid == 'lines': req['bid'][0]['end'] = 999
    if invalid == 'role': req['bid'][0]['document_id'] = 'tender-0'
    if invalid == 'requirement': value['issues'][0]['requirement_id'] = 'unknown'
    if invalid == 'no-bid': req['bid'] = []
    if invalid == 'incomplete-catalog': value['requirements'][0]['basis_id'] = 'unknown-basis'
    if invalid == 'no-issue': value['issues'] = []
    with pytest.raises(ValueError): review.save_report(tid, 'review-test', value)
    assert review.snapshot(tid)['status'] != 'reviewed'

def test_saved_candidate_is_not_reviewed_until_completed(setup):
    tid, value, path, record = setup
    review.save_report(tid, 'review-test', value)
    assert review.snapshot(tid)['report'] is None
    path.write_text(json.dumps({**record, 'finished_state': 'failed'}))
    assert review.snapshot(tid)['status'] == 'failed'
    path.write_text(json.dumps({**record, 'stopped_at': 2}))
    with pytest.raises(BusinessError): review.save_report(tid, 'review-test', value)
    assert review.snapshot(tid)['status'] == 'stopped'

def test_basis_only_changes_for_selected_review_inputs(setup):
    tid, *_ = setup
    before = review.basis(tid)
    (artifact_store.sources_dir(tid) / '未参与文件.docx').write_bytes(b'not selected')
    assert before == review.basis(tid)
    (artifact_store.work_dir(tid) / 'analysis/structure.md').write_text('已更新招标解读')
    assert before != review.basis(tid)


def test_external_unknown_does_not_force_fabricated_bid_evidence_or_issue(setup):
    tid, value, *_ = setup
    value['requirements'][0]['status'] = 'uncertain'
    value['requirements'][0]['bid'] = []
    value['issues'] = []
    value['limitations'] = ['外部真实行为与证照真实性无法由投标文字证明，须人工核验。']
    assert review.compile_report(tid, value)['requirements'][0]['status'] == 'uncertain'


def test_quotes_are_extracted_from_original_lines_and_saved_tampering_rejected(setup):
    tid, value, *_ = setup
    report = review.compile_report(tid, value)
    e = report['requirements'][0]['bid'][0]
    doc = next(d for d in review.documents(tid) if d['id'] == e['document_id'])
    assert e['quote'] == '\n'.join(Path(doc['path']).read_text().splitlines()[e['start'] - 1:e['end']])
    assert review.validate_saved_report(tid, report) == report
    e['quote'] = '模型编造的摘录'
    with pytest.raises(ValueError, match='不一致'):
        review.validate_saved_report(tid, report)


def test_report_uses_original_catalog_and_cannot_rewrite_tender(setup):
    tid, value, *_ = setup
    report = review.compile_report(tid, value)
    assert report['requirements'][0]['category'] == '商务技术要求'
    assert '维护服务期限为三年' in report['requirements'][0]['text']
    assert len(report['coverage']) == 7
    bad = deepcopy(value)
    bad['requirements'][0]['tender'] = []
    with pytest.raises(ValueError): review.compile_report(tid, bad)


def test_incomplete_original_analysis_cannot_produce_review(setup):
    tid, value, *_ = setup
    (artifact_store.work_dir(tid) / 'review/analysis/evaluation.md').unlink()
    with pytest.raises(ValueError, match='评分标准'): review.compile_report(tid, value)


def test_original_checks_are_saved_and_external_process_limits_explicit(setup):
    tid = finish(setup)
    report = review.snapshot(tid)['report']
    assert report['method']['source_commit']
    assert report['analysis']['validation'].startswith('[校验通过]')
    assert any(c['status'] == 'not_applicable' for c in report['checks'])


def test_original_docx_notes_are_marked_for_review(setup):
    tid, *_ = setup
    path = artifact_store.sources_dir(tid) / '投标.docx'
    doc = Document(path)
    level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '1')
    doc.paragraphs[0]._p.get_or_add_pPr().append(level)
    doc.save(path)
    check = review_method.original_checks(tid, review.inputs(tid), 'external')[0]
    assert check['status'] == 'warning'
    assert '大纲级别' in check['detail']


def test_modified_review_analysis_requires_new_review(setup):
    tid = finish(setup)
    p = artifact_store.work_dir(tid) / 'review/analysis/evaluation.md'
    p.write_text(p.read_text() + '\n人工核对说明\n')
    assert review.snapshot(tid)['status'] == 'stale'


def test_platform_analysis_reuse_requires_same_source_and_current_notes(setup):
    tid, *_ = setup
    root = artifact_store.work_dir(tid)
    docs = review.documents(tid)
    tender = docs[0]
    parse = root / 'parse' / tender['name']
    shutil.copytree(Path(tender['path']).parent, parse)
    shutil.copytree(root / 'review/analysis', root / 'analysis', dirs_exist_ok=True)
    (setup[2].parent / 'analysis.json').write_text(json.dumps({'kind': 'analyse', 'created_at': 0,
        'analysis_sources': page_flow.analysis_source_version(tid), 'analysis_progress': [{'done': 8, 'total': 8}]}))
    assert review_method.analysis_directory(tid, docs) == (root / 'analysis', True)
    assert review_method.analysis_context(tid, docs, 'external')['reused_platform_analysis']
    assert review_method.analysis_context(tid, docs, 'external')['sections'][0]['path'].startswith(str(root / 'analysis'))
    # Changed source bytes and changed line mapping must both invalidate reused citations.
    changed = deepcopy(docs); changed[0]['sha256'] = 'different-source'
    assert not review_method.platform_analysis_current(tid, changed)
    original = parse / (tender['name'] + '.md')
    original.write_text(original.read_text() + '\n新行\n')
    assert not review_method.platform_analysis_current(tid, docs)
    shutil.copyfile(tender['path'], original)
    os.utime(root / 'analysis/evaluation.md', (1, 1))
    assert not review_method.platform_analysis_current(tid, docs)
    (parse / (tender['name'] + '.meta.json')).write_text('{}')
    assert not review_method.platform_analysis_current(tid, docs)
    (root / 'analysis/clarifications.md').unlink()
    previous = review_method.previous_flow(tid, 'platform')
    assert len(previous['analysis_files']) == 7
    assert all(item['sha256'] for item in previous['analysis_files'])
    assert '未齐备' in previous['analysis_note']


def test_reparsed_tender_cannot_reuse_old_external_notes(setup):
    tid, value, *_ = setup
    os.utime(artifact_store.work_dir(tid) / 'review/analysis/evaluation.md', (1, 1))
    with pytest.raises(ValueError, match='旧审查要点过期'):
        review.compile_report(tid, value)


def test_platform_review_version_tracks_previous_planning(setup):
    tid, *_ = setup
    (artifact_store.work_dir(tid) / 'review/input.json').unlink()
    before = review.basis(tid)
    path = artifact_store.work_dir(tid) / 'body/关键事实与承诺.md'
    path.write_text('维护年限：三年')
    assert review.basis(tid) != before
    previous = review_method.previous_flow(tid, 'platform')
    assert previous['bidder']['snapshot'] is None
    assert previous['planning'][1]['status'] == 'missing'
    assert previous['planning'][1]['path'] == str(path)


def test_failed_comparisons_never_reused(setup):
    tid, value, *_ = setup
    folder = artifact_store.work_dir(tid) / 'review/drafts'; folder.mkdir()
    (folder / 'failed-model.json').write_text(json.dumps({'report': value}))
    assert 'drafts' not in review.review_basis(tid)
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


def comparison(setup):
    tid, value, *_ = setup
    batch = review_batches.next_batch(tid, 'review-test')
    problem = {k: v for k, v in value['issues'][0].items() if k not in ('id', 'requirement_id')}
    return batch, [{'slot': 1, 'tender': tender_citations(tid), 'status': 'conflict', 'assessment': '投标一年，要求三年。', 'bid': value['requirements'][0]['bid'], 'problems': [problem]}]


@pytest.mark.parametrize('bad', ['foreign-id', 'inline-line', 'inline-page', 'wrong-slot', 'tampered-quote'])
def test_batches_reject_binding_and_citation_failures(setup, bad):
    tid, *_ = setup
    batch, results = comparison(setup)
    if bad == 'foreign-id': results[0]['problems'][0]['requirement_id'] = 'other'
    if bad == 'inline-line': results[0]['problems'][0]['explanation'] = '招标L1047要求三年。'
    if bad == 'inline-page': results[0]['assessment'] = '第123行证明。'
    if bad == 'wrong-slot': results[0]['slot'] = 2
    if bad in ('foreign-id', 'inline-line', 'inline-page', 'wrong-slot'):
        with pytest.raises(ValueError): fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    else:
        proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
        if bad == 'tampered-quote': proposed['comparisons'][0]['bid'][0]['quote'] = '伪造内容'
        with pytest.raises(ValueError): review_batches.accept(tid, 'review-test', proposed)
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


def test_batch_binding_and_completion_are_immutable(setup):
    tid = finish(setup)
    envelope = review.snapshot(tid)['report']
    req = envelope['report']['requirements'][0]
    assert envelope['report']['issues'][0]['requirement_id'] == req['id']
    assert envelope['report']['issues'][0]['id'] == req['id'] + ':issue:1'
    p = review_batches.path(tid, 'review-test'); data = json.loads(p.read_text())
    data['batches'][0]['comparisons'][0]['status'] = 'responded'; p.write_text(json.dumps(data))
    with pytest.raises(ValueError, match='绑定被修改'): review.finish(tid, 'review-test')


def test_full_report_without_complete_batches_cannot_finish(setup):
    tid, value, *_ = setup
    review.save_report(tid, 'review-test', value)
    with pytest.raises(ValueError, match='尚未完整'): review.finish(tid, 'review-test')


def test_background_fields_do_not_accept_unrelated_bid_evidence(setup, monkeypatch):
    tid, value, *_ = setup
    row = review.review_basis(tid, True)['catalog'][0] | {'category': '结构事实', 'text': '字段：代理机构；值：测试机构'}
    monkeypatch.setattr(review_batches, 'context', lambda *_: ('background-version', [[row]]))
    batch = review_batches.next_batch(tid, 'review-test')
    assert batch['requirements'][0]['response_duty'] == 'information_only'
    with pytest.raises(ValueError, match='招标背景信息无需'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], [{'slot': 1, 'tender': tender_citations(tid), 'status': 'responded', 'assessment': '投标封面写了项目名称，因此响应代理机构。', 'bid': value['requirements'][0]['bid'], 'problems': []}])
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], [{'slot': 1, 'tender': tender_citations(tid), 'status': 'informational', 'assessment': '仅为招标程序背景，无投标文件响应义务。', 'bid': [], 'problems': []}])
    assert proposed['comparisons'][0]['bid'] == []


def test_bid_duty_cannot_be_hidden_as_background(setup):
    tid, *_ = setup
    batch = review_batches.next_batch(tid, 'review-test')
    with pytest.raises(ValueError, match='投标义务不能'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], [{'slot': 1, 'tender': tender_citations(tid), 'status': 'informational', 'assessment': '无需核查维护年限。', 'bid': [], 'problems': []}])


def test_old_severity_cannot_be_mapped_to_business_classification(setup):
    tid, value, *_ = setup
    value = deepcopy(value)
    value['issues'][0].pop('risk_type')
    value['issues'][0]['severity'] = 'high'
    with pytest.raises(ValueError): review.compile_report(tid, value)
    batch, results = comparison(setup)
    results[0]['problems'][0].pop('risk_type')
    results[0]['problems'][0]['severity'] = 'high'
    with pytest.raises(ValueError): fixture_candidate(tid, 'review-test', batch['batch_id'], results)


def test_same_auditor_can_correct_prior_batch_before_completion(setup, monkeypatch):
    tid, value, *_ = setup
    row = review.review_basis(tid, True)['catalog'][0]
    monkeypatch.setattr(review_batches, 'context', lambda *_: ('current-version', [[row], [row | {'id': 'other-requirement'}]]))
    batch, results = comparison(setup)
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', batch['batch_id'], results))
    following = review_batches.next_batch(tid, 'review-test')
    assert following['accepted'] == 1
    results[0]['assessment'] = '后文出现更完整的响应，但范围仍需确认。'
    results[0]['status'] = 'uncertain'
    results[0]['problems'][0]['risk_type'] = 'confirmation'
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', batch['batch_id'], results))
    assert review_batches.next_batch(tid, 'review-test') == following
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', following['batch_id'], results))
    rows = review_batches.completed(tid, 'review-test')
    assert len(rows) == 2
    assert rows[0]['assessment'] == results[0]['assessment']
    assert rows[0]['problems'][0]['risk_type'] == 'confirmation'
    with pytest.raises(ValueError): fixture_candidate(tid, 'review-test', 'foreign-version:0', results)
    with pytest.raises(ValueError): fixture_candidate(tid, 'review-test', 'current-version:2', results)


def test_completed_report_cannot_be_corrected_in_place(setup):
    tid = finish(setup)
    with pytest.raises(ValueError, match='本轮已结束'):
        fixture_candidate(tid, 'review-test', 'old:0', [])


def test_expanded_and_separate_tender_citations_survive_full_report(setup):
    tid, value, *_ = setup
    batch, results = comparison(setup)
    doc = next(d for d in review.documents(tid) if d['kind'] == 'tender')
    lines = Path(doc['path']).read_text().splitlines()
    original = deepcopy(results[0]['tender'])
    results[0]['tender'][0]['start'] -= 1
    extra_line = next(i + 1 for i, text in enumerate(lines) if '特殊情况下' in text)
    results[0]['tender'].append({'document_id': doc['id'], 'start': extra_line, 'end': extra_line})
    results[0]['assessment'] = '维护期限不同，另有特殊情况下协商安排的约定，适用范围待澄清。'
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    review_batches.accept(tid, 'review-test', proposed)
    review_batches.finalize(tid, 'review-test', value['limitations'])
    review.finish(tid, 'review-test')
    envelope = json.loads(review._report_path(tid, 'review-test').read_text())
    req = envelope['report']['requirements'][0]
    assert len(req['tender']) == 2
    assert req['tender'][-1]['quote'] == lines[extra_line - 1]
    assert req['tender'][0]['start'] == original[0]['start'] - 1
    assert review.review_basis(tid, True)['catalog'][0]['tender'] == original
    assert review.validate_saved_report(tid, envelope['report']) == envelope['report']
    # Even a forged batch digest cannot make omitted original evidence valid.
    ledger = json.loads(review_batches.path(tid, 'review-test').read_text())
    ledger['batches'][0]['comparisons'][0]['tender'] = [req['tender'][-1]]
    ledger['batches'][0]['digest'] = page_flow.digest({k: ledger['batches'][0][k] for k in ('batch_id', 'comparisons')})
    review_batches.path(tid, 'review-test').write_text(json.dumps(ledger))
    with pytest.raises(ValueError, match='保留完整要点依据'):
        review.finish(tid, 'review-test')


@pytest.mark.parametrize('setup', [True], indirect=True)
def test_main_and_supplement_conflict_sources_are_both_preserved(setup):
    tid, value, *_ = setup
    batch, results = comparison(setup)
    doc = next(d for d in review.documents(tid) if d['name'] == '补遗.docx')
    lines = Path(doc['path']).read_text().splitlines()
    n = next(i + 1 for i, text in enumerate(lines) if '异地支持期限' in text)
    results[0]['tender'].append({'document_id': doc['id'], 'start': n, 'end': n})
    results[0]['status'] = 'uncertain'
    results[0]['problems'][0]['risk_type'] = 'confirmation'
    results[0]['assessment'] = '主文件与补遗的适用范围待澄清，两条原文并列，不决定优先级。'
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    review_batches.accept(tid, 'review-test', proposed)
    review_batches.finalize(tid, 'review-test', value['limitations'])
    review.finish(tid, 'review-test')
    report = json.loads(review._report_path(tid, 'review-test').read_text())['report']
    assert {e['document_id'] for e in report['requirements'][0]['tender']} == {'tender-0', 'tender-1'}
    assert '异地支持期限为五年' in report['requirements'][0]['tender'][1]['quote']
    assert report['requirements'][0]['status'] == 'uncertain'
    assert report['issues'][0]['risk_type'] == 'confirmation'


@pytest.mark.parametrize('failure', ['omit', 'role', 'unknown', 'range', 'quote', 'truncate', 'tamper'])
def test_tender_citation_failures_are_rejected(setup, failure):
    tid, *_ = setup
    batch, results = comparison(setup)
    e = results[0]['tender'][0]
    if failure == 'omit': results[0].pop('tender')
    if failure == 'role': e['document_id'] = 'bid-0'
    if failure == 'unknown': e['document_id'] = 'foreign-tender'
    if failure == 'range': e['end'] = 999
    if failure == 'quote': e['quote'] = '编造的原文'
    if failure == 'truncate': e['start'] = e['end'] = 1
    if failure != 'tamper':
        with pytest.raises(ValueError): fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    else:
        proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
        proposed['comparisons'][0]['tender'][0]['quote'] = '编造的原文'
        with pytest.raises(ValueError, match='保存期间引用'):
            review_batches.accept(tid, 'review-test', proposed)
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


@pytest.mark.parametrize('change', ['add', 'remove', 'main'])
def test_source_set_changes_invalidate_platform_and_review_analysis(setup, change):
    tid, _, actions, _ = setup
    root = artifact_store.work_dir(tid)
    tender = review.documents(tid)[0]
    shutil.copytree(Path(tender['path']).parent, root / 'parse' / tender['name'])
    shutil.copytree(root / 'review/analysis', root / 'analysis', dirs_exist_ok=True)
    extra = artifact_store.sources_dir(tid) / '已解析补充.docx'
    d = Document(); d.add_paragraph('补遗新要求：服务改为五年，须提交新的认证证明。' * 8); d.save(extra)
    from app import document
    document.prepare(extra, root / 'parse' / extra.name)
    if change == 'remove':
        page_flow.set_sources(tid, tender['name'], [extra.name])
        review.prepare(tid); review.review_basis(tid)
        for p in (root / 'review/analysis').glob('*.md'): p.touch()
    actions.with_name('analysis.json').write_text(json.dumps({'kind': 'analyse', 'created_at': 0,
        'analysis_sources': page_flow.analysis_source_version(tid), 'analysis_progress': [{'done': 8, 'total': 8}]}))
    for p in (root / 'analysis').glob('*.md'): p.touch()
    review.prepare(tid)
    if change != 'remove': assert review_method.analysis_directory(tid, review.documents(tid))[1]
    if change == 'add': page_flow.set_sources(tid, tender['name'], [extra.name])
    elif change == 'remove': page_flow.set_sources(tid, tender['name'], [])
    else: page_flow.set_sources(tid, extra.name, [])
    review.prepare(tid)
    assert not review_method.analysis_directory(tid, review.documents(tid))[1]
    context = review.review_basis(tid)
    assert all(s['stale'] and not s['exists'] for s in context['sections'])
    with pytest.raises(ValueError, match='旧审查要点过期'): review.review_basis(tid, True)
    # Rewriting one file cannot certify the other seven against the new source set.
    (root / 'review/analysis/structure.md').touch()
    with pytest.raises(ValueError, match='旧审查要点过期'): review.review_basis(tid, True)


@pytest.mark.parametrize('name', ['L2026招标.docx', '招标L12-L15（同文件）.docx', '招标;L4补遗.docx'])
def test_catalog_ignores_line_like_tokens_inside_real_filenames(setup, name):
    tid, _, _, _ = setup
    root = artifact_store.work_dir(tid)
    shutil.copyfile(artifact_store.sources_dir(tid) / '招标.docx', artifact_store.sources_dir(tid) / name)
    page_flow.set_sources(tid, name, [])
    review.prepare(tid); review.review_basis(tid)
    for p in (root / 'review/analysis').glob('*.md'):
        p.write_text(p.read_text().replace('招标.docx', name))
    catalog = review.review_basis(tid, True)['catalog']
    assert len(catalog[0]['tender']) == 1
    assert catalog[0]['tender'][0] == {'document_id': 'tender-0', 'start': 2, 'end': 2}


def test_terminal_findings_are_visible_decidable_and_validated(setup):
    tid, _, _, _ = setup
    source = artifact_store.sources_dir(tid) / '投标.docx'
    d = Document(source); p = d.add_paragraph('授权代表联系方式：【待补】')
    d.add_comment(p.runs, text='工程验收：签章附件待核验', author='fixture'); d.save(source)
    review.prepare(tid)
    setup[3]['input_version'] = review.basis(tid); setup[2].write_text(json.dumps(setup[3]))
    finish(setup)
    report = review.snapshot(tid)['report']['report']
    issue = next(i for i in report['document_issues'] if '待补' in i['title'])
    assert issue['document_ids'] == ['bid-0'] and issue['location'] == 'P3'
    assert issue['risk_type'] == 'suggestion'
    review.decide(tid, 'review-test', issue['id'], 'pending', '终稿需要补齐')
    assert review.snapshot(tid)['report']['decisions'][issue['id']]['note'] == '终稿需要补齐'
    assert any('工程验收：签章附件待核验' in n for c in report['check_notes'] for n in c['notes'])
    assert not any('未传 block_ids' in n for c in report['check_notes'] for n in c['notes'])  # 实际素材记录已校验，不能再保留旧的跳过提示
    report['document_issues'] = []
    with pytest.raises(ValueError, match='不一致'): review.validate_saved_report(tid, report)


def test_unreviewed_list_does_not_hash_inputs(setup, monkeypatch):
    tid, _, record, _ = setup
    record.unlink()
    monkeypatch.setattr(page_flow, 'file_hash', lambda *_: pytest.fail('列表不应读取大文件'))
    assert task_stage.enrich([db.get_task(tid)])[0]['review_status'] == 'unreviewed'


def test_file_digest_survives_process_cache_reset_and_same_size_edit(setup, monkeypatch):
    import builtins
    tid, *_ = setup
    path = artifact_store.sources_dir(tid) / '大文件数据.bin'
    path.write_bytes(b'a' * 1024 * 1024)
    original_open = builtins.open
    reads = []
    def counted_open(filename, *args, **kwargs):
        if Path(filename) == path.resolve(): reads.append(str(filename))
        return original_open(filename, *args, **kwargs)
    monkeypatch.setattr(builtins, 'open', counted_open)
    first = page_flow.file_hash(path)
    page_flow._file_hash.cache_clear()
    assert page_flow.file_hash(path) == first and len(reads) == 1
    previous = path.stat()
    path.write_bytes(b'b' * previous.st_size)
    os.utime(path, ns=(previous.st_atime_ns, previous.st_mtime_ns))
    page_flow._file_hash.cache_clear()
    assert page_flow.file_hash(path) != first and len(reads) == 2


def test_reviewed_list_uses_persistent_digests(setup, monkeypatch):
    import builtins
    tid = finish(setup)
    page_flow._file_hash.cache_clear()
    original_open = builtins.open
    def no_source_open(filename, *args, **kwargs):
        if Path(filename).parent == artifact_store.sources_dir(tid): pytest.fail('未变化的原件不应重复读取')
        return original_open(filename, *args, **kwargs)
    monkeypatch.setattr(builtins, 'open', no_source_open)
    assert task_stage.enrich([db.get_task(tid)])[0]['review_status'] == 'reviewed'


@pytest.mark.parametrize('continuous', [False, True])
def test_status_digest_rereads_a_changed_file_without_caching_mixed_bytes(setup, monkeypatch, continuous):
    import builtins, hashlib
    tid, *_ = setup
    path = artifact_store.work_dir(tid) / '执行中文件.md'
    path.write_bytes(b'initial content')
    original_open = builtins.open
    reads = []
    class ConcurrentWrite:
        def __enter__(self):
            self.stream = original_open(path, 'rb')
            reads.append(True)
            return self.stream.__enter__()
        def __exit__(self, *args):
            result = self.stream.__exit__(*args)
            if continuous or len(reads) == 1:
                path.write_bytes(('completed version ' + str(len(reads))).encode())
            return result
    def changing_open(filename, *args, **kwargs):
        return ConcurrentWrite() if Path(filename) == path.resolve() else original_open(filename, *args, **kwargs)
    monkeypatch.setattr(builtins, 'open', changing_open)
    key = 'file-digest:' + hashlib.sha256(str(path.resolve()).encode()).hexdigest()
    if continuous:
        with pytest.raises(page_flow.FileChangedDuringRead, match='文件在读取期间已变化'):
            page_flow.file_hash(path)
        assert len(reads) == 3
        assert db.get_setting(key) is None
    else:
        digest = page_flow.file_hash(path)
        assert len(reads) == 2
        assert digest == hashlib.sha256(path.read_bytes()).hexdigest()
        assert json.loads(db.get_setting(key))['sha256'] == digest


def test_file_digest_is_reused_by_an_actual_second_python_process(setup):
    import subprocess, sys
    tid, *_ = setup
    path = artifact_store.sources_dir(tid) / '跨进程数据.bin'
    path.write_bytes(b'x' * 1024 * 1024)
    code = '''import sys, builtins, json
from pathlib import Path
from app import db
from app.services.page_flow import file_hash
db.init_db()
target=Path(sys.argv[1]).resolve(); original=builtins.open; reads=[]
def counted(name,*args,**kwargs):
    if Path(name)==target: reads.append(str(name))
    return original(name,*args,**kwargs)
builtins.open=counted
digest=file_hash(target)
print(json.dumps({'digest':digest,'reads':len(reads)}))'''
    env = {**os.environ, 'PYTHONPATH': str(Path(__file__).parents[1])}
    results = [json.loads(subprocess.check_output([sys.executable, '-c', code, str(path)], env=env)) for _ in range(2)]
    assert results[0]['reads'] == 1 and results[1]['reads'] == 0
    assert results[0]['digest'] == results[1]['digest']


@pytest.mark.parametrize('title', [r'问题\n补充', '问题\n补充', r'问题\r补充', '1. 标题', '2、标题', '- 标题'])
def test_review_title_format_rejected_before_batch_save(setup, title):
    tid, *_ = setup
    batch, results = comparison(setup)
    results[0]['problems'][0]['title'] = title
    with pytest.raises(ValueError, match='单行纯文本'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    assert not review_batches.path(tid, 'review-test').exists()


@pytest.mark.parametrize('title', ['3.5小时的安排需核对', '2026年资格证明需核对', '费用范围需明确'])
def test_review_title_keeps_business_numbers(setup, title):
    tid, *_ = setup
    batch, results = comparison(setup)
    results[0]['problems'][0]['title'] = title
    assert fixture_candidate(tid, 'review-test', batch['batch_id'], results)['comparisons'][0]['problems'][0]['title'] == title


def test_review_overview_derived_from_saved_items_and_checks(setup, monkeypatch):
    tid, value, *_ = setup
    monkeypatch.setattr(review_method, 'original_checks', lambda *_: [{'name': '终稿', 'status': 'warning', 'document_ids': ['bid-0'], 'issues': [], 'warnings': ['需修改占位符'], 'notes': []}])
    result = review.compile_report(tid, value)
    assert 'summary' not in result
    assert len(result['summary_items']) == 2
    assert '1 项招标要点' in result['summary_items'][0]
    assert '2 条审查意见' in result['summary_items'][1]
    assert '建议修改 2 条' in result['summary_items'][1]
    assert all(r'\n' not in text and '\n' not in text for text in result['summary_items'])
    result['summary_items'][1] = '本次没有问题。'
    with pytest.raises(ValueError, match='已保存报告'):
        review.validate_saved_report(tid, result)
    with pytest.raises(ValueError):
        review.compile_report(tid, value | {'summary': '旧模型概要不得进入新报告'})


@pytest.mark.parametrize('limitations', [[r'签章\n人工确认'], ['1. 签章人工确认']])
def test_review_limitations_format_rejected(setup, limitations):
    with pytest.raises(ValueError, match='单行纯文本'):
        review.compile_report(setup[0], setup[1] | {'limitations': limitations})


def test_tender_conditions_are_generated_from_source_and_cannot_be_dropped(setup):
    tid, *_ = setup
    batch, results = comparison(setup)
    doc = next(d for d in review.documents(tid) if d['kind'] == 'tender')
    lines = Path(doc['path']).read_text().splitlines()
    n = next(i + 1 for i, text in enumerate(lines) if '特殊情况下' in text)
    conditional = {'document_id': doc['id'], 'start': n, 'end': n}
    results[0]['tender'].append(conditional)
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    assert proposed['comparisons'][0]['qualifications'] == proposed['comparisons'][0]['tender']
    tampered = deepcopy(proposed)
    tampered['comparisons'][0]['qualifications'] = []
    with pytest.raises(ValueError, match='引用或要点已变化'):
        review_batches.accept(tid, 'review-test', tampered)
    review_batches.accept(tid, 'review-test', proposed)
    review_batches.finalize(tid, 'review-test', ['仅验证工程取证约束。'])
    envelope = json.loads(review._report_path(tid, 'review-test').read_text())
    rid = batch['requirements'][0]['id']
    assert lines[n - 1] in envelope['audit']['assessments'][rid]
    assert envelope['audit']['evidence_notes'][rid]['qualifications'][-1]['quote'] == lines[n - 1]
    review.finish(tid, 'review-test')
    envelope['audit']['assessments'][rid] = results[0]['assessment']
    review._report_path(tid, 'review-test').write_text(json.dumps(envelope))
    with pytest.raises(ValueError, match='说明或原文限定'):
        review.finish(tid, 'review-test')


@pytest.mark.parametrize('status,risk', [('conflict','confirmation'), ('responded','confirmation'), ('uncertain','scoring'), ('uncertain','disqualification'), ('uncertain','suggestion')])
def test_ambiguous_applicability_cannot_be_saved_as_established_consequence(setup, status, risk):
    tid, *_ = setup
    batch, results = comparison(setup)
    results[0].update(applicability='uncertain', status=status)
    results[0]['problems'][0]['risk_type'] = risk
    with pytest.raises(ValueError, match='条款适用范围或条件未明确'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


def test_ambiguous_scope_keeps_actual_evidence_and_confirmation(setup):
    tid, *_ = setup
    batch, results = comparison(setup)
    results[0].update(applicability='uncertain', status='uncertain')
    results[0]['problems'][0]['risk_type'] = 'confirmation'
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    review_batches.accept(tid, 'review-test', proposed)
    row = review_batches.next_batch(tid, 'review-test', batch['batch_id'])['saved_results'][0]
    assert row['status'] == 'uncertain' and row['problems'][0]['risk_type'] == 'confirmation'
    assert row['tender'] and row['bid'] and row['bid_facts']


@pytest.mark.parametrize('bad', ['tender-as-bid', 'invented-fact', 'wrong-role', 'outside-card', 'no-facts'])
def test_literal_bid_facts_reject_source_confusion_and_unanchored_claims(setup, bad):
    tid, *_ = setup
    batch, results = comparison(setup)
    ref = results[0]['bid'][0]
    fact = {'evidence': ref, 'text': review_batches.quote(tid, ref)['quote']}
    if bad == 'tender-as-bid': fact['text'] = '维护服务期限为三年'
    if bad == 'invented-fact': fact['text'] = '投标人提供了所有资格查询截图'
    if bad == 'wrong-role': fact['evidence'] = results[0]['tender'][0]
    if bad == 'outside-card': fact['evidence'] = ref | {'start': 1, 'end': 1}
    results[0]['bid_facts'] = [] if bad == 'no-facts' else [fact]
    with pytest.raises(ValueError):
        fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


def test_required_new_evidence_contract_has_no_old_payload_fallback(setup):
    tid, *_ = setup
    batch, results = comparison(setup)
    with pytest.raises(ValueError, match='applicability'):
        review_batches.candidate(tid, 'review-test', batch['batch_id'], results)


def test_round_evidence_index_and_prior_batch_read_support_later_correction(setup, monkeypatch):
    tid, *_ = setup
    row = review.review_basis(tid, True)['catalog'][0]
    monkeypatch.setattr(review_batches, 'context', lambda *_: ('same-round', [[row], [row | {'id': 'later-requirement'}]]))
    first, results = comparison(setup)
    results[0].update(status='uncertain', bid=[], bid_facts=[], assessment='已查早期范围未定位，后文尚待精读。')
    results[0]['problems'][0]['risk_type'] = 'confirmation'
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', first['batch_id'], results))
    second = review_batches.next_batch(tid, 'review-test')
    bid = next(d for d in review.documents(tid) if d['kind'] == 'bid')
    lines = Path(bid['path']).read_text().splitlines()
    n = next(i + 1 for i, text in enumerate(lines) if '维护服务期限' in text)
    ref = {'document_id': bid['id'], 'start': n, 'end': n}
    later = deepcopy(results)
    later[0].update(status='responded', bid=[ref], bid_facts=[{'evidence': ref, 'text': '按季度提交维护记录'}], problems=[], assessment='本项季度记录有文字回应。')
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', second['batch_id'], later))
    memo = Path(second['evidence_index_path']).read_text()
    assert '按季度提交维护记录' in memo and second['batch_id'] in memo
    assert '本项季度记录有文字回应' not in memo  # Receipts contain original text, not an interpretation reused as fact.
    saved = review_batches.next_batch(tid, 'review-test', first['batch_id'])['saved_results']
    corrected = [{k: saved[0][k] for k in ('slot','status','assessment','tender','bid','applicability','applicability_reason','bid_facts','conflicts','problems')}]
    corrected[0].update(status='uncertain', conflicts=[], bid=[ref], bid_facts=later[0]['bid_facts'], assessment='后文已定位季度记录，期限差异仍需核对。')
    review_batches.accept(tid, 'review-test', review_batches.candidate(tid, 'review-test', first['batch_id'], corrected))
    rows = review_batches.completed(tid, 'review-test')
    assert len(rows) == 2 and rows[0]['bid_facts'] == later[0]['bid_facts']
    assert review_batches.next_batch(tid, 'review-test')['complete']
    with pytest.raises(ValueError, match='本轮已经保存'):
        review_batches.next_batch(tid, 'review-test', 'other-round:0')


def test_model_batch_uses_original_clauses_instead_of_compressed_analysis_claims(setup, monkeypatch):
    tid, *_ = setup
    row = review.review_basis(tid, True)['catalog'][0]
    misleading = row | {'text': '概括错误：任何服务均适用最低课时要求'}
    monkeypatch.setattr(review_batches, 'context', lambda *_: ('source-first', [[misleading]]))
    batch = review_batches.next_batch(tid, 'review-test')
    assert 'text' not in batch['requirements'][0]
    assert batch['requirements'][0]['analysis_locator'] == misleading['text']
    assert '最低课时' not in batch['requirements'][0]['tender'][0]['quote']
    assert batch['requirements'][0]['tender'][0]['quote'] == review_batches.quote(tid, row['tender'][0])['quote']


def test_shared_source_receipt_links_exact_intersections_without_rewriting_findings(setup, monkeypatch):
    tid, *_ = setup
    row = review.review_basis(tid, True)['catalog'][0]
    monkeypatch.setattr(review_batches, 'context', lambda *_: ('source-links', [[row], [row | {'id':'other-duty'}]]))
    first, results = comparison(setup)
    review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', first['batch_id'], results))
    second = review_batches.next_batch(tid, 'review-test')
    changed = deepcopy(results)
    changed[0].update(status='uncertain', applicability='uncertain')
    changed[0]['problems'][0]['risk_type'] = 'confirmation'
    receipt = review_batches.accept(tid, 'review-test', fixture_candidate(tid, 'review-test', second['batch_id'], changed))
    assert receipt['shared_evidence']
    for link in receipt['shared_evidence']:
        assert link['earlier_items'][0]['batch_id'] == first['batch_id']
        assert link['source']['quote'] == review_batches.quote(tid, {k:link['source'][k] for k in ('document_id','start','end')})['quote']
    assert review_batches.next_batch(tid, 'review-test', first['batch_id'])['saved_results'][0]['status'] == 'conflict'
    assert review_batches.next_batch(tid, 'review-test', second['batch_id'])['saved_results'][0]['status'] == 'uncertain'


def test_applicability_reason_is_required_before_a_consequence_is_accepted(setup):
    tid, *_ = setup
    batch, results = comparison(setup)
    candidate = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    payload = review_batches.comparison_input(candidate['comparisons'][0])
    del payload['applicability_reason']
    with pytest.raises(ValueError, match='applicability_reason'):
        review_batches.candidate(tid, 'review-test', batch['batch_id'], [payload])


@pytest.mark.parametrize('left,right', [('time_window','effort'), ('response_delay','effort'), ('date','period'), ('quantity','categorical')])
def test_conflict_rejects_incomparable_parameter_dimensions(setup, left, right):
    tid, *_ = setup
    batch, results = comparison(setup)
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    result = review_batches.comparison_input(proposed['comparisons'][0])
    result['conflicts'][0]['tender']['dimension'] = left
    result['conflicts'][0]['bid']['dimension'] = right
    with pytest.raises(ValueError, match='不能比较不同参数类型'):
        review_batches.candidate(tid, 'review-test', batch['batch_id'], [result])
    assert review_batches.next_batch(tid, 'review-test')['accepted'] == 0


def test_conflict_needs_opposing_literal_source_pair(setup):
    tid, *_ = setup
    batch, results = comparison(setup)
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], results)
    result = review_batches.comparison_input(proposed['comparisons'][0])
    result['conflicts'] = []
    with pytest.raises(ValueError, match='成对原文'):
        review_batches.candidate(tid, 'review-test', batch['batch_id'], [result])
    result['conflicts'] = deepcopy(proposed['comparisons'][0]['conflicts'])
    result['conflicts'][0]['bid']['fact']['text'] = '本项目只提供单次培训且不承担费用'
    with pytest.raises(ValueError, match='逐字摘录'):
        review_batches.candidate(tid, 'review-test', batch['batch_id'], [result])


def test_new_review_in_same_main_session_keeps_prior_report_and_batches(setup):
    tid = finish(setup)
    _, _, action_path, original = setup
    report_path = review._report_path(tid, original['session_id'])
    batch_path = review_batches.path(tid, original['session_id'])
    old_report, old_batches = report_path.read_bytes(), batch_path.read_bytes()
    new_record = {**original, 'created_at': original['created_at'] + 1}
    action_path.with_name('two.json').write_text(json.dumps(new_record))
    assert review._report_path(tid, original['session_id']) != report_path
    assert review_batches.path(tid, original['session_id']) != batch_path
    assert review.snapshot(tid)['report'] is None
    assert review_batches.read(tid, original['session_id'])[0]['batches'] == []
    assert report_path.read_bytes() == old_report
    assert batch_path.read_bytes() == old_batches
