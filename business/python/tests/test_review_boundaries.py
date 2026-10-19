"""本轮五项复核缺陷的边界用例；不替代评分语义或既有正文完整性验收。"""
import asyncio
import hashlib
import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app import artifact_store, db, runctx
from app.contracts.tender_directory import ResponseDocument
from app.services import build_state, document_format, outline_contract as oc, page_flow
from app.services.artifacts import ContentUpdate, update_artifact_content
from app.tools.assemble_tender import assemble_tender
from app.tools.document_format import docx_format_rule
from app.tools.docx_ops import docx_section_create, docx_assemble_volume, templates_dir, TEMPLATE_SETTING_KEY
from tests.test_assemble_tender import env
from tests.test_auto_export import export, verify_downloads
from tests.test_document_pagination import setup, seed, node
from tests.test_docx_ops import _tiny_png


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fragment(doc):
    """只用于测试构造页面当前目录的实际子稿，不代替生产分类逻辑。"""
    flat = list(oc._nodes(doc['directory']))
    return '\n'.join([
        '# 响应文件：'+doc['name'], doc.get('scope', ''), '', '## 目录',
        *['  '*(len(path)-1)+'- '+path[-1] for path, _ in flat], '', '## 来源标注',
        *['- '+path[-1]+' :: '+', '.join(n.get('来源位置', [])) for path, n in flat if n.get('来源位置')],
        '', '## 目录说明',
        *['- '+path[-1]+' :: 交付形态='+n.get('交付形态', '')+' | 归位理由='+n.get('归位理由', '')+
          ' | 理由来源='+', '.join(n.get('理由来源', []))+' | 概述='+n.get('节点概述', '') for path, n in flat], ''])


@pytest.mark.parametrize('volumes', [1, 2], ids=['single-tender', 'sealed-multi-volume-tender'])
def test_saved_directory_is_baseline_despite_old_markdown_and_fragments(env, volumes):
    tid, root = env['task']['id'], env['out']
    file = root/'outline/tender-response-docs.md'
    if volumes == 2:
        file.write_text(file.read_text()+'\n'+file.read_text().replace('响应文件：技术部分', '响应文件：商务部分'))
    assert assemble_tender().startswith('[组装发布成功]')
    current = page_flow.directory(tid)
    edited = deepcopy(current['content'])
    for doc in edited['response_documents']:
        doc['directory'].reverse()
        doc['directory'].append({'目录名称': '人工补充培训方案', 'level': 1, 'children': [],
                                '交付形态': '正文编写', '归位理由': '用户按当前招标要求人工归集',
                                '节点概述': '原有培训安排、人数及实施时间全部保留'})
    asyncio.run(update_artifact_content(current['artifact_id'], ContentUpdate(content=edited, base_content_seq=current['content_seq'])))
    body = root/'body'
    for doc in edited['response_documents']:
        folder = body/doc['name'] if volumes == 2 else body
        folder.mkdir(parents=True, exist_ok=True)
        d = Document(); d.add_heading('人工补充培训方案', 1); d.add_paragraph('人工已填写的完整培训正文。')
        d.save(folder/'人工补充培训方案.docx')
        old = oc.fragment_path(tid, doc['name']); old.parent.mkdir(parents=True, exist_ok=True)
        old.write_text(fragment(current['content']['response_documents'][edited['response_documents'].index(doc)]))
    originals = {p: p.read_bytes() for p in body.rglob('*.docx')}
    old_markdown = file.read_bytes()
    oc.begin(tid, 'review-current-saved-directory')
    value = oc._load(tid)
    assert value['baseline'] == [ResponseDocument.model_validate(d).model_dump() for d in edited['response_documents']]
    assert [d['directory'][-1]['目录名称'] for d in value['baseline']] == ['人工补充培训方案']*volumes
    assert [d['directory'][0]['目录名称'] for d in value['baseline']] == [d['directory'][0]['目录名称'] for d in edited['response_documents']]
    assert {b['volume'] for b in value['body_inventory']} == {d['name'] for d in edited['response_documents']}
    assert json.loads(oc.context_path(tid).read_text())['baseline_source'] == '页面当前已保存目录'
    oc.plan_batch(tid, oc.PlanBatch())
    for doc in edited['response_documents']:
        target = Path(oc.dispatch(tid, doc['name'])['fragment_path'])
        with pytest.raises(ValueError, match='一级标题|原有节点'):
            oc.check(tid, doc['name'])
        target.write_text(fragment(doc))
        assert oc.check(tid, doc['name'])['ok']
    assert file.read_bytes() == old_markdown
    assert page_flow.directory(tid)['content'] == edited
    assert all(p.read_bytes() == data for p, data in originals.items())


@pytest.mark.parametrize('landscape,cancelled', [(False, False), (True, True)], ids=['portrait-new-explicit-format', 'landscape-requirement-cancelled'])
def test_custom_template_export_and_stale_format_recovery_preserve_sources(tmp_path, monkeypatch, landscape, cancelled):
    tid, root = setup(tmp_path, monkeypatch)
    try:
        template = Document()
        for style in list(template.styles):
            if style.name == 'TOC Heading' or style.name.lower() in ('toc 1', 'toc 2', 'toc 3'):
                style.delete()
        if landscape:
            template.sections[0].page_width = Inches(11)
            template.sections[0].page_height = Inches(8.5)
        template_path = templates_dir()/'自定义无目录样式.docx'
        template.save(template_path); template_hash = digest(template_path)
        db.set_setting(TEMPLATE_SETTING_KEY, template_path.name)
        parent = '实施组织方案' if landscape else '项目管理方案'
        child = '验收与运维' if landscape else '人员培训'
        seed(tid, [node('封面'), node('目录'), node(parent, [{'目录名称':child, 'level':2, 'children':[]}]), node('证明附件')])
        evidence = root/'parse/tender/format.md'; evidence.parent.mkdir(parents=True)
        quote = '正文采用仿宋，三号字，固定28磅行距。一级标题采用黑体，二号字。'
        evidence.write_text(quote)
        for target, values in [('body', {'font':'仿宋', 'font_size':16, 'line_spacing':28, 'spacing_rule':'exact'}),
                               ('heading1', {'font':'黑体', 'font_size':22})]:
            docx_format_rule(target, **values, evidence_path='parse/tender/format.md', evidence_quote=quote)
        for title in ['封面', parent, child, '证明附件']:
            docx_section_create('body/'+title, title, paragraphs=title+'原始正文与具体事实12345。')
        source = root/('body/'+parent+'.docx'); d = Document(source)
        table = d.add_table(rows=1, cols=2); table.cell(0,0).text='原表格'; table.cell(0,1).text='原事实98765'
        png = tmp_path/'original.png'; _tiny_png(png); d.add_picture(str(png))
        d.add_comment(d.paragraphs[1].runs[0], text='原有核对批注完整保留', author='原作者'); d.save(source)
        originals = {p:p.read_bytes() for p in (root/'body').glob('*.docx')}
        first = export(tid); assert 'result' in first, first
        row, _, word = verify_downloads(tid)
        previous = word.read_bytes()
        exported = Document(word)
        assert all(name in exported.styles for name in ('TOC Heading','toc 1','toc 2','toc 3'))
        assert len(exported.tables) == 1 and len(exported.inline_shapes) == 1
        assert [c.text for c in exported.comments] == ['原有核对批注完整保留']
        evidence.write_text('正文及标题未作特别排版规定，沿用项目默认格式。' if cancelled else
                            '正文采用宋体，小四字，2倍行距。一级标题采用黑体，三号字。')
        assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'stale'
        inspected = docx_format_rule()
        assert inspected['status'] == 'stale' and {i['target'] for i in inspected['issues']} == {'body','heading1'}
        failed = docx_assemble_volume(); assert '格式依据已变化' in failed
        assert word.read_bytes() == previous
        reason = '已逐项回读当前招标原文，原排版要求已明确取消'
        if cancelled:
            with pytest.raises(ValueError, match='核对理由'):
                docx_format_rule('body', reset=True, evidence_path='parse/tender/format.md', evidence_quote=evidence.read_text())
            partial = docx_format_rule('body', reset=True, reason=reason, evidence_path='parse/tender/format.md', evidence_quote=evidence.read_text())
        else:
            partial = docx_format_rule('body', reset=True, reason='回读当前招标原文，按最新明确字号和行距重新设置', font='宋体', font_size=12, line_spacing=2,
                                      evidence_path='parse/tender/format.md', evidence_quote=evidence.read_text())
        assert partial['status'] == 'stale' and {i['target'] for i in partial['issues']} == {'heading1'}
        docx_format_rule('heading1', reset=True, reason=reason if cancelled else '回读当前原文，一级标题已经明确改为三号黑体',
                         **({} if cancelled else {'font':'黑体','font_size':16}), evidence_path='parse/tender/format.md', evidence_quote=evidence.read_text())
        assert docx_format_rule()['status'] == 'ready'
        second = export(tid); assert 'result' in second, second
        updated, _, word = verify_downloads(tid)
        assert updated['content_seq'] == row['content_seq']+1
        assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'
        final = Document(word)
        p = next(p for p in final.paragraphs if p.text == parent+'原始正文与具体事实12345。')
        assert p.runs[0].font.size == Pt(12) and p.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia')) == '宋体'
        assert p.paragraph_format.line_spacing == (1.5 if cancelled else 2)
        assert final.tables[0].cell(0,1).text == '原事实98765'
        assert len(final.inline_shapes) == 1 and [c.text for c in final.comments] == ['原有核对批注完整保留']
        assert digest(template_path) == template_hash
        assert all(p.read_bytes() == data for p,data in originals.items())
        assert build_state.read(tid)['document_format_reviews']
    finally:
        runctx.clear_run()


def test_assembly_replacement_titles_and_old_toc_keep_every_comment_anchor(tmp_path, monkeypatch):
    tid, root = setup(tmp_path, monkeypatch)
    try:
        seed(tid,[node('封面'),node('目录'),node('管理方案'),node('附件证明')])
        expected=[]
        for name in ['封面','目录','管理方案','附件证明']:
            d=Document();heading=d.add_heading(name,1)
            d.add_paragraph('管理方案' if name=='目录' else name+'原始正文。')
            title_note=name+'原标题批注'
            d.add_comment(heading.runs[0],text=title_note,author='原作者');expected.append(title_note)
            if name in ('目录','管理方案'):
                cross=name+'跨原标题及正文的批注'
                d.add_comment((heading.runs[0],d.paragraphs[1].runs[0]),text=cross,author='原作者');expected.append(cross)
            d.save(root/('body/'+name+'.docx'))
        before={p:p.read_bytes() for p in (root/'body').glob('*.docx')}
        result=docx_assemble_volume();assert '已发布为整本' in result,result
        d=Document(root/'body/整本-格式验证册.docx')
        assert sorted(c.text for c in d.comments)==sorted(expected)
        starts=[e.get(qn('w:id')) for e in d.element.body.iter(qn('w:commentRangeStart'))]
        ends=[e.get(qn('w:id')) for e in d.element.body.iter(qn('w:commentRangeEnd'))]
        refs=[e.get(qn('w:id')) for e in d.element.body.iter(qn('w:commentReference'))]
        assert sorted(starts)==sorted(ends)==sorted(refs)==sorted(str(c.comment_id) for c in d.comments)
        assert all(p.read_bytes()==b for p,b in before.items())
    finally:
        runctx.clear_run()
