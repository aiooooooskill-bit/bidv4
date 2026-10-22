"""Real contracts for parent declarations, score classification and unchanged evidence gates."""
import pytest
from docx import Document
from app import artifact_store
from app.services import build_state, review_batches
from app.tools import body_contract
from app.tools.docx_ops import docx_section_create, docx_assemble_volume
from app.tools.validate_body import _validate_guide, validate_docx_file
from app.tools.validate_analysis import validate_analysis, score_value_issue
from tests.test_docx_ops import env, _seed_dir_artifact
from tests.test_review_flow import setup as review_setup, fixture_candidate, tender_citations
from tests.test_validate_analysis import _all_compliant, _seed


def declaration_directory(mode='模板或附件填充'):
    return {'response_documents': [{'name': '响应文件', 'directory': [
        {'目录名称': '响应声明', '交付形态': mode, '来源位置': ['TPL-01'],
         'children': [{'目录名称': '附件证明', '交付形态': '模板或附件填充'}]},
        {'目录名称': '纯容器', '交付形态': '目录容器',
         'children': [{'目录名称': '服务方案', '交付形态': '正文编写'}]}]}],
        'registry': {'TPL-01': {'type': '模板', 'text': '响应声明全文', '出处': '招标.docx L1-L3'}}}


@pytest.mark.parametrize('mode', ['模板或附件填充', '目录容器', ''])
def test_first_build_includes_template_parent_and_guide_requires_it(env, mode):
    tid = env['task']['id']
    _seed_dir_artifact(env, declaration_directory(mode))
    assert [s['title'] for s in build_state.sections(tid)] == ['响应声明', '附件证明', '服务方案']
    guide = '| 节 | 模式 | 依据 | 素材 | 缺口/备注 |\n|---|---|---|---|---|\n| 附件证明 | — | TPL-01 | — | 缺证明 |\n| 服务方案 | 推理撰写 | — | — | — |\n'
    issues, _, _ = _validate_guide(guide, tid)
    assert any('指引缺行：「响应声明」' in i for i in issues)
    guide += '| 响应声明 | 格式跟随 | TPL-01 | — | 按原格式完整填报 |\n'
    assert _validate_guide(guide, tid)[0] == []


def test_missing_parent_cannot_publish_volume_then_real_declaration_is_merged(env):
    tid = env['task']['id']; _seed_dir_artifact(env, declaration_directory())
    docx_section_create(path='body/附件证明', title='附件证明', paragraphs='（此处贴入证明）')
    docx_section_create(path='body/服务方案', title='服务方案', paragraphs='按本项目要求提供服务。')
    result = docx_assemble_volume()
    assert '响应声明' in result and '未发布整册' in result
    docx_section_create(path='body/响应声明', title='响应声明', paragraphs='本响应有效期为九十日，保证所附资料真实有效。')
    docx_assemble_volume()
    volume = Document(artifact_store.work_dir(tid) / 'body/整本-响应文件.docx')
    text = '\n'.join(p.text for p in volume.paragraphs)
    assert '本响应有效期为九十日' in text
    assert text.index('本响应有效期为九十日') < text.rindex('附件证明')


def test_title_only_is_not_declaration_content(tmp_path):
    p = tmp_path / '响应声明.docx'; d = Document(); d.add_heading('响应声明', 1); d.save(p)
    assert any('正文仅有节标题' in i for i in validate_docx_file(p)[0])
    d.add_paragraph('保证所附资料真实有效。'); d.save(p)
    assert not validate_docx_file(p)[0]


@pytest.mark.parametrize('value', ['未知（原文未设分值，未给出评分表）', '通过制（不能证明合理性按无效响应处理）'])
def test_analysis_rejects_non_scoring_rules_without_dropping_price_obligations(env, value):
    products = _all_compliant()
    products['evaluation'] = products['evaluation'].replace('| 技术方案 | 30 |', f'| 价格审查 | {value} |')
    _seed(env['task'], products)
    assert validate_analysis().startswith('[校验未通过]')
    products['evaluation'] = products['evaluation'].replace(f'| 价格审查 | {value} | 完整性 | 招标文件.docx 评标办法 附表（L171-L172） |', '')
    products['evaluation'] += '\n## 评标办法概述\n价格审查按原文处理，不能证明合理性按无效响应处理，依据招标文件.docx L171-L172。\n'
    _seed(env['task'], products)
    assert validate_analysis().startswith('[校验通过]')
    assert '价格审查按原文处理' in (artifact_store.work_dir(env['task']['id']) / 'analysis/evaluation.md').read_text()


@pytest.mark.parametrize('value', ['30', '8~12', '优/良/中/差', '通过制（满足得基础分）', '通过制（满足得5分，不满足无得分）', '通过制（满足得5分，不满足无分值）', '通过制（不通过扣5分）', '未知（原件分值不可读）'])
def test_original_scoring_types_remain_available(value):
    assert score_value_issue(value) is None


def test_no_score_catalog_rejects_scoring_issue_but_keeps_uncertainty(review_setup):
    tid, value, *_ = review_setup
    batch = review_batches.next_batch(tid, 'review-test')
    problem = dict(risk_type='scoring', title='服务期限待确认', explanation='实际期限不明。', suggestion='核对期限。')
    rows = [dict(slot=1, tender=tender_citations(tid), bid=value['requirements'][0]['bid'],
                 status='uncertain', assessment='期限须确认。', problems=[problem])]
    with pytest.raises(ValueError, match='无评分项'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], rows)
    problem['risk_type'] = 'confirmation'
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], rows)
    assert proposed['comparisons'][0]['problems'][0]['risk_type'] == 'confirmation'
