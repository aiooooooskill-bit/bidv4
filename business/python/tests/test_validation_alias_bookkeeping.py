import pytest
from tests.test_section_write_atomicity import env, request
from app.services import build_state

@pytest.mark.parametrize('alias', ['body/实施方案.docx', './body/实施方案.docx', '.\\body\\实施方案.docx', 'body/../body/实施方案.docx', 'body//实施方案.docx'])
def test_success_uses_same_canonical_section_for_validation_and_ledger(env, alias):
    tid, root, _ = env
    code, result = request(env)
    assert code == 0, result
    assert 'body/实施方案.docx' in build_state.read(tid)['pending_sections']
    code, result = request(env, 'validate_body', {'section': alias})
    assert code == 0 and result['result'].startswith('[校验通过]'), result
    value = build_state.read(tid)
    assert 'body/实施方案.docx' not in value.get('pending_sections', {})
    assert 'body/实施方案.docx' in value['sections']
    assert all(key == 'body/实施方案.docx' for key in value['sections'])
    # Successful validation replay does not rewrite the ledger.
    before = (root / 'body/.build-state.json').read_bytes()
    request(env, 'validate_body', {'section': alias})
    assert (root / 'body/.build-state.json').read_bytes() == before

@pytest.mark.parametrize('target', ['body', '.\\body\\不存在.docx', '../body/实施方案.docx', '/body/实施方案.docx'])
def test_global_missing_or_invalid_path_never_marks_pending_section_ready(env, target):
    tid, root, _ = env
    before = (root / 'body/.build-state.json').read_bytes()
    request(env, 'validate_body', {'section': target})
    assert (root / 'body/.build-state.json').read_bytes() == before
    assert 'body/实施方案.docx' in build_state.read(tid)['pending_sections']


def test_failed_alias_validation_never_marks_invalid_docx_ready(env):
    tid, root, _ = env
    code, result = request(env, 'docx_section_create', {'path': 'body/实施方案.docx', 'title': '实施方案', 'paragraphs': '联系人：【待补】', 'replace': True})
    assert code == 0, result
    before = (root / 'body/.build-state.json').read_bytes()
    code, result = request(env, 'validate_body', {'section': '.\\body\\实施方案.docx'})
    assert code == 0 and not result['result'].startswith('[校验通过]'), result
    assert (root / 'body/.build-state.json').read_bytes() == before


def test_validation_of_untracked_file_cannot_complete_another_section(env):
    from docx import Document
    tid, root, _ = env
    other = root / 'body/其他.docx'
    doc = Document(); doc.add_paragraph('另一个现有章节'); doc.save(other)
    before = (root / 'body/.build-state.json').read_bytes()
    code, result = request(env, 'validate_body', {'section': './body/其他.docx'})
    assert code == 0 and result['result'].startswith('[校验通过]'), result
    assert (root / 'body/.build-state.json').read_bytes() == before


def test_changed_pending_hash_cannot_be_promoted_by_alias(env):
    from docx import Document
    tid, root, _ = env
    path = root / 'body/实施方案.docx'
    doc = Document(path); doc.add_paragraph('在记账后由其他入口更改'); doc.save(path)
    before = (root / 'body/.build-state.json').read_bytes()
    code, result = request(env, 'validate_body', {'section': './body/实施方案.docx'})
    assert code == 0 and result['result'].startswith('[校验通过]'), result
    assert (root / 'body/.build-state.json').read_bytes() == before


def test_same_title_in_two_volumes_only_completes_selected_canonical_path(env):
    from app import publish
    tid, root, _ = env
    publish.publish_artifact('tender.directory/tender-response-docs@1', {'response_documents': [
        {'name': name, 'directory': [{'目录名称': '实施方案', '交付形态': '正文编写'}]}
        for name in ['技术册', '商务册']]}, task_id=tid)
    for name in ['技术册', '商务册']:
        code, result = request(env, 'docx_section_create', {'path': f'body/{name}/实施方案.docx', 'title': '实施方案', 'paragraphs': f'{name}的独立方案'})
        assert code == 0, result
    code, result = request(env, 'validate_body', {'section': r'.\body\技术册\实施方案.docx'})
    assert code == 0 and result['result'].startswith('[校验通过]'), result
    pending = build_state.read(tid)['pending_sections']
    assert 'body/技术册/实施方案.docx' not in pending
    assert 'body/商务册/实施方案.docx' in pending
