import asyncio
import json
from pathlib import Path
import zipfile
import pytest
from openpyxl import Workbook
from app import artifact_store, db, document
from app.parse import convert, sha256_file
from app.services import files, page_flow, projects, review
from app.services.errors import IncomingFile, BusinessError


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.setenv('DATA_DIR', str(tmp_path / 'data'))
    db.init_db()
    tid = db.create_task('原项目')['id']
    artifact_store.ensure_task_skeleton(tid)
    return tid


def upload(tid, path, name=None, **options):
    incoming = IncomingFile(path, name or path.name)
    try:
        return asyncio.run(files.upload_file(tid, incoming, **options))
    finally:
        asyncio.run(incoming.close())


def test_xlsx_values_formulas_hidden_sheets_and_short_workbooks(tmp_path):
    path = tmp_path / '清单.xlsx'
    book = Workbook();sheet = book.active;sheet.title = '采购清单'
    sheet.append(['编码', '数量', '说明', '日期', '比例'])
    from datetime import date
    sheet.append([123, 0, '含|竖线\n换行', date(2026, 10, 26), 0.125])
    sheet['A2'].number_format = '00000';sheet['E2'].number_format = '0.0%'
    sheet['F2'] = '=B2+1';sheet['G2'] = '=B2+2'
    sheet.merge_cells('A4:C4');sheet['A4'] = '合并范围条款'
    hidden = book.create_sheet('隐藏补充');hidden.sheet_state = 'hidden';hidden['A1'] = '隐藏条款'
    book.save(path)
    with zipfile.ZipFile(path) as z:
        content = {name:z.read(name) for name in z.namelist()}
    content['xl/worksheets/sheet1.xml'] = content['xl/worksheets/sheet1.xml'].replace(b'<f>B2+1</f><v></v>', b'<f>B2+1</f><v>1</v>')
    with zipfile.ZipFile(path, 'w') as z:
        for name, value in content.items():z.writestr(name, value)
    digest = sha256_file(path)
    result = convert(path)
    assert '00123' in result.md and '| 0 |' in result.md and '12.5%' in result.md
    assert '2026-10-26' in result.md and '含\\|竖线<br>换行' in result.md
    assert '合并范围条款' in result.md and '隐藏条款' in result.md
    assert result.info['warnings'][0].startswith('采购清单!G2')
    assert 'F2' not in '\n'.join(result.info['warnings'])
    assert sha256_file(path) == digest
    response = document.prepare(path, tmp_path / 'parse')
    assert response['meta']['conversion'] == 'excel-values'
    assert document.prepare(path, tmp_path / 'parse')['cached']
    tiny = Workbook();tiny.active['A1'] = '小型补充';tiny.save(tmp_path / '小表.xlsx')
    assert document.prepare(tmp_path / '小表.xlsx', tmp_path / 'tiny')['meta']['tables'] == 1
    empty = Workbook();empty.save(tmp_path / '空表.xlsx')
    with pytest.raises(ValueError, match='没有可读取'):convert(tmp_path / '空表.xlsx')


def test_xls_real_biff_values():
    source = Path(__file__).parents[3] / 'tests/fixtures/excel/source.xls'
    result = convert(source)
    assert '00123' in result.md and '2026-10-26' in result.md and '| 0 |' in result.md


@pytest.mark.parametrize('epoch', [1900, 1904])
def test_xls_clock_duration_and_calendar_values(epoch):
    from app.parse.excel import read_workbook
    source = Path(__file__).parents[3] / f'tests/fixtures/excel/time-values-{epoch}.xls'
    sheet = read_workbook(source)[0]
    values = {row['cells'][0]: row['cells'][1] for row in sheet['rows']}
    assert values == {'开始时间': '12:00:00', '含整数天的钟点': '12:00:00',
                      '响应时长': '0:30:00', '跨天时长': '1 day, 12:00:00',
                      '日期': '2026-10-27', '日期时间': '2026-10-27T12:30:00', '仅月份': '2026-10-27'}
    result = convert(source)
    assert '1899-' not in result.md and '1904-' not in result.md
    assert '| 开始时间 | 12:00:00 |' in result.md and '| 响应时长 | 0:30:00 |' in result.md
    assert not sheet['warnings']


@pytest.mark.parametrize('epoch', [1900, 1904])
def test_xls_text_colour_and_multiple_format_sections_do_not_change_values(epoch, project, tmp_path):
    from app.parse.excel import read_workbook
    source = Path(__file__).parents[3] / f'tests/fixtures/excel/formatted-time-values-{epoch}.xls'
    expected = {
        '月份文字标签': '2026-10-27', '日期文字标签': '2026-10-27',
        '钟点文字标签': '12:00:00', '文字中的时长标记': '12:00:00',
        '文字中的分号': '12:00:00', '转义年份字母': '12:00:00',
        '转义小时字母': '2026-10-27', '转义分号': '12:00:00',
        '留白与填充': '12:00:00', '红色钟点': '12:00:00',
        '索引颜色钟点': '12:00:00', '大写钟点': '12:00:00',
        '区域钟点': '12:00:00', '多段钟点': '12:00:00',
        '不同类型的后续段': '12:00:00', '多段日期': '2026-10-27',
        '多段月份': '2026-10-27', '多段时长': '1 day, 12:00:00',
        '彩色分钟时长': '0:30:00', '带标签的跨天时长': '1 day, 12:00:00',
        '说明': '保留的文字与数字123', '数量': '0',
    }
    digest = sha256_file(source)
    sheet = read_workbook(source)[0]
    assert {row['cells'][0]: row['cells'][1] for row in sheet['rows']} == expected
    assert not sheet['warnings']
    result = convert(source)
    for label, value in expected.items():
        assert f'| {label} | {value} |' in result.md
    assert '1899-' not in result.md and '1904-' not in result.md
    prepared = document.prepare(source, tmp_path / 'formatted-parse')
    assert prepared['meta']['conversion'] == 'excel-values'
    # Unchanged originals must not keep the previous parser's incorrect values.
    parsed, _, meta_path = map(Path, prepared['paths'])
    previous_meta = json.loads(meta_path.read_text())
    previous_meta.pop('text_view')
    meta_path.write_text(json.dumps(previous_meta))
    parsed.write_text('| 月份文字标签 | 00:00:00 |\n| 红色钟点 | 1899-12-31T12:00:00 |')
    refreshed = document.prepare(source, tmp_path / 'formatted-parse')
    assert not refreshed.get('cached')
    assert parsed.read_text() == result.md
    assert document.prepare(source, tmp_path / 'formatted-parse')['cached']
    upload(project, source)
    preview = asyncio.run(files.preview_excel(source.name, project))
    assert preview['rows'] == sheet['rows']
    assert sha256_file(source) == digest


def test_xlsx_clock_and_duration_match_legacy_values(tmp_path):
    from datetime import time, timedelta
    from app.parse.excel import read_workbook
    source = tmp_path / '时间.xlsx'
    book = Workbook();sheet = book.active
    sheet.append(['开始时间', time(12)]);sheet['B1'].number_format = 'hh:mm'
    sheet.append(['响应时长', timedelta(minutes=30)]);sheet['B2'].number_format = '[h]:mm'
    sheet.append(['跨天时长', timedelta(hours=36)]);sheet['B3'].number_format = '[h]:mm:ss'
    book.save(source)
    assert [row['cells'][1] for row in read_workbook(source)[0]['rows']] == ['12:00:00', '0:30:00', '1 day, 12:00:00']


def test_file_lifecycle_default_sources_replacement_and_removal(project, tmp_path):
    tid = project
    source = tmp_path / '主文件.pdf';source.write_bytes(b'original')
    upload(tid, source)
    source = tmp_path / '补充.xlsx';book = Workbook();book.active['A1'] = '补充要求';book.save(source)
    upload(tid, source)
    assert page_flow.source_selection(tid)['supplements'] == [{'file':'补充.xlsx','role':'补充文件'}]
    old_version = page_flow.input_versions(tid)['analyse']
    preview = asyncio.run(files.preview_excel('补充.xlsx', tid))
    assert preview['rows'][0] == {'number':1,'cells':['补充要求']}
    page_flow.set_sources(tid, '主文件.pdf', [])
    upload(tid, source, replace_name='补充.xlsx')
    assert not page_flow.source_selection(tid)['supplements'], 'Replacing an excluded file must preserve exclusion'
    page_flow.set_sources(tid, '主文件.pdf', ['补充.xlsx'])
    newer = tmp_path / '新主文件.pdf';newer.write_bytes(b'new main')
    upload(tid, newer, replace_name='主文件.pdf')
    assert page_flow.source_selection(tid)['main'] == '新主文件.pdf'
    assert not (artifact_store.sources_dir(tid) / '主文件.pdf').exists()
    assert page_flow.input_versions(tid)['analyse'] != old_version
    asyncio.run(files.delete_file('新主文件.pdf', tid))
    assert page_flow.source_selection(tid)['main'] == ''
    assert page_flow.source_selection(tid)['supplements'][0]['file'] == '补充.xlsx'
    upload(tid, newer)
    assert page_flow.source_selection(tid)['main'] == '', 'Never silently promote a replacement when other tender sources remain'
    asyncio.run(files.delete_file('补充.xlsx', tid))
    asyncio.run(files.delete_file('新主文件.pdf', tid))
    assert page_flow.source_selection(tid) == {'main':'','supplements':[],'excluded':[]}
    upload(tid, newer)
    assert page_flow.source_selection(tid)['main'] == '新主文件.pdf'


def test_failed_upload_preserves_original_and_bid_files_do_not_join_sources(project, tmp_path, monkeypatch):
    tid = project;source = tmp_path / '招标.pdf';source.write_bytes(b'old');upload(tid, source)
    original = page_flow.source_selection(tid)
    source.write_bytes(b'x' * 2048);monkeypatch.setattr(files, 'MAX_UPLOAD_BYTES', 1024)
    with pytest.raises(BusinessError, match='100MB'):upload(tid, source, replace_name='招标.pdf')
    assert (artifact_store.sources_dir(tid) / '招标.pdf').read_bytes() == b'old'
    assert page_flow.source_selection(tid) == original
    upload(tid, source, name='受审.pdf', purpose='bid')
    assert page_flow.source_selection(tid) == original
    review.set_input(tid, ['受审.pdf'])
    source.write_bytes(b'small file')
    with pytest.raises(BusinessError, match='不同名称'):upload(tid, source, name='受审.pdf')
    asyncio.run(files.delete_file('受审.pdf', tid))
    assert review.selection(tid)['bids'] == []
    assert not list(artifact_store.sources_dir(tid).glob('.*.part'))


def test_rename_metadata_only_with_concurrent_change_guard(project, tmp_path):
    tid = project;source = tmp_path / '招标.pdf';source.write_bytes(b'original');upload(tid, source)
    versions = page_flow.input_versions(tid)
    result = asyncio.run(projects.rename_project(tid, projects.RenameBody(title=' 手工修改的项目 ', previous_title='原项目')))
    assert result['title'] == '手工修改的项目'
    assert page_flow.input_versions(tid) == versions
    assert (artifact_store.sources_dir(tid) / '招标.pdf').read_bytes() == b'original'
    with pytest.raises(BusinessError, match='已变化'):
        asyncio.run(projects.rename_project(tid, projects.RenameBody(title='旧页覆盖', previous_title='原项目')))
