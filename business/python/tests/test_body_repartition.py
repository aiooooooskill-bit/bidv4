import hashlib
import json
from collections import Counter
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from app import artifact_store, db, publish, runctx
from app.services import build_state
from app.tools.body_repartition import docx_section_repartition, _blocks
from app.tools.docx_ops import docx_assemble_volume
from app.tools.validate_body import validate_body
from tests.test_docx_ops import _tiny_png
from tests.util import init_env


@pytest.fixture
def env(tmp_path, monkeypatch):
    task, conv = init_env(tmp_path, monkeypatch)
    tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    root = artifact_store.work_dir(tid)
    runctx.set_run(conv['id'], 'body-conservation', tid)
    nodes = [{'目录名称': '实施方案', 'level': 1, 'children': [
        {'目录名称': '接口对接', 'level': 2, 'children': []},
        {'目录名称': '故障处置', 'level': 2, 'children': []}]}]
    content = {'numbering': 'decimal', 'response_documents': [{'name': '技术标', 'directory': nodes}]}
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    source = root / 'body/实施方案.docx'
    doc = Document()
    doc.add_heading('实施方案', 1)                         # B1
    doc.add_paragraph('原有总说明完整保留。')                # B2
    doc.add_heading('接口对接', 2)                         # B3
    doc.add_paragraph('原有接口说明，服务期三年。')           # B4
    table = doc.add_table(rows=2, cols=2)                   # B5
    table.cell(0, 0).merge(table.cell(0, 1)).text = '原有合并表头'
    table.cell(1, 0).text = '鉴权方式'; table.cell(1, 1).text = '双向认证'
    doc.add_heading('故障处置', 2)                         # B6
    p = doc.add_paragraph('原有故障说明。')                  # B7
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to('https://example.org/original-evidence', RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '原始链接'; r.append(t); link.append(r); p._p.append(link)
    png = tmp_path / 'figure.png'; _tiny_png(png)
    doc.add_picture(str(png))                              # B8
    doc.add_comment(p.runs[0], text='原有核对批注', author='原作者')
    doc.save(source)
    yield tid, root, content, source
    runctx.clear_run()


def assignments():
    return [{'title': '实施方案', 'ranges': [[1, 2]]},
            {'title': '接口对接', 'ranges': [[3, 5]]},
            {'title': '故障处置', 'ranges': [[6, 8]]}]


def test_split_and_real_assembly_preserve_all_blocks_tables_images_links_and_comments(env):
    tid, root, _, source = env
    before = source.read_bytes()
    view = json.loads(docx_section_repartition('body/实施方案.docx'))
    assert view['block_count'] == 8 and view['blocks'][4]['kind'] == 'tbl'
    result = json.loads(docx_section_repartition('body/实施方案.docx', json.dumps(assignments()), view['sha256']))
    assert result['preserved_blocks'] == 8
    assert (root / result['restore_point']).read_bytes() == before
    expected = Counter(etree.tostring(e) for e in _blocks(Document(BytesIO(before))))
    actual = Counter(etree.tostring(e) for item in result['outputs'] for e in _blocks(Document(root / item['path'])))
    assert actual == expected
    assert build_state.snapshot(tid)['orphan_body_files'] == []
    assert all(s['status'] == 'draft' for s in build_state.snapshot(tid)['body_sections'])
    assert '未发布整册' in docx_assemble_volume()
    assert not [r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume']
    for item in result['outputs']:
        assert validate_body(item['path']).startswith('[校验通过]')
        build_state.record_validation(tid, item['path'])
    assert '合并 3 节' in docx_assemble_volume()
    row = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    volume = Document(artifact_store.package_files(row['artifact_id'], row)[0])
    text = '\n'.join(p.text for p in volume.paragraphs)
    for sentence in ['原有总说明完整保留。', '原有接口说明，服务期三年。', '原有故障说明。']:
        assert text.count(sentence) == 1
    assert len(volume.tables) == 1
    assert volume.tables[0].cell(0, 0)._tc is volume.tables[0].cell(0, 1)._tc
    assert volume.tables[0].cell(1, 1).text == '双向认证'
    assert len(volume.inline_shapes) == 1
    link = next(volume.element.body.iter(qn('w:hyperlink')))
    assert volume.part.rels[link.get(qn('r:id'))].target_ref == 'https://example.org/original-evidence'
    assert [c.text for c in volume.comments] == ['原有核对批注']
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'
    changed = Document(source); changed.add_paragraph('父章节补充。'); changed.save(source)
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'stale'


def test_existing_parent_body_is_consumed_even_without_repartition(env):
    tid, root, _, source = env
    before = source.read_bytes()
    for title in ['接口对接', '故障处置']:
        d = Document(); d.add_heading(title, 1); d.add_paragraph(title + '新增补充。'); d.save(root / ('body/' + title + '.docx'))
    assert '实施方案' in [s['title'] for s in build_state.sections(tid)]
    assert '合并 3 节' in docx_assemble_volume()
    volume = Document(root / 'body/整本-技术标.docx')
    assert source.read_bytes() == before
    assert '\n'.join(p.text for p in volume.paragraphs).count('原有总说明完整保留。') == 1
    assert len(volume.tables) == 1 and len(volume.inline_shapes) == 1


@pytest.mark.parametrize('bad', [
    [{'title': '实施方案', 'ranges': [[1, 7]]}],
    [{'title': '实施方案', 'ranges': [[1, 8]]}, {'title': '故障处置', 'ranges': [[8, 8]]}],
    [{'title': '其他章节', 'ranges': [[1, 8]]}],
    [{'title': '接口对接', 'ranges': [[0, 8]]}],
])
def test_rejects_loss_duplication_unrelated_targets_and_bad_ranges_without_writes(env, bad):
    _, root, _, source = env
    original = source.read_bytes()
    with pytest.raises(ValueError):
        docx_section_repartition('body/实施方案.docx', json.dumps(bad), hashlib.sha256(original).hexdigest())
    assert source.read_bytes() == original
    assert sorted(p.name for p in (root / 'body').glob('*.docx')) == ['实施方案.docx']


def test_rejects_stale_read_and_existing_target(env):
    _, root, _, source = env
    before = source.read_bytes()
    with pytest.raises(ValueError, match='已变化'):
        docx_section_repartition('body/实施方案.docx', json.dumps(assignments()), 'old')
    target = root / 'body/故障处置.docx'; target.write_bytes(before)
    with pytest.raises(ValueError, match='已有正文'):
        docx_section_repartition('body/实施方案.docx', json.dumps(assignments()), hashlib.sha256(before).hexdigest())
    assert source.read_bytes() == before and target.read_bytes() == before


def test_failed_save_rolls_back_every_destination_and_build_state(env, monkeypatch):
    tid, root, _, source = env
    from app.tools import docx_ops
    original, ledger = source.read_bytes(), build_state.read(tid)
    real_save = docx_ops._atomic_save
    def fail_on_second(doc, target):
        if target.name == '接口对接.docx':
            raise OSError('simulated full disk')
        real_save(doc, target)
    monkeypatch.setattr(docx_ops, '_atomic_save', fail_on_second)
    with pytest.raises(OSError, match='full disk'):
        docx_section_repartition('body/实施方案.docx', json.dumps(assignments()), hashlib.sha256(original).hexdigest())
    assert source.read_bytes() == original
    assert not (root / 'body/接口对接.docx').exists() and not (root / 'body/故障处置.docx').exists()
    assert build_state.read(tid) == ledger


def test_unassigned_original_body_blocks_publication_and_remains_intact(env):
    tid, root, content, source = env
    original = source.read_bytes()
    content['response_documents'][0]['directory'] = [{'目录名称': '新章节', 'level': 1, 'children': []}]
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    d = Document(); d.add_heading('新章节', 1); d.add_paragraph('新的内容。'); d.save(root / 'body/新章节.docx')
    result = docx_assemble_volume()
    assert '未归集的既有正文：body/实施方案.docx' in result
    assert '未发布整册' in result
    assert not [r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume']
    assert source.read_bytes() == original


def test_missing_owned_parent_cannot_disappear_from_required_sections(env):
    tid, root, _, source = env
    build_state.record_section(tid, 'body/实施方案.docx')
    for title in ['接口对接', '故障处置']:
        d = Document(); d.add_heading(title, 1); d.add_paragraph(title + '独立正文。'); d.save(root / ('body/' + title + '.docx'))
    assert '未发布整册' not in docx_assemble_volume()
    row = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    published = artifact_store.package_files(row['artifact_id'], row)[0].read_bytes()
    source.unlink()
    assert any(s['title'] == '实施方案' and s['status'] == 'missing' for s in build_state.snapshot(tid)['body_sections'])
    result = docx_assemble_volume()
    assert '缺失 1 节未并入：实施方案' in result and '未发布整册' in result
    assert artifact_store.package_files(row['artifact_id'], row)[0].read_bytes() == published


def test_full_repartition_releases_only_verified_original_ownership(env):
    tid, root, _, source = env
    build_state.record_section(tid, 'body/实施方案.docx')
    before = source.read_bytes()
    assigned = [{'title': '接口对接', 'ranges': [[1, 5]]}, {'title': '故障处置', 'ranges': [[6, 8]]}]
    result = json.loads(docx_section_repartition('body/实施方案.docx', json.dumps(assigned), hashlib.sha256(before).hexdigest()))
    assert not source.exists() and (root / result['restore_point']).read_bytes() == before
    assert {s['title'] for s in build_state.sections(tid)} == {'接口对接', '故障处置'}
    assert result['preserved_blocks'] == 8


@pytest.mark.parametrize('marker', ['comment', 'field'])
def test_cross_paragraph_ranges_cannot_be_torn_across_destinations(env, marker):
    _, _, _, source = env
    d = Document(source)
    if marker == 'comment':
        d.add_comment((d.paragraphs[1].runs[0], d.paragraphs[5].runs[0]), text='跨段原批注', author='审阅人')
    else:
        for p, kind in [(d.paragraphs[1], 'begin'), (d.paragraphs[3], 'end')]:
            field = OxmlElement('w:fldChar'); field.set(qn('w:fldCharType'), kind); p.add_run()._r.append(field)
    d.save(source);before = source.read_bytes()
    with pytest.raises(ValueError, match='不能跨目标拆开'):
        docx_section_repartition('body/实施方案.docx', json.dumps(assignments()), hashlib.sha256(before).hexdigest())
    assert source.read_bytes() == before


def test_vml_preview_and_embedded_object_parts_are_preserved_without_bitmap_reencoding(env):
    from docx.opc.part import Part
    from app.tools.docx_ops import _migrate_images
    _, root, _, source = env
    d = Document(source)
    preview = d.part.get_or_add_image(BytesIO(d.part.related_parts[next(d.element.body.iter(qn('a:blip'))).get(qn('r:embed'))].blob))[0]
    blob = b'opaque embedded object bytes retained exactly'
    package = d.part.package
    embedded = Part(package.next_partname('/word/embeddings/oleObject%d.bin'), 'application/vnd.openxmlformats-officedocument.oleObject', blob, package)
    rid = d.part.relate_to(embedded, RT.OLE_OBJECT)
    obj = OxmlElement('w:object')
    shape = etree.SubElement(obj, '{urn:schemas-microsoft-com:vml}shape')
    image = etree.SubElement(shape, '{urn:schemas-microsoft-com:vml}imagedata'); image.set(qn('r:id'), preview)
    ole = etree.SubElement(obj, '{urn:schemas-microsoft-com:office:office}OLEObject'); ole.set(qn('r:id'), rid)
    target = Document(); copied = deepcopy(obj)
    assert _migrate_images(d, target, copied) == 1
    target.add_paragraph().add_run()._r.append(copied)
    path = root / 'object-preservation.docx'; target.save(path); reopened = Document(path)
    mapped = next(reopened.element.body.iter('{urn:schemas-microsoft-com:office:office}OLEObject')).get(qn('r:id'))
    assert reopened.part.related_parts[mapped].blob == blob
    mapped_preview = next(reopened.element.body.iter('{urn:schemas-microsoft-com:vml}imagedata')).get(qn('r:id'))
    assert reopened.part.related_parts[mapped_preview].blob == d.part.related_parts[preview].blob


def test_missing_image_relationship_stops_assembly_instead_of_publishing_loss(env):
    tid, root, _, source = env
    for title in ['接口对接', '故障处置']:
        d = Document(); d.add_heading(title, 1); d.add_paragraph(title + '已写正文。'); d.save(root / ('body/' + title + '.docx'))
    assert '未发布整册' not in docx_assemble_volume()
    row = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    package = artifact_store.package_files(row['artifact_id'], row)[0]
    previous = package.read_bytes()
    damaged = Document(source);next(damaged.element.body.iter(qn('a:blip'))).set(qn('r:embed'), 'rIdMissing');damaged.save(source)
    assert docx_assemble_volume().startswith('[合册失败]')
    assert package.read_bytes() == previous
