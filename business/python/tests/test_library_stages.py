"""Recovery must retain real content without inventing facts or crossing companies."""
import asyncio
import json
import pytest
from docx import Document
from app import db, company_store as cs, runctx, document
from app.knowledge import store, ingest, images
from app.services import company_library as lib, knowledge, bidder
from app.tools.search_knowledge import search_company_assets
from app.tools.company_asset import read_company_asset

@pytest.fixture
def library(tmp_path,monkeypatch):
    monkeypatch.setenv('DATA_DIR',str(tmp_path));db.init_db();store.ensure_dirs()
    co=asyncio.run(lib.save_company(lib.CompanyBody(name='虚构星河科技',profile={'legal_name':'虚构法人'})))
    tid=db.create_task('虚构采购')['id']
    asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=co['id'],signer_id='legal',sign_date='2026-10-16')))
    runctx.set_run('stages','stages',tid)
    yield co,tid
    runctx.clear_run()

def insert(co,name):
    item=db.kb_insert_item(file_name=name,file_hash=name,title=name.rsplit('.',1)[0],ext='.'+name.rsplit('.',1)[1])
    cs.attach_asset(item['id'],co['id'])
    return item

def test_invalid_json_retains_parsed_text_index_and_retry_cache(library,monkeypatch):
    co,_=library;i=insert(co,'虚构财务报表.docx')
    path=store.kb_files_dir()/i['file_name'];doc=Document();doc.add_paragraph('虚构财报：星河公司年度净资产为九百万元。');doc.save(path)
    document.prepare(path,store.kb_parse_dir(i['file_name']));ingest.parsed(i['id'])
    with pytest.raises(ValueError):ingest.finish(i['id'],json.loads('{"doc_type":"financial_report" invalid}'))
    ingest.fail(i['id'],'extract','虚构 JSON 错误')
    detail=asyncio.run(knowledge.get_item(i['id']))
    assert detail['parse_status']=='ready' and detail['extract_status']=='failed'
    assert detail['availability']['usable'] and '净资产' in detail['content']
    result=search_company_assets('净资产');assert '九百万元' in result and 'AI整理失败' in result
    assert ingest.schedule_ingest(i['id'])
    monkeypatch.setattr(document,'convert',lambda *_:pytest.fail('Retry must reuse successfully parsed text'))
    assert document.prepare(path,store.kb_parse_dir(i['file_name']))['cached']
    ingest.parsed(i['id']);ingest.finish(i['id'],{'doc_type':'financial_report','fields':{}})
    assert asyncio.run(knowledge.get_item(i['id']))['error'] is None

def test_failed_raw_file_locator_manual_recovery_and_strict_type(library):
    co,_=library;i=insert(co,'虚构营业执照.png');ingest.fail(i['id'],'parse','模型不支持图片')
    assert asyncio.run(knowledge.list_items(q='营业执照',company_id=co['id']))['total']==1
    result=search_company_assets('营业执照')
    assert '仅原件' in result and '不是事实证据' in result and 'raw-locator' in result
    assert '尚未上传' not in result and '模型不支持图片' in result
    typed=search_company_assets('营业执照','business_license')
    assert '类型筛选无命中' in typed and '未采用其他类型' in typed
    item=db.kb_get_item(i['id'])
    empty=asyncio.run(lib.review_asset(i['id'],lib.ReviewBody(revision=item['revision'],company_id=co['id'],title=i['title'],business={'doc_type':'business_license'})))
    assert not empty['availability']['usable'] and empty['capability']=='stored'
    reviewed=asyncio.run(lib.review_asset(i['id'],lib.ReviewBody(revision=empty['revision'],company_id=co['id'],title=i['title'],business={'doc_type':'business_license','extra':{'虚构证号':{'value':'TEST2026','source':'人工校验样例'}}},text='人工核对的虚构许可证 TEST2026')))
    assert reviewed['availability']['source']=='human' and reviewed['parse_status']=='failed'
    assert reviewed['error']=='模型不支持图片' # accurate machine history, not fake OCR success
    assert 'TEST2026' in search_company_assets('TEST2026','business_license')
    read=json.loads(read_company_asset(i['id']));assert read['availability']['usable']
    other=asyncio.run(lib.save_company(lib.CompanyBody(name='虚构乙公司')))
    foreign=insert(other,'私有外国许可证.png');ingest.fail(foreign['id'],'parse','不得泄漏的错误')
    nohit=search_company_assets('XYZUNMATCHED')
    assert '不得泄漏' not in nohit and foreign['file_name'] not in nohit
    assert '不能据此断言未上传' in nohit

def test_image_failure_is_independent_and_success_clears_only_image_warning(library,monkeypatch):
    co,_=library;i=insert(co,'虚构证明.docx')
    doc=Document();doc.add_paragraph('虚构证明的正文始终保留');src=store.kb_files_dir()/i['file_name'];doc.save(src)
    document.prepare(src,store.kb_parse_dir(i['file_name']));ingest.parsed(i['id'])
    def fail(*_):raise ValueError('虚构图片渲染失败')
    monkeypatch.setattr(images,'_extract_docx',fail)
    assert '图片' in ingest.prepare_images(i['id'])['error']
    ingest.finish(i['id'],{'doc_type':'other'})
    detail=asyncio.run(knowledge.get_item(i['id']))
    assert detail['parse_status']=='ready' and detail['error'] is None
    assert detail['meta']['image_error'] and detail['availability']['usable']
    monkeypatch.setattr(images,'_extract_docx',lambda *_:[])
    assert ingest.prepare_images(i['id'])['error'] is None
    assert 'image_error' not in asyncio.run(knowledge.get_item(i['id']))['meta']

def test_raw_name_locator_survives_dense_factual_matches(library):
    co,_=library
    for n in range(40):
        item=insert(co,f'虚构合同{n}.docx')
        md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True,exist_ok=True)
        md.write_text('营业执照用于办理本合同；这是合同正文，不是营业执照原件。')
        outline.write_text('[]');ingest.parsed(item['id'])
        ingest.finish(item['id'],{'doc_type':'contract_case'})
    target=insert(co,'虚构营业执照.pdf');ingest.fail(target['id'],'parse','虚构 OCR 故障')
    result=search_company_assets('营业执照')
    evidence=json.loads(result.split('[evidence] ')[-1])['items']
    assert any(h['file']==target['file_name'] and h['role']=='raw-locator' for h in evidence)
    assert any(h['role']=='fact-verified' for h in evidence), 'Name location must not discard factual search'
    assert len(evidence)<=8
    assert '不是事实证据' in result
    # The same dense corpus must not hide the candidate in strict-type diagnostics.
    assert target['file_name'] in search_company_assets(target['file_name'],'business_license')

def test_symbol_only_filename_search_is_scoped_and_does_not_call_fts(library,monkeypatch):
    co,_=library
    target=insert(co,'★虚构营业执照.pdf');ingest.fail(target['id'],'parse','虚构故障')
    other=asyncio.run(lib.save_company(lib.CompanyBody(name='虚构其他主体')))
    foreign=insert(other,'★其他主体.pdf');ingest.fail(foreign['id'],'parse','不可泄漏')
    monkeypatch.setattr(db,'kb_search_segments',lambda *_args,**_kwargs:pytest.fail('Symbol-only query has no FTS expression'))
    result=asyncio.run(knowledge.list_items(q='★',company_id=co['id']))
    assert [i['id'] for i in result['items']]==[target['id']]
    assert asyncio.run(knowledge.list_items(q='★★',company_id=co['id']))['total']==0
    assert '不可泄漏' not in search_company_assets('★')
    assert target['file_name'] in search_company_assets('★')

def test_dense_name_hits_do_not_displace_factual_content_before_limit(library):
    co,_=library
    for n in range(40):
        item=insert(co,f'营业执照{n}.pdf');ingest.fail(item['id'],'parse','虚构故障')
    item=insert(co,'虚构合同.docx')
    md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True,exist_ok=True)
    md.write_text('营业执照是合同办理所需的材料。'+'虚构合同正文。'*500)
    outline.write_text('[]');ingest.parsed(item['id']);ingest.finish(item['id'],{'doc_type':'contract_case'})
    result=search_company_assets('营业执照')
    evidence=json.loads(result.split('[evidence] ')[-1])['items']
    assert any(h['file']==item['file_name'] and h['role']=='fact-verified' for h in evidence)
    assert any(h['role']=='raw-locator' for h in evidence)

def test_list_reads_content_once_only_for_filtered_page(library,monkeypatch):
    from pathlib import Path
    from collections import Counter
    co,_=library
    other=asyncio.run(lib.save_company(lib.CompanyBody(name='虚构其他主体')))
    for n in range(40):
        item=insert(co if n<12 else other,f'虚构资料{n}.docx')
        md,_,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True,exist_ok=True)
        md.write_text('虚构财务正文')
        db.kb_update_item(item['id'],doc_type='financial_report' if n<10 else 'other',review_status='confirmed' if n%2==0 else 'processed')
    reads=Counter();original=Path.read_text
    def read(path,*args,**kwargs):
        if path.suffix=='.md':reads[path]+=1
        return original(path,*args,**kwargs)
    monkeypatch.setattr(Path,'read_text',read)
    page=asyncio.run(knowledge.list_items(company_id=co['id'],doc_type='financial_report',review_status='confirmed',page_size=5))
    assert page['total']==5 and len(page['items'])==5
    assert sum(reads.values())==5
    assert {p.name for p in reads}=={i['file_name']+'.md' for i in page['items']}
    assert all(i['availability']['has_text'] and i['capability']=='searchable' for i in page['items'])
    reads.clear()
    page=asyncio.run(knowledge.list_items(company_id=co['id'],page=2,page_size=5))
    assert page['total']==12 and sum(reads.values())==5
    assert {p.name for p in reads}=={i['file_name']+'.md' for i in page['items']}
    reads.clear()
    assert asyncio.run(knowledge.list_items(company_id=co['id'],doc_type='business_license'))['total']==0
    assert not reads

def test_exact_title_keeps_name_priority_when_it_also_matches_body(library):
    co,_=library
    def indexed(name,body):
        item=insert(co,name)
        md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True,exist_ok=True)
        md.write_text(body);outline.write_text('[]')
        ingest.parsed(item['id']);ingest.finish(item['id'],{'doc_type':'other'})
        return item
    for n in range(12):
        indexed(f'虚构办理事项{n}.docx','营业执照作为办理事项材料，请核对原件。')
    target=indexed('营业执照.pdf','营业执照\n'+'虚构测试数据 '*1000)
    from app.knowledge import fts
    raw=db.kb_search_segments(fts.build_match_expr('营业执照'),limit=100,domain='kb',company_id=co['id'],exclude_identity=True)
    assert raw[-1]['item_id']==target['id'], 'Long target must rank below short body matches in FTS alone'
    page=asyncio.run(knowledge.list_items(q='营业执照',company_id=co['id'],page_size=5))
    assert page['total']==13
    assert page['items'][0]['id']==target['id'], 'Exact title must retain its priority even with a lower-scoring body hit'

def test_unicode_body_search_uses_index_tokenization_and_names_remain_caseless(library):
    co,_=library
    item=insert(co,'虚构说明.docx')
    md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True,exist_ok=True)
    md.write_text('虚构地址 Straße / Maße / Geschäftsbericht，仅用于Unicode测试。');outline.write_text('[]')
    ingest.parsed(item['id']);ingest.finish(item['id'],{'doc_type':'other'})
    for query in ('Straße','Maße','Geschäftsbericht'):
        result=asyncio.run(knowledge.list_items(q=query,company_id=co['id']))
        assert [i['id'] for i in result['items']]==[item['id']],query
    named=insert(co,'Straße虚构原件.pdf');ingest.fail(named['id'],'parse','虚构失败')
    result=asyncio.run(knowledge.list_items(q='STRASSE',company_id=co['id']))
    assert [i['id'] for i in result['items']]==[named['id']]
