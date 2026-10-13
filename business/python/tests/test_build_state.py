import asyncio
import json
import os
import subprocess
import sys
from pathlib import Path
import pytest
from docx import Document
from app import artifact_store, db, publish, runctx
from app.services import build_state, page_flow, workbench
from app.services.errors import BusinessError
from tests.util import init_env


@pytest.fixture
def env(tmp_path, monkeypatch):
    task, _ = init_env(tmp_path, monkeypatch)
    tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    root = artifact_store.work_dir(tid)
    content = {'numbering': 'decimal', 'page_target': 120, 'response_documents': [{'name': '技术标', 'directory': [
        {'目录名称': '实施方案', 'level': 1, 'children': [], '节点概述': '按需求实施'},
        {'目录名称': '服务保障', 'level': 1, 'children': []},
        {'目录名称': '目录', 'level': 1, 'children': [], '交付形态': '目录容器'}]}]}
    monkeypatch.setattr(page_flow, 'directory', lambda _: {'artifact_id': 'dir', 'content_seq': 1, 'content': content})
    (root / 'body/写作指引.md').write_text('# 写作指引\n\n| 节 | 模式 | 依据 | 素材 | 缺口/备注 |\n|---|---|---|---|---|\n| 实施方案 | 推理撰写 | — | — | 无 |\n| 服务保障 | 推理撰写 | — | — | 无 |\n| 目录 | — | — | — | 合册机械生成 |')
    (root / 'body/关键事实与承诺.md').write_text('| 事项 | 值 |\n|---|---|\n| 维护 | 三年 |')
    build_state.record_preparation(tid, 'guide')
    build_state.record_preparation(tid, 'promises')
    path = 'body/实施方案.docx'
    document = Document()
    document.add_heading('实施方案', 1)
    document.add_paragraph('按需求实施。')
    document.save(root / path)
    build_state.record_write(tid, path)
    build_state.record_validation(tid, path)
    return tid, root, content


def test_actual_chapters_exclude_mechanical_toc_and_keep_missing(env):
    tid, _, _ = env
    state = build_state.snapshot(tid)
    assert [(s['title'], s['status']) for s in state['body_sections']] == [('实施方案', 'ready'), ('服务保障', 'missing')]
    assert state['build_preparation'] == {'guide': 'ready', 'promises': 'ready'}


def test_numbering_changes_only_volume_not_body_or_preparation(env):
    tid, root, content = env
    complete_sections(tid, root)
    publish_volume(tid, root, '技术标')
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'
    before = (root / 'body/实施方案.docx').read_bytes()
    content['numbering'] = 'gov'
    state = build_state.snapshot(tid)
    assert state['body_sections'][0]['status'] == 'ready'
    assert state['build_preparation'] == {'guide': 'ready', 'promises': 'ready'}
    assert state['volume_states'][0]['status'] == 'stale'
    assert (root / 'body/实施方案.docx').read_bytes() == before


def test_directory_change_flags_only_affected_body_then_rename_retains_file(env):
    tid, root, content = env
    original = (root / 'body/实施方案.docx').read_bytes()
    content['response_documents'][0]['directory'][0]['节点概述'] = '按新需求实施'
    state = build_state.snapshot(tid)
    assert state['body_sections'][0]['status'] == 'stale'
    assert state['build_preparation']['guide'] == 'stale'
    # 只读校验不能把旧正文的依据更新。
    build_state.record_validation(tid, 'body/实施方案.docx')
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'stale'
    content['response_documents'][0]['directory'][0]['目录名称'] = '实施安排'
    state = build_state.snapshot(tid)
    assert state['body_sections'][0]['status'] == 'missing'
    assert state['orphan_body_files'] == ['body/实施方案.docx']
    assert (root / 'body/实施方案.docx').read_bytes() == original


def test_user_saved_guide_is_reused_and_promises_need_update(env):
    tid, root, _ = env
    path = 'body/写作指引.md'
    content = (root / path).read_text()
    asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path=path, content=content+'\n\n用户补充要求', base_hash=workbench._hash(content))))
    state = build_state.snapshot(tid)
    assert state['build_preparation'] == {'guide': 'ready', 'promises': 'stale'}
    assert state['body_sections'][0]['status'] == 'stale'
    assert '用户补充要求' in (root / path).read_text()


def test_invalid_guide_cannot_be_recorded_ready_or_replace_user_file(env):
    tid, root, _ = env
    path = root / 'body/写作指引.md'
    before = path.read_text()
    with pytest.raises(Exception, match='校验未通过'):
        asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path='body/写作指引.md', content='# invalid', base_hash=workbench._hash(before))))
    assert path.read_text() == before
    path.write_text('# invalid')
    with pytest.raises(ValueError, match='校验未通过'):
        build_state.record_preparation(tid, 'guide')
    assert build_state.snapshot(tid)['build_preparation']['guide'] == 'stale'


@pytest.mark.parametrize('prefix', ['./body/', r'.\body' + chr(92), 'body/../body/'])
@pytest.mark.parametrize('name,stage', [('写作指引.md', 'guide'), ('关键事实与承诺.md', 'promises')])
def test_preparation_aliases_validate_before_write_and_record_valid_saves(env, prefix, name, stage):
    tid, root, _ = env
    file = root / 'body' / name
    before = file.read_text()
    with pytest.raises(BusinessError, match='校验未通过|承诺清单须包含') as error:
        asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path=prefix + name, content='# invalid', base_hash=workbench._hash(before))))
    assert error.value.status_code == 422
    assert file.read_text() == before
    assert not file.with_name(name + '.restorepoints').exists()
    assert build_state.snapshot(tid)['build_preparation'][stage] == 'ready'
    asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path=prefix + name, content=before, base_hash=workbench._hash(before))))
    assert build_state.snapshot(tid)['build_preparation'][stage] == 'ready'


def test_preparation_case_alias_cannot_bypass_validation(env):
    tid, root, _ = env
    file = root / 'body/写作指引.md'
    if not (root / 'BODY/写作指引.md').is_file():
        pytest.skip('当前文件系统区分大小写，没有此别名')
    before = file.read_text()
    with pytest.raises(BusinessError, match='校验未通过'):
        asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path='BODY/写作指引.md', content='# invalid', base_hash=workbench._hash(before))))
    assert file.read_text() == before


def test_preparation_hash_detects_unrecorded_changes_and_names_cause(env):
    tid, root, content = env
    content['page_target'] = 200
    state = build_state.snapshot(tid)
    assert '目录内容或目标篇幅已变化' in state['build_preparation_changes']['guide']
    build_state.record_preparation(tid, 'guide')
    assert build_state.snapshot(tid)['build_preparation'] == {'guide': 'ready', 'promises': 'stale'}
    path = root / 'body/写作指引.md'
    path.write_text(path.read_text() + '\n未经确认的修改')
    state = build_state.snapshot(tid)
    assert state['build_preparation']['guide'] == 'stale'
    assert '准备文件已修改，尚未通过校验' in state['build_preparation_changes']['guide']


def test_model_cannot_change_remove_or_duplicate_existing_promises(env):
    tid, root, _ = env
    current = (root / 'body/关键事实与承诺.md').read_text()
    build_state.check_preparation_write(tid, 'promises', current.replace('三年', '三 年'))
    for proposal in (current.replace('三年', '一年'), '# deleted', current + '\n| 维护 | 三年 |'):
        with pytest.raises(ValueError):
            build_state.check_preparation_write(tid, 'promises', proposal)
    assert (root / 'body/关键事实与承诺.md').read_text() == current
    asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path='body/关键事实与承诺.md', content=current.replace('三年', '五年'), base_hash=workbench._hash(current))))
    assert '五年' in (root / 'body/关键事实与承诺.md').read_text()
    assert build_state.snapshot(tid)['build_preparation']['promises'] == 'ready'


def test_full_escaped_promise_value_is_protected_and_dispatched(env):
    from app.dispatch_enrich import _promise_lines
    tid, root, _ = env
    content = '| 事项 | 值 | 说明 |\n|---|---|---|\n| 服务时间 | 电话支持\\|7×24小时 | 用户确认 |\n'
    path = root / 'body/关键事实与承诺.md'
    path.write_text(content)
    build_state.record_preparation(tid, 'promises')
    assert build_state.promise_values(content)[0] == {'服务时间': '电话支持|7×24小时'}
    assert _promise_lines(tid) == ['- 服务时间：电话支持|7×24小时（用户确认）']
    with pytest.raises(ValueError, match='7×24小时 → 电话支持\\|5×8小时'):
        build_state.check_preparation_write(tid, 'promises', content.replace('7×24', '5×8'))
    assert path.read_text() == content


def test_guide_readiness_and_dispatch_share_standard_table_semantics(env):
    from app.dispatch_enrich import _guide_row
    tid, root, directory = env
    directory['response_documents'][0]['directory'][0]['目录名称'] = '实施|方案'
    path = root / 'body/写作指引.md'
    guide = path.read_text().replace('实施方案', r'实施\|方案').replace('推理撰写', '**推理撰写**')
    guide = guide.replace('| 无 |', r'| 补充 A\|B 资料 |')
    # 代码示例中的非法节名不能参与就绪对账或派发。
    guide += '\n\n~~~md\n| 节 | 模式 |\n|---|---|\n| 假章 | 非法模式 |\n~~~\n'
    path.write_text(guide)
    build_state.record_preparation(tid, 'guide')
    assert build_state.snapshot(tid)['build_preparation']['guide'] == 'ready'
    cells, columns = _guide_row(tid, directory, '技术标', '实施|方案')
    assert cells[columns['节']] == '实施|方案'
    assert cells[columns['模式']] == '推理撰写'
    assert cells[columns['缺口/备注']] == '补充 A|B 资料'


def test_repair_duplicate_intermediate_promises_preserves_each_confirmed_value(env):
    tid, root, _ = env
    path = root / 'body/关键事实与承诺.md'
    good = path.read_text()
    duplicated = good + '\n| 维护 | 三 年 |'
    path.write_text(duplicated)
    with pytest.raises(ValueError, match='重复'):
        build_state.record_preparation(tid, 'promises')
    build_state.check_preparation_write(tid, 'promises', good)
    path.write_text(good)
    build_state.record_preparation(tid, 'promises')
    assert build_state.snapshot(tid)['build_preparation']['promises'] == 'ready'
    path.write_text(good + '\n| 维护 | 五年 |')
    for proposal in (good, good.replace('三年', '五年')):
        with pytest.raises(ValueError, match='已有承诺值需要用户处理'):
            build_state.check_preparation_write(tid, 'promises', proposal)


@pytest.mark.parametrize('bad_row', ['| | 三年 |', '| 维护期 | |', '| 维护期 |'])
def test_invalid_promise_rows_cannot_be_written_saved_or_recorded_ready(env, bad_row):
    tid, root, _ = env
    path = root / 'body/关键事实与承诺.md'
    before = path.read_text()
    invalid = before + '\n' + bad_row
    with pytest.raises(ValueError, match='事项和值不能为空'):
        build_state.check_preparation_write(tid, 'promises', invalid)
    with pytest.raises(Exception, match='事项和值不能为空'):
        asyncio.run(workbench.write_content(workbench.WorkbenchWrite(task_id=tid, path='body/关键事实与承诺.md', content=invalid, base_hash=workbench._hash(before))))
    assert path.read_text() == before
    path.write_text(invalid)
    with pytest.raises(ValueError, match='事项和值不能为空'):
        build_state.record_preparation(tid, 'promises')
    assert build_state.snapshot(tid)['build_preparation']['promises'] != 'ready'
    path.unlink()
    with pytest.raises(ValueError, match='事项和值不能为空'):
        build_state.check_preparation_write(tid, 'promises', invalid)


def test_written_draft_needs_validation_and_read_only_check_does_not_refresh_basis(env):
    tid, root, content = env
    path = 'body/实施方案.docx'
    build_state.record_write(tid, path)
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'draft'
    content['page_target'] = 200
    build_state.record_validation(tid, path)
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'draft'
    # 明确重新写入并通过校验，才算当前依据对应的已完成章节。
    build_state.record_write(tid, path)
    build_state.record_validation(tid, path)
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'ready'


def test_page_target_contract_rejects_invalid_values():
    from app.contracts.tender_directory import TenderDirectoryModel
    from pydantic import ValidationError
    for n in (0, 1001, -1):
        with pytest.raises(ValidationError):
            TenderDirectoryModel.model_validate({'response_documents': [], 'page_target': n})
    assert TenderDirectoryModel.model_validate({'response_documents': [], 'page_target': None}).page_target is None


def publish_volume(tid, root, name):
    from app.tools.docx_ops import _publish_volume
    file = root / f'body/整本-{name}.docx'
    doc = Document()
    doc.add_heading(name, 1)
    doc.add_paragraph('工程整册正文')
    doc.save(file)
    result = _publish_volume(tid, file, name, 1, 0, 0)
    assert '发布失败' not in result, result
    return next(r for r in db.list_artifact_index(task_id=tid) if r['display_name'] == name)


def complete_sections(tid, root):
    for section in build_state.sections(tid):
        path = root / section['path']
        if not path.is_file():
            path.parent.mkdir(parents=True, exist_ok=True)
            doc = Document(); doc.add_heading(section['title'], 1); doc.add_paragraph('章节正文。')
            doc.save(path)
        build_state.record_section(tid, section['path'])


def test_renamed_volume_keeps_old_artifact_but_requires_new_publication(env):
    tid, root, content = env
    complete_sections(tid, root)
    old = publish_volume(tid, root, '技术标')
    old_file = artifact_store.package_files(old['artifact_id'], old)[0]
    original = old_file.read_bytes()
    content['response_documents'][0]['name'] = '新技术标'
    state = build_state.snapshot(tid)['volume_states']
    assert state == [{'name': '新技术标', 'path': 'body/整本-新技术标.docx', 'artifact_id': None, 'status': 'missing'}]
    new = publish_volume(tid, root, '新技术标')
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'
    assert new['artifact_id'] != old['artifact_id']
    assert old_file.read_bytes() == original


def test_each_volume_tracks_own_sources_and_requires_published_package(env):
    tid, root, content = env
    content['response_documents'].append({'name': '商务标', 'directory': [{'目录名称': '资质材料', 'level': 1, 'children': []}]})
    for name, title in [('技术标', '实施方案'), ('商务标', '资质材料')]:
        folder = root / 'body' / name
        folder.mkdir()
        doc = Document(); doc.add_heading(title, 1); doc.save(folder / (title+'.docx'))
    complete_sections(tid, root)
    publish_volume(tid, root, '技术标')
    business = publish_volume(tid, root, '商务标')
    assert [v['status'] for v in build_state.snapshot(tid)['volume_states']] == ['ready', 'ready']
    doc = Document(); doc.add_heading('实施方案', 1); doc.add_paragraph('修改技术正文'); doc.save(root / 'body/技术标/实施方案.docx')
    assert [v['status'] for v in build_state.snapshot(tid)['volume_states']] == ['stale', 'ready']
    artifact_store.package_files(business['artifact_id'], business)[0].unlink()
    assert [v['status'] for v in build_state.snapshot(tid)['volume_states']] == ['stale', 'missing']


def test_failed_volume_publication_cannot_mark_current_volume_ready(env, monkeypatch):
    from app.tools.docx_ops import _publish_volume
    tid, root, _ = env
    complete_sections(tid, root)
    old = publish_volume(tid, root, '技术标')
    (root / 'body/实施方案.docx').write_bytes(b'changed section')
    before = build_state.read(tid)['volumes'][old['artifact_id']]
    def fail(*args, **kwargs):
        raise publish.PublishError('工程发布失败')
    monkeypatch.setattr(publish, 'publish_file_artifact', fail)
    assert '工程发布失败' in _publish_volume(tid, root / 'body/整本-技术标.docx', '技术标', 1, 0, 0)
    assert build_state.read(tid)['volumes'][old['artifact_id']] == before
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'stale'


def test_repeated_direct_outline_publish_keeps_consumed_options(tmp_path, monkeypatch):
    import json
    from app.tools.publish import publish_artifact
    task, _ = init_env(tmp_path, monkeypatch)
    tid = task['id']; artifact_store.ensure_task_skeleton(tid)
    root = artifact_store.work_dir(tid)
    options = root / 'outline/.generation-options.json'
    options.write_text(json.dumps({'page_target': 200, 'numbering': 'gov'}))
    staging = root.parent / '_meta/staging'; staging.mkdir(parents=True, exist_ok=True)
    runctx.set_run('fixture-session', 'fixture-run', tid)
    try:
        for index in (1, 2):
            draft = staging / f'draft-{index}.json'
            draft.write_text(json.dumps({'response_documents': [{'name': '技术标', 'directory': [{'目录名称': f'方案{index}', 'level': 1, 'children': []}]}]}))
            assert '[发布成功]' in publish_artifact('tender.directory/tender-response-docs@1', str(draft))
            current = page_flow.directory(tid)['content']
            assert current['page_target'] == 200 and current['numbering'] == 'gov'
            assert not options.exists()
        options.write_text(json.dumps({'page_target': None, 'numbering': 'none'}))
        draft = staging / 'new-settings.json'; draft.write_text(json.dumps(current))
        assert '[发布成功]' in publish_artifact('tender.directory/tender-response-docs@1', str(draft))
        current = page_flow.directory(tid)['content']
        assert current['page_target'] is None and current['numbering'] == 'none'
    finally:
        runctx.clear_run()


def export_result(tid):
    request = {'operation': 'export', 'data_root': os.environ['DATA_DIR'], 'task_id': tid}
    result = subprocess.run([sys.executable, str(Path(__file__).parents[1] / 'runner.py')],
                            input=json.dumps(request), text=True, capture_output=True, check=False)
    return json.loads(result.stdout)


def test_missing_prose_cannot_be_ready_even_with_published_stamp(env):
    tid, root, _ = env
    publish_volume(tid, root, '技术标')
    assert build_state.snapshot(tid)['volume_states'][0]['status'] != 'ready'


def test_pending_chapter_requires_validation_before_volume_is_ready(env):
    tid, root, _ = env
    complete_sections(tid, root)
    publish_volume(tid, root, '技术标')
    build_state.record_write(tid, 'body/实施方案.docx')
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'stale'
    build_state.record_validation(tid, 'body/实施方案.docx')
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'


def test_partial_reassembly_keeps_previous_published_package(env):
    tid, root, content = env
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    complete_sections(tid, root)
    assert 'result' in export_result(tid)
    row = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    original = artifact_store.package_files(row['artifact_id'], row)[0].read_bytes()
    (root / 'body/服务保障.docx').unlink()
    assert export_result(tid)['status'] == 422
    current = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    assert current['content_seq'] == row['content_seq']
    assert artifact_store.package_files(current['artifact_id'], current)[0].read_bytes() == original
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'stale'


@pytest.mark.parametrize('already_published', [False, True])
def test_pending_revision_cannot_publish_until_validated(env, already_published):
    tid, root, content = env
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    complete_sections(tid, root)
    previous = None
    if already_published:
        assert 'result' in export_result(tid)
        previous = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
        previous_bytes = artifact_store.package_files(previous['artifact_id'], previous)[0].read_bytes()
        previous_stamp = build_state.read(tid)['volumes'][previous['artifact_id']]

    path = 'body/实施方案.docx'
    doc = Document(); doc.add_heading('实施方案', 1); doc.add_paragraph('尚未校验的修改正文。')
    doc.save(root / path)
    build_state.record_write(tid, path)
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'draft'
    assert export_result(tid)['status'] == 422
    rows = [r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume']
    if previous:
        assert len(rows) == 1
        assert rows[0]['artifact_id'] == previous['artifact_id']
        assert rows[0]['content_seq'] == previous['content_seq']
        assert artifact_store.package_files(rows[0]['artifact_id'], rows[0])[0].read_bytes() == previous_bytes
        assert build_state.read(tid)['volumes'][previous['artifact_id']] == previous_stamp
    else:
        assert rows == []
        assert build_state.read(tid).get('volumes', {}) == {}
    assert '尚未校验的修改正文。' in '\n'.join(p.text for p in Document(root / 'body/整本-技术标.docx').paragraphs)

    build_state.record_validation(tid, path)
    assert build_state.snapshot(tid)['body_sections'][0]['status'] == 'ready'
    assert 'result' in export_result(tid)
    current = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
    assert current['content_seq'] == (previous['content_seq'] + 1 if previous else 1)
    if previous:
        assert current['artifact_id'] == previous['artifact_id']
    published = Document(artifact_store.package_files(current['artifact_id'], current)[0])
    assert '尚未校验的修改正文。' in '\n'.join(p.text for p in published.paragraphs)
    assert build_state.snapshot(tid)['volume_states'][0]['status'] == 'ready'


def test_pending_revision_blocks_only_its_own_volume(env):
    from app.tools.docx_ops import docx_assemble_volume
    tid, root, content = env
    content['response_documents'] = [
        {'name': name, 'directory': [{'目录名称': '实施方案', 'level': 1, 'children': []}]}
        for name in ('技术标', '商务标')]
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    complete_sections(tid, root)
    assert 'result' in export_result(tid)
    previous = {r['display_name']: r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume'}
    path = 'body/商务标/实施方案.docx'
    doc = Document(); doc.add_heading('实施方案', 1); doc.add_paragraph('商务标待校验修改。')
    doc.save(root / path)
    build_state.record_write(tid, path)
    runctx.set_run('fixture-session', 'fixture-run', tid)
    try:
        result = docx_assemble_volume()
        assert '商务标：待校验章节（实施方案），未发布整册' in result
        assert '技术标：待校验' not in result
    finally:
        runctx.clear_run()
    assert [v['status'] for v in build_state.snapshot(tid)['volume_states']] == ['ready', 'stale']
    assert export_result(tid)['status'] == 422
    build_state.record_validation(tid, path)
    assert 'result' in export_result(tid)
    current = {r['display_name']: r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume'}
    assert current['技术标']['content_seq'] == previous['技术标']['content_seq']
    assert current['商务标']['content_seq'] == previous['商务标']['content_seq'] + 1
    assert all(v['status'] == 'ready' for v in build_state.snapshot(tid)['volume_states'])


def test_real_partial_assembly_rejects_export_then_recovers_after_completion(env):
    from app.tools.docx_ops import docx_assemble_volume
    tid, root, content = env
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    runctx.set_run('fixture-session', 'fixture-run', tid)
    try:
        result = docx_assemble_volume()
        assert '服务保障' in result and '缺失' in result
        assert (root / 'body/整本-技术标.docx').is_file()  # 保留部分成果供继续制作
        assert not any(r['kind'] == 'tender.volume' for r in db.list_artifact_index(task_id=tid))
        assert build_state.snapshot(tid)['volume_states'][0]['status'] != 'ready'
        assert export_result(tid)['status'] == 422

        doc = Document(); doc.add_heading('服务保障', 1); doc.add_paragraph('提供维护保障。')
        doc.save(root / 'body/服务保障.docx')
        build_state.record_section(tid, 'body/服务保障.docx')
        assert 'result' in export_result(tid)
        state = build_state.snapshot(tid)
        assert state['volume_states'][0]['status'] == 'ready'
        row = next(r for r in db.list_artifact_index(task_id=tid) if r['kind'] == 'tender.volume')
        merged = Document(artifact_store.package_files(row['artifact_id'], row)[0])
        assert '提供维护保障。' in '\n'.join(p.text for p in merged.paragraphs)
    finally:
        runctx.clear_run()


@pytest.mark.parametrize('name, expected', [(' 技术标 ', '技术标'), ('\t技术标\n', '技术标'), ('', '主册')])
def test_volume_name_identity_matches_real_assembly_and_section_recording(env, name, expected):
    from app.tools.docx_ops import docx_assemble_volume
    tid, root, content = env
    content['response_documents'][0]['name'] = name
    doc = Document(); doc.add_heading('服务保障', 1); doc.add_paragraph('维护保障。')
    doc.save(root / 'body/服务保障.docx')
    publish.publish_artifact('tender.directory/tender-response-docs@1', content, task_id=tid)
    runctx.set_run('fixture-session', 'fixture-run', tid)
    try:
        result = docx_assemble_volume()
        assert '发布失败' not in result, result
        for section in build_state.sections(tid):
            build_state.record_section(tid, section['path'])
        state = build_state.snapshot(tid)
        assert {s['volume'] for s in state['body_sections']} == {expected}
        assert all(s['status'] == 'ready' for s in state['body_sections'])
        assert state['volume_states'][0]['name'] == expected
        assert state['volume_states'][0]['status'] == 'ready'
        assert state['volume_states'][0]['artifact_id']
        assert 'result' in export_result(tid)
    finally:
        runctx.clear_run()
