import json
from pathlib import Path
import pytest
from docx import Document
from tests.test_section_write_atomicity import env, request
from app.tools.docx_ops import docx_section_revise
from app.services import build_state


def test_mixed_target_cannot_silently_edit_paragraph(env):
    _, root, _ = env
    p = root / 'body/实施方案.docx'; before = p.read_bytes()
    result = docx_section_revise('body/实施方案.docx', [{'para': 1, 'table': 'T1', 'row': 1, 'col': 1, 'action': 'replace', 'find': '原始服务方案', 'text': '误改段落'}])
    assert result.startswith('[修订失败]') and p.read_bytes() == before
    assert 'P1' in result and 'para=4' in result

@pytest.mark.parametrize('bad', ['[{"para":1,"action":"delete"}]', [{'table': 'T1', 'row': 1, 'col': 1, 'action': 'fill', 'text': '错误'}], [{'para': True, 'action': 'delete'}], [{'para': 0, 'action': 'delete'}]])
def test_native_contract_rejects_invalid_target_before_write(env, bad):
    _, root, _ = env
    p = root / 'body/实施方案.docx'; before = p.read_bytes(); ledger = (root/'body/.build-state.json').read_bytes()
    code, result = request(env, parameters={'path': 'body/实施方案.docx', 'edits': bad})
    assert code != 0 and 'error' in result
    assert p.read_bytes() == before and (root/'body/.build-state.json').read_bytes() == ledger


def test_failed_batch_view_is_original_and_recovery_uses_real_location(env):
    tid, root, _ = env; p = root/'body/实施方案.docx'
    d = Document(p); d.add_paragraph('真正的定位原文'); d.save(p); before = p.read_bytes()
    code, result = request(env, parameters={'path': 'body/实施方案.docx', 'edits': [
        {'para': 2, 'action': 'replace', 'find': '真正的定位原文', 'text': '不应发布的中间内容'},
        {'para': 1, 'action': 'replace', 'find': '猜测不存在的内容', 'text': '错误'}]})
    assert code == 0 and result['result'].startswith('[修订失败]')
    assert p.read_bytes() == before and '不应发布的中间内容' not in result['result']
    assert '[P2]' in result['result'] and '真正的定位原文' in result['result']
    code, result = request(env, parameters={'path': 'body/实施方案.docx', 'edits': [{'para': 2, 'action': 'replace', 'find': '真正的定位原文', 'text': '已按实际位置恢复'}]})
    assert code == 0 and result['result'].startswith('[已修订]')
    code, result = request(env, 'validate_body', {'section': './body/实施方案.docx'})
    assert code == 0 and result['result'].startswith('[校验通过]')
    assert 'body/实施方案.docx' not in build_state.read(tid).get('pending_sections', {})


def test_nonempty_cell_fill_returns_original_view_and_replace_recovers(env):
    _, root, _ = env; p = root/'body/实施方案.docx'; d = Document(p)
    d.add_table(rows=1, cols=1).cell(0,0).text='原格值'; d.save(p); before=p.read_bytes()
    result=docx_section_revise('body/实施方案.docx',[{'table':1,'row':1,'col':1,'action':'fill','text':'新格值'}])
    assert result.startswith('[修订失败]') and '原格值' in result and p.read_bytes()==before
    assert docx_section_revise('body/实施方案.docx',[{'table':1,'row':1,'col':1,'action':'replace','find':'原格值','text':'新格值'}]).startswith('[已修订]')


def test_comments_are_separate_metadata_not_paragraph_replace_text(env):
    from app.tools.docx_ops import docx_section_read
    _,root,_=env;p=root/'body/实施方案.docx';d=Document(p)
    d.add_comment(d.paragraphs[0].runs,'旧批注：1小时到场，联系人误用签字人',author='历史记录');d.save(p)
    view=docx_section_read('body/实施方案.docx')
    body=next(line for line in view.splitlines() if line.startswith('[P1]'))
    assert '旧批注' not in body and '原始服务方案' in body
    assert '[批注@P1] 旧批注' in view and '非正文' in view
    before=p.read_bytes();result=docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'旧批注','text':''}])
    assert result.startswith('[修订失败]') and p.read_bytes()==before
    assert docx_section_revise('body/实施方案.docx',[{'para':1,'action':'replace','find':'原始服务方案','text':'已修订的真实正文'}]).startswith('[已修订]')
    assert '旧批注' in docx_section_read('body/实施方案.docx')


def test_legacy_mixed_target_payload_is_rejected_without_silent_paragraph_edit(env):
    _, root, _ = env; path = root/'body/实施方案.docx'; before = path.read_bytes(); ledger = (root/'body/.build-state.json').read_bytes()
    legacy = json.dumps([{'para': 1, 'table': 'T1', 'row': 1, 'col': 1, 'action': 'replace', 'find': '原始服务方案', 'text': '错误改到段落'}])
    code, result = request(env, parameters={'path': 'body/实施方案.docx', 'edits': legacy})
    assert code != 0 and 'error' in result
    assert path.read_bytes() == before and (root/'body/.build-state.json').read_bytes() == ledger
