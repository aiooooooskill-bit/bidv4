"""逐页视觉验收中发现的合册边界；只覆盖新增修复。"""
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

from app.services import document_format
from app.tools.docx_ops import _remove_leading_page_breaks


def test_assembly_boundary_removes_only_breaks_before_first_content():
    source = Document()
    p = source.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p = source.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run('附件原文')
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run('内部下一页')
    source.add_paragraph('后续段落').paragraph_format.page_break_before = True
    original = source.element.xml
    elements = [deepcopy(p._p) for p in source.paragraphs]
    _remove_leading_page_breaks(elements)
    assert len(elements[0].xpath('.//w:br[@w:type="page"]')) == 0
    assert len(elements[1].xpath('.//w:br[@w:type="page"]')) == 1
    assert ''.join(elements[1].itertext())
    assert [t.text for t in elements[1].iter(qn('w:t'))] == ['附件原文','内部下一页']
    assert elements[1].pPr.pageBreakBefore.val is False
    assert elements[2].pPr.pageBreakBefore.val is True
    assert source.element.xml == original


def test_assembly_boundary_does_not_strip_page_breaks_after_image_or_table():
    source = Document()
    p = source.add_paragraph()
    p.add_run()._r.append(OxmlElement('w:drawing'))
    p.add_run().add_break(WD_BREAK.PAGE)
    after = source.add_paragraph('下一页')
    after.paragraph_format.page_break_before = True
    elements = [deepcopy(p._p),deepcopy(after._p)]
    _remove_leading_page_breaks(elements)
    assert len(elements[0].xpath('.//w:br[@w:type="page"]')) == 1
    assert elements[1].pPr.pageBreakBefore.val is True
    table = source.add_table(rows=1,cols=1)
    elements = [deepcopy(table._tbl),deepcopy(after._p)]
    _remove_leading_page_breaks(elements)
    assert elements[1].pPr.pageBreakBefore.val is True


def test_native_heading_black_without_changing_colored_body():
    doc = Document()
    heading = doc.add_heading('三级标题',3)
    body = doc.add_paragraph('原稿红色提示')
    for p in (heading,body):
        color = OxmlElement('w:color')
        color.set(qn('w:val'),'FF0000')
        color.set(qn('w:themeColor'),'accent1')
        p.runs[0]._r.get_or_add_rPr().append(color)
    document_format.apply(doc,{},[heading._p,body._p])
    assert heading.style.name == 'Heading 3'
    assert heading.runs[0]._r.rPr.color.attrib == {qn('w:val'):'000000'}
    assert doc.styles['Heading 3'].element.rPr.color.attrib == {qn('w:val'):'000000'}
    assert body.runs[0]._r.rPr.color.get(qn('w:val')) == 'FF0000'


def test_oversize_merged_table_fits_custom_body_without_changing_content():
    doc = Document()
    table = doc.add_table(rows=3,cols=3)
    for i,column in enumerate(table.columns):
        column.width = 635*4000
        for row in table.rows: row.cells[i].width = 635*4000
    table.cell(0,0).merge(table.cell(0,1)).text = '合并表头不丢失'
    table.cell(1,2).text = '原始事实123'
    nested = table.cell(2,0).add_table(rows=1,cols=2)
    nested.cell(0,0).text='内表原文'
    before = [t.text for t in table._tbl.iter(qn('w:t'))]
    copy = deepcopy(table._tbl)
    document_format.fit_table(copy,9000)
    assert sum(int(c.get(qn('w:w'))) for c in copy.tblGrid) == 9000
    assert [t.text for t in copy.iter(qn('w:t'))] == before
    assert copy.xpath('./w:tr[1]/w:tc[1]/w:tcPr/w:tcW')[0].get(qn('w:w')) == '6000'
    assert copy.xpath('./w:tr[1]/w:tc[1]/w:tcPr/w:gridSpan')[0].get(qn('w:val')) == '2'
    original = table._tbl.xml
    document_format.fit_table(table._tbl,13000)
    assert table._tbl.xml == original


def test_short_label_stays_with_following_image_without_changing_text():
    doc = Document()
    label = doc.add_paragraph('原件证书图片说明')
    picture = doc.add_paragraph()
    picture.add_run()._r.append(OxmlElement('w:drawing'))
    next_text = doc.add_paragraph('正文不强行绑定后续段落')
    document_format.apply(doc,{},[label._p,picture._p,next_text._p])
    assert label.paragraph_format.keep_with_next is True
    assert label.text == '原件证书图片说明'
    assert next_text.paragraph_format.keep_with_next is None


def test_short_source_labels_and_table_rows_use_native_pagination():
    doc = Document()
    label = doc.add_paragraph('服务端安全'); label.runs[0].bold = True
    ordinary = doc.add_paragraph('普通正文不改变层级。')
    table = doc.add_table(rows=3,cols=2)
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells): cell.text = f'原文{i}-{j}'
    caption = doc.add_paragraph('图 1 原表格示意图')
    before = [t.text for t in table._tbl.iter(qn('w:t'))]
    document_format.apply(doc,{},[label._p,ordinary._p,table._tbl,caption._p])
    assert label.paragraph_format.keep_with_next is True
    assert label.style.name == 'Normal'
    assert ordinary.paragraph_format.keep_with_next is None
    assert len(table._tbl.xpath('./w:tr/w:trPr/w:cantSplit')) == 3
    assert all(p.find('./'+qn('w:pPr')+'/'+qn('w:keepNext')) is not None for p in table._tbl.iter(qn('w:p')))
    assert [t.text for t in table._tbl.iter(qn('w:t'))] == before


def test_inline_raster_ratio_body_fit_and_cropped_original_preserved():
    from io import BytesIO
    from PIL import Image
    from docx.shared import Inches
    stream = BytesIO(); Image.new('RGB',(400,200),'white').save(stream,format='PNG'); stream.seek(0)
    doc = Document(); doc.sections[0].page_width = Inches(5)
    picture = doc.add_picture(stream,width=Inches(6),height=Inches(6))
    cropped = doc.add_picture(stream,width=Inches(2),height=Inches(3))
    crop = cropped._inline.find('.//'+qn('a:srcRect'))
    if crop is None:
        crop = OxmlElement('a:srcRect'); cropped._inline.find('.//'+qn('pic:blipFill')).append(crop)
    crop.set('t','10000')
    crop_xml = cropped._inline.xml
    part = next(p for p in doc.part.related_parts.values() if p.content_type == 'image/png')
    original = part.blob
    document_format.fit_inline_images(doc,doc.element.body)
    assert picture.width == 2*picture.height
    assert picture.width <= doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    assert cropped._inline.xml == crop_xml
    assert part.blob == original
    doc.sections[0].page_width = Inches(4)
    document_format.fit_inline_images(doc,doc.element.body)
    assert cropped.width <= doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    assert cropped.height/cropped.width == 1.5
    assert crop.get('t') == '10000' and part.blob == original


def test_embedded_own_title_requires_unique_native_heading():
    from app.tools.docx_ops import _embedded_own_title
    doc = Document(); doc.add_paragraph('上一节完整落款')
    doc.add_heading('数据更新承诺函',4); doc.add_paragraph('本节正文')
    assert _embedded_own_title(list(doc.element.body)[:-1],doc,'数据更新承诺函（格式6-2）') == 1
    doc.add_heading('数据更新承诺函',4)
    assert _embedded_own_title(list(doc.element.body)[:-1],doc,'数据更新承诺函（格式6-2）') is None


def test_material_index_uses_native_targets_and_keeps_unresolved_cells():
    from app.services import page_references,document_pagination
    doc = Document(); table = doc.add_table(rows=3,cols=2)
    for row,values in zip(table.rows,[['资料名称','页码范围'],['方案（格式6-2）','8-9'],['未提供的材料','50-51']]):
        for cell,text in zip(row.cells,values):cell.text=text
    h=doc.add_heading('方案（格式6-2）',1)
    document_pagination.bookmark(h._p,'_TenderHeading1',1)
    doc.add_paragraph('原文内容')
    refs,unresolved=page_references.prepare(doc,[{'name':'_TenderHeading1','depth':1,'title':'1 方案（格式6-2）'}])
    assert len(refs) == len(unresolved) == 1
    assert table.cell(2,1).text == '50-51'
    assert 'PAGEREF TenderRefStart' in table.cell(1,1)._tc.xml
    assert refs[0]['original']=='8-9'
    assert '8-9' not in table.cell(1,1).text and 'PAGEREF TenderRefEnd' in table.cell(1,1)._tc.xml


def test_assembly_title_follows_original_tail_and_keeps_both_title_comments(tmp_path,monkeypatch):
    from tests.test_document_pagination import setup,seed,node
    from app.tools.docx_ops import docx_section_create,docx_assemble_volume
    tid,root=setup(tmp_path,monkeypatch)
    seed(tid,[node('目录'),node('上一章'),node('本章方案（格式6-2）')])
    docx_section_create('body/上一章','上一章',paragraphs='上一章原文。')
    docx_section_create('body/本章方案（格式6-2）','本章方案（格式6-2）',paragraphs='上一章完整落款。')
    source=root/'body/本章方案（格式6-2）.docx';d=Document(source)
    d.add_comment(d.paragraphs[0].runs[0],text='人工章标题批注',author='原作者')
    h=d.add_heading('本章方案',4);d.add_comment(h.runs[0],text='原件章标题批注',author='原作者')
    d.add_paragraph('本章原始正文不得丢失。');d.save(source);original=source.read_bytes()
    result=docx_assemble_volume();assert '已发布为整本' in result,result
    out=Document(root/'body/整本-格式验证册.docx');ps=out.paragraphs
    tail=next(i for i,p in enumerate(ps) if p.text=='上一章完整落款。')
    heading=next(i for i,p in enumerate(ps) if p.text=='本章方案（格式6-2）' and p.style.name=='Heading 1')
    body=next(i for i,p in enumerate(ps) if p.text=='本章原始正文不得丢失。')
    assert tail < heading < body
    assert sorted(c.text for c in out.comments)==['人工章标题批注','原件章标题批注']
    ids={e.get(qn('w:id')) for e in ps[heading]._p.iter(qn('w:commentReference'))}
    assert len(ids)==2
    assert source.read_bytes()==original


def test_heading_stack_spacing_stops_at_table_and_explicit_page_break():
    from docx.shared import Pt
    d=Document();parent=d.add_heading('父标题',1);child=d.add_heading('子标题',2)
    d.add_table(rows=1,cols=1);after_table=d.add_heading('表后新标题',2)
    separate=d.add_heading('另页标题',2);separate.paragraph_format.page_break_before=True
    document_format.compact_heading_stacks(d)
    assert child.paragraph_format.space_before == Pt(0)
    assert parent.paragraph_format.space_before is None
    assert after_table.paragraph_format.space_before is None
    assert separate.paragraph_format.space_before is None


def test_typed_source_labels_and_header_after_merged_instructions():
    d=Document();labels=[d.add_paragraph(t) for t in ['（1）智能监控系统','d. 定制化数据源抓取','【概要】']]
    table=d.add_table(rows=4,cols=3);table.cell(0,0).merge(table.cell(0,2)).text='响应表的完整原始说明，不应误识别为表头，也不能删除。'
    for c,t in zip(table.rows[1].cells,['序号','响应内容','说明']):c.text=t
    for c,t in zip(table.rows[2].cells,['1','原始响应','完整保留']):c.text=t
    document_format.apply(d,{},[p._p for p in labels]+[table._tbl])
    assert all(p.paragraph_format.keep_with_next is True and p.style.name=='Normal' for p in labels)
    assert all(p.paragraph_format.keep_with_next is True for c in table.rows[1].cells for p in c.paragraphs)
    assert all(p.paragraph_format.keep_with_next is None for p in table.cell(0,0).paragraphs)


def test_rotated_picture_retains_inner_extent_ratio_when_fitted():
    from io import BytesIO
    from PIL import Image
    from docx.shared import Inches
    stream=BytesIO();Image.new('RGB',(400,200),'white').save(stream,format='PNG');stream.seek(0)
    doc=Document();picture=doc.add_picture(stream,width=Inches(6),height=Inches(3))
    transform=picture._inline.find('.//'+qn('a:xfrm'));transform.set('rot','5400000')
    # 90度旋转图：外框与内部坐标的宽高互换，缩放不能覆盖内部比例。
    picture.width=Inches(3);picture.height=Inches(6)
    inner=transform.find(qn('a:ext'));inner.set('cx',str(Inches(6)));inner.set('cy',str(Inches(3)))
    doc.sections[0].page_width=Inches(4)
    document_format.fit_inline_images(doc,doc.element.body)
    assert int(inner.get('cx'))==2*int(inner.get('cy'))
    assert picture.height==2*picture.width
    assert int(inner.get('cx'))==picture.height and int(inner.get('cy'))==picture.width
    assert transform.get('rot')=='5400000'


def test_material_index_matches_whole_punctuated_heading_before_grouping():
    from app.services import page_references,document_pagination
    d=Document();table=d.add_table(rows=2,cols=2)
    table.cell(0,0).text='资料名称';table.cell(0,1).text='页码范围'
    table.cell(1,0).text='集成测试、联调验收';table.cell(1,1).text='99-100'
    h=d.add_heading('集成测试、联调验收',1);d.add_paragraph('实际完整章节正文')
    document_pagination.bookmark(h._p,'_TenderHeading1',1)
    refs,unresolved=page_references.prepare(d,[{'name':'_TenderHeading1','title':'1 集成测试、联调验收','depth':1}])
    assert len(refs)==1 and not unresolved
    assert refs[0]['original']=='99-100' and '99-100' not in table.cell(1,1).text
    assert 'PAGEREF TenderRefStart' in table.cell(1,1)._tc.xml
