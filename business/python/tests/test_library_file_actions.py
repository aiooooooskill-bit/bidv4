"""Replacement/delete must change files, facts and retrieval as one lifecycle."""
import asyncio
from io import BytesIO
import pytest
from docx import Document
from app import db, company_store as cs, document
from app.knowledge import store, ingest, materials_lib as mt, fts
from app.services import knowledge, materials, templates, company_library
from app.services.errors import IncomingFile, BusinessError
from app.services.library_files import mutation_lock
from app.tools import docx_ops

@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setenv('DATA_DIR', str(tmp_path / 'data'))
    db.init_db()
    co = asyncio.run(company_library.save_company(company_library.CompanyBody(name='测试公司')))
    return tmp_path, co


def docx(text):
    d = Document()
    d.add_heading('实施方案', 1)
    d.add_paragraph(text)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def upload(tmp_path, fn, name, data, **kw):
    src = tmp_path / 'incoming'
    src.write_bytes(data)
    file = IncomingFile(src, name)
    try:
        return asyncio.run(fn(file=file, **kw))
    finally:
        asyncio.run(file.close())


def kb(library):
    root, co = library
    result = upload(root, knowledge.upload_file, '旧资料.docx', docx('旧事实：银行账号 OLDSECRET'), company_id=co['id'])
    kid = result['id']
    document.prepare(store.kb_files_dir()/result['file_name'], store.kb_parse_dir(result['file_name']))
    ingest.parsed(kid)
    ingest.finish(kid, {'doc_type': 'other', 'fields': {}})
    current = db.kb_get_item(kid)
    asyncio.run(company_library.review_asset(kid, company_library.ReviewBody(revision=current['revision'], company_id=co['id'], title='固定资料名称', business={'doc_type': 'other', 'statement': '人工核对 OLDSECRET 的旧事实说明，必须随新原件失效'}, text='人工核对旧全文 OLDSECRET')))
    return db.kb_get_item(kid)


def hits(word):
    return db.kb_search_segments(fts.build_match_expr(word), limit=100)


def test_company_replacement_runs_from_new_original_and_removes_old_facts(library):
    root, co = library
    old = kb(library)
    assert hits('OLDSECRET')
    result = upload(root, knowledge.replace_file, '新资料.docx', docx('新事实：银行账号 NEWFACT'), kid=old['id'], revision=old['revision'])
    item = db.kb_get_item(old['id'])
    assert item['id'] == old['id'] and item['company_id'] == co['id']
    assert item['title'] == old['title'] and item['created_at'] == old['created_at']
    assert item['revision'] > old['revision'] and item['review_status'] == 'pending_review'
    assert item['parse_status'] == 'pending' and item['corrected_text'] is None
    assert item['business_metadata'] is None and item['suggested_metadata'] is None
    assert not hits('OLDSECRET')
    assert not (store.kb_files_dir()/old['file_name']).exists()
    assert not store.kb_parse_dir(old['file_name']).exists()
    assert asyncio.run(knowledge.get_item_raw(item['id'])).path == str(store.kb_files_dir()/result['file_name'])
    parsed = document.prepare(store.kb_files_dir()/item['file_name'], store.kb_parse_dir(item['file_name']))
    assert not parsed.get('cached')
    ingest.parsed(item['id'])
    ingest.finish(item['id'], {'doc_type': 'other', 'fields': {}})
    assert hits('NEWFACT') and not hits('OLDSECRET')
    latest = db.kb_get_item(item['id'])
    asyncio.run(knowledge.delete_item(item['id'], latest['revision']))
    assert db.kb_get_item(item['id']) is None and not hits('NEWFACT')
    assert cs.asset_state(item['id'])['revision'] == 0
    assert not (store.kb_files_dir()/latest['file_name']).exists()
    assert not store.kb_parse_dir(latest['file_name']).exists()


def test_full_recognition_discards_cache_and_human_review(library):
    old = kb(library)
    asyncio.run(knowledge.retrigger(old['id'], old['revision']))
    item = db.kb_get_item(old['id'])
    assert item['parse_status'] == 'pending' and item['corrected_text'] is None
    assert item['review_status'] == 'pending_review' and item['company_id'] == old['company_id']
    assert not store.kb_parse_dir(old['file_name']).exists() and not hits('OLDSECRET')
    prepared = document.prepare(store.kb_files_dir()/old['file_name'], store.kb_parse_dir(old['file_name']))
    assert not prepared.get('cached')


def test_rejected_company_replacements_keep_existing_file_review_and_index(library):
    root, _ = library
    old = kb(library)
    before = (store.kb_files_dir()/old['file_name']).read_bytes()
    for name, data, revision in [('损坏.docx', b'broken', old['revision']), ('错误.txt', b'text', old['revision']), ('新.docx', docx('版本冲突'), old['revision']-1)]:
        with pytest.raises(BusinessError):
            upload(root, knowledge.replace_file, name, data, kid=old['id'], revision=revision)
        assert db.kb_get_item(old['id']) == old
        assert (store.kb_files_dir()/old['file_name']).read_bytes() == before
        assert hits('OLDSECRET')
    assert not list(store.kb_files_dir().glob('.replacement-*'))
    db.kb_update_item(old['id'], extract_status='running')
    latest = db.kb_get_item(old['id'])
    for fn in (knowledge.delete_item, knowledge.retrigger):
        with pytest.raises(BusinessError, match='正在识别'):
            asyncio.run(fn(old['id'], latest['revision']))


def test_duplicate_company_upload_does_not_replace_target(library):
    root, co = library
    old = kb(library)
    data = docx('另一份已上传的资料')
    upload(root, knowledge.upload_file, '另一份.docx', data, company_id=co['id'])
    with pytest.raises(BusinessError, match='相同内容'):
        upload(root, knowledge.replace_file, '替换.docx', data, kid=old['id'], revision=old['revision'])
    assert db.kb_get_item(old['id']) == old and hits('OLDSECRET')


def test_material_replacement_requires_new_chapters_and_rebuilds_retrieval(library):
    root, _ = library
    first = upload(root, materials.upload_file, '旧方案.docx', docx('旧素材 OLDWRITING'))
    fid = first['id']
    outline = mt.read_outline(fid)[0]
    block = asyncio.run(materials.create_block(fid, {'title': '旧素材', 'ranges': [[outline['start_line'], outline['end_line']]]}))
    old = db.mt_get_file(fid)
    assert hits('OLDWRITING')
    result = upload(root, materials.replace_file, '新方案.docx', docx('新素材 NEWWRITING'), fid=fid, file_hash=old['file_hash'])
    current = db.mt_get_file(fid)
    assert current['id'] == fid and current['parse_status'] == 'ready'
    assert current['file_hash'] != old['file_hash']
    assert not mt.mt_parse_dir(old['file_name']).exists()
    assert not (mt.mt_files_dir()/old['file_name']).exists()
    assert db.mt_get_block(block['id']) is None and not hits('OLDWRITING')
    assert asyncio.run(materials.list_items(file_id=fid))['items'][0]['kind'] == 'file'
    with pytest.raises(BusinessError, match='已被更换'):
        asyncio.run(materials.delete_file(fid, old['file_hash']))
    new_outline = mt.read_outline(fid)[0]
    new_block = asyncio.run(materials.create_block(fid, {'title': '新素材', 'ranges': [[new_outline['start_line'], new_outline['end_line']]]}))
    assert hits('NEWWRITING')
    asyncio.run(materials.delete_block(new_block['id']))
    assert db.mt_get_file(fid) and not hits('NEWWRITING')
    asyncio.run(materials.create_block(fid, {'title': '新素材二', 'ranges': [[new_outline['start_line'], new_outline['end_line']]]}))
    asyncio.run(materials.delete_file(fid, current['file_hash']))
    assert db.mt_get_file(fid) is None and not db.mt_list_blocks(fid) and not hits('NEWWRITING')
    assert not (mt.mt_files_dir()/result['file_name']).exists()


def test_material_mutations_cannot_overlap_parser_or_replacement(library):
    root, _ = library
    result = upload(root, materials.upload_file, '方案.docx', docx('原方案'))
    fid = result['id']
    with mutation_lock('mt', fid):
        with pytest.raises(BusinessError, match='正在处理'):
            asyncio.run(materials.reparse_file(fid))
        with pytest.raises(BusinessError, match='正在处理'):
            asyncio.run(materials.delete_file(fid, db.mt_get_file(fid)['file_hash']))
    assert db.mt_get_file(fid)['parse_status'] == 'ready'


def test_template_replace_keeps_key_and_default_and_delete_resets_default(library):
    root, _ = library
    first = upload(root, templates.upload_template, '自定义.docx', docx('旧版式'))
    asyncio.run(templates.activate_template(first.key))
    old = next(t for t in asyncio.run(templates.list_templates()) if t.key == first.key)
    upload(root, templates.replace_template, '不同文件名.docx', docx('新版式'), key=old.key, version=old.version)
    new = next(t for t in asyncio.run(templates.list_templates()) if t.key == old.key)
    assert new.version != old.version and new.active
    assert '新版式' in '\n'.join(p.text for p in Document(docx_ops.templates_dir()/old.key).paragraphs)
    with pytest.raises(BusinessError, match='已更新'):
        asyncio.run(templates.delete_template(old.key, old.version))
    asyncio.run(templates.delete_template(new.key, new.version))
    assert not (docx_ops.templates_dir()/new.key).exists()
    assert docx_ops.active_template_name() is None
    with pytest.raises(BusinessError, match='内置版式'):
        asyncio.run(templates.delete_template(docx_ops.BUILTIN_TEMPLATE_KEY, ''))
