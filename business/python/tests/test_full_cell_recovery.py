import pytest
from docx import Document
from app.tools.docx_ops import docx_section_read, docx_section_revise
from tests.test_section_write_atomicity import env

def make_doc(root, text):
 p=root/'body/实施方案.docx';d=Document();d.add_paragraph(text);d.save(p);return p


def table(root):
 p=make_doc(root,'真实响应表');d=Document(p);t=d.add_table(rows=2,cols=2);t.cell(0,0).text='采购包';t.cell(0,1).text='要求';t.cell(1,0).text='包3';t.cell(1,1).text='故障处理：安排相应技术人员共同处理故障';d.save(p);return p


def test_wrong_cell_returns_actual_content_and_exact_existing_locations(env):
 _,root,_=env;p=table(root);before=p.read_bytes();out=docx_section_revise('body/实施方案.docx',[{'table':1,'row':1,'col':1,'action':'replace','find':'故障处理','text':'故障处理'}])
 assert out.startswith('[修订失败]') and '选定单元格T1R1C1当前原文' in out and '采购包' in out
 assert '精确匹配位置：T1R2C2' in out;assert p.read_bytes()==before


def test_selected_cell_read_exposes_complete_long_text_without_changing_file(env):
 _,root,_=env;p=table(root);d=Document(p);text='完整开头'+('真实采购要求与字段'*40)+'完整尾部';d.tables[0].cell(1,1).text=text;d.save(p);before=p.read_bytes()
 assert '…（截断）' in docx_section_read('body/实施方案.docx')
 out=docx_section_read('body/实施方案.docx',table=1,row=2,col=2)
 assert text in out and 'T1R2C2' in out and '…（截断）' not in out;assert p.read_bytes()==before


def test_failed_batch_witness_uses_unmodified_original_not_partial_memory(env):
 _,root,_=env;p=table(root);before=p.read_bytes();out=docx_section_revise('body/实施方案.docx',[{'table':1,'row':2,'col':2,'action':'replace','find':'故障处理','text':'新段落值'},{'table':1,'row':1,'col':1,'action':'replace','find':'新段落值','text':'不应写入'}])
 assert out.startswith('[修订失败]') and '精确匹配位置：原稿中未找到' in out
 assert '精确匹配位置：T1R2C2' not in out;assert p.read_bytes()==before


def test_exact_match_witness_does_not_join_separate_cell_paragraphs(env):
 _,root,_=env;p=table(root);d=Document(p);d.tables[0].cell(1,1).text='故障';d.tables[0].cell(1,1).add_paragraph('处理');d.save(p)
 out=docx_section_revise('body/实施方案.docx',[{'table':1,'row':1,'col':1,'action':'replace','find':'故障处理','text':'不应写入'}]);assert '精确匹配位置：原稿中未找到' in out


@pytest.mark.parametrize('kwargs',[{'table':1},{'table':1,'row':99,'col':1}])
def test_partial_or_out_of_range_cell_read_is_rejected(env,kwargs):
 _,root,_=env;table(root);assert docx_section_read('body/实施方案.docx',**kwargs).startswith('[读取失败]')


def test_exported_tool_schema_preserves_positive_indices_and_nonempty_edits(tmp_path,monkeypatch):
 import json,subprocess,sys
 from pathlib import Path
 monkeypatch.setenv('DATA_DIR',str(tmp_path))
 payload={'operation':'schemas','data_root':str(tmp_path)}
 x=subprocess.run([sys.executable,str(Path(__file__).parents[1]/'runner.py')],input=json.dumps(payload),text=True,capture_output=True,timeout=15);assert x.returncode==0,x.stdout
 schemas=json.loads(x.stdout)['result'];read=next(s for s in schemas if s['name']=='docx_section_read')['parameters']
 assert read['properties']['table']['anyOf'][0]['exclusiveMinimum']==0
 assert read['properties']['para']['anyOf'][0]['exclusiveMinimum']==0
 revise=next(s for s in schemas if s['name']=='docx_section_revise')['parameters'];assert revise['properties']['edits']['minItems']==1
