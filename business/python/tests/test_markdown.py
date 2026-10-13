"""跨端 Markdown 契约与原文行号，覆盖业务实际依赖的语法边界。"""
import json
from pathlib import Path

import pytest

from app.markdown import headings, iter_tables, list_items, markdown_text
from app.parse import outline_with_lines
from app.knowledge.segmenter import segments_from
from app.tools.assemble_tender import _md_to_tree, _tree_skipped_lines

CASES = json.loads((Path(__file__).parents[3] / 'tests/fixtures/markdown-contract.json').read_text())


@pytest.mark.parametrize('case', CASES, ids=lambda case: case['name'])
def test_shared_markdown_contract(case):
    assert [[h.line, h.level, h.title] for h in headings(case['content'])] == case['headings']
    actual = [[line, header, [[n, cells] for n, cells in rows]]
              for line, header, rows in iter_tables(case['content'])]
    assert actual == case['tables']


def test_outline_and_search_segments_keep_original_lines():
    content = '# **总章**\n\n~~~md\n## 伪章节\n~~~\n\n子章\n-----\n正文\n\n# 末章\n末章正文\n'
    outline = outline_with_lines(content)
    assert [(n['标题'], n['start_line'], n['end_line']) for n in outline] == [('总章', 1, 10), ('末章', 11, 12)]
    assert outline[0]['children'][0] == {'标题': '子章', 'level': 2, 'start_line': 7, 'end_line': 10, 'children': []}
    segments = segments_from(content, outline)
    assert [(s['section_path'], s['line_start'], s['line_end']) for s in segments] == [('总章', 1, 10), ('总章 / 子章', 7, 10), ('末章', 11, 12)]
    assert segments[1]['raw'] == '子章\n-----\n正文\n'


def test_directory_uses_real_list_depth_and_ignores_fenced_and_ordered_items():
    content = '- **父章**\n    + 子章\n        * 孙章\n\n~~~md\n- 示例假章\n~~~\n\n1. 不合目录契约\n\n> - 引用示例\n'
    tree = _md_to_tree(content)
    assert len(tree) == 1 and tree[0]['目录名称'] == '父章'
    child = tree[0]['children'][0]
    assert child['目录名称'] == '子章' and child['level'] == 2
    assert child['children'][0]['目录名称'] == '孙章' and child['children'][0]['level'] == 3
    skipped = _tree_skipped_lines(content)
    assert '- 示例假章' in skipped and '1. 不合目录契约' in skipped and '> - 引用示例' in skipped
    assert '+ 子章' not in skipped
    assert [(item.line, item.depth) for item in list_items(content) if not item.ordered] == [(1, 1), (2, 2), (3, 3)]


def test_semantic_text_for_material_comparison():
    content = '<!-- 内部元信息 -->\n# **服务保障**\n\n| 项 | 值 |\n|---|---|\n| 维护 | [三年](https://example.com) |\n'
    assert markdown_text(content) == '服务保障\n项\n值\n维护\n三年'


def test_docx_and_pdf_generated_cells_roundtrip_literal_pipe_and_backslash(tmp_path):
    from docx import Document
    from app.parse import convert
    from app.parse.pdf import _pipe_table

    value = r'电话支持|7×24，路径 C:\service'
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '事项', '值'
    table.cell(1, 0).text, table.cell(1, 1).text = '服务', value
    path = tmp_path / '服务.docx'
    document.save(path)
    for content in (convert(path).md, _pipe_table([['事项', '值'], ['服务', value]])):
        assert list(iter_tables(content))[0][2][0][1] == ['服务', value]


def test_knowledge_and_material_excerpts_read_semantic_text_without_metadata(tmp_path, monkeypatch):
    from app.tools.search_knowledge import _excerpt, _mt_block_excerpt
    from app.knowledge import store, materials_lib
    from app import company_store
    monkeypatch.setattr(company_store,'asset_state',lambda _: {'corrected_text':None})

    path = tmp_path / '资料.md'
    path.write_text('<!-- 内部元信息 -->\n# **[服务保障][service]**\n\n| 项 | 值 |\n|---|---|\n| 支持 | A\\|B，7×24 |\n\n[service]: https://example.com/service\n')
    monkeypatch.setattr(store, 'kb_parse_paths', lambda _: (path, None, None, None))
    monkeypatch.setattr(materials_lib, 'mt_parse_paths', lambda _: (path, None, None))
    for text in (_excerpt({'id':'fixture','file_name': '资料.docx'}, 1, 6), _mt_block_excerpt('资料.docx', [[1, 6]])):
        assert '服务保障' in text and 'A|B，7×24' in text
        assert '内部元信息' not in text and '**' not in text and '---' not in text and '[service]' not in text


def test_pdf_heading_recovery_does_not_rewrite_markdown_blocks():
    from app.parse.pdf import _apply_numbered
    content = '# 已有标题\n正文\n\n~~~\n第一章 假章\n内容\n~~~\n\n项目 | 值\n---|---\n第一章 伪表 | 值\n\n> 第二章 引用\n\n第一章 总述\n正文详细说明。\n\n第二章 实施\n正文详细。\n\n第三章 服务\n正文详细。\n'
    marks = _apply_numbered(content.splitlines(), set())
    assert [text for _, text in marks.values()] == ['# 第一章 总述', '# 第二章 实施', '# 第三章 服务']
