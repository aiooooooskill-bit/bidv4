"""Safety regressions for scoped freshness, complete annotations and dispatch binding."""
import json
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app import artifact_store, runctx
from app.services import page_flow, build_state, review, source_usage
from app.services import review_batches
from tests.util import init_env
from tests.test_dispatch_enrich import env, _seed
from tests.test_review_flow import setup, fixture_candidate, tender_citations


def test_unselected_changes_do_not_expire_selected_inputs(tmp_path, monkeypatch):
    task, _ = init_env(tmp_path, monkeypatch); tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    root = artifact_store.sources_dir(tid)
    for name in ('main.txt', 'supplement.txt', 'unused.txt'):
        (root / name).write_text(name * 20)
    page_flow.set_sources(tid, 'main.txt', [])
    initial = page_flow.input_versions(tid)
    stamps = {k: build_state.preparation_stamp(tid, k) for k in ('guide', 'promises')}
    (root / 'unused.txt').write_text('changed unrelated upload')
    (root / 'new-unused.txt').write_text('new unrelated upload')
    page_flow.set_sources(tid, 'main.txt', [])  # excluded inventory itself is not a dependency.
    assert page_flow.input_versions(tid) == initial
    assert {k: build_state.preparation_stamp(tid, k) for k in stamps} == stamps
    (root / 'main.txt').write_text('changed actual tender')
    assert all(v != initial[k] for k, v in page_flow.input_versions(tid).items())
    before_selection = page_flow.input_versions(tid)
    page_flow.set_sources(tid, 'main.txt', ['supplement.txt'])
    assert all(v != before_selection[k] for k, v in page_flow.input_versions(tid).items())


def test_consumed_reference_changes_still_expire_only_consuming_stages(tmp_path, monkeypatch):
    task, _ = init_env(tmp_path, monkeypatch); tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    root = artifact_store.sources_dir(tid)
    (root / 'main.txt').write_text('selected tender')
    (root / 'reference.txt').write_text('consumed reference')
    page_flow.set_sources(tid, 'main.txt', [])
    source_usage.record_read(tid, 'writer', source_usage.access(tid, root / 'reference.txt'))
    value = build_state.read(tid)
    source_usage.attach(tid, value, 'guide', 'writer', True); build_state.save(tid, value)
    initial = page_flow.input_versions(tid)
    (root / 'reference.txt').write_text('new reference version')
    current = page_flow.input_versions(tid)
    assert all(current[k] == initial[k] for k in ('analyse', 'outline'))
    assert all(current[k] != initial[k] for k in ('guide', 'promises', 'body', 'revise', 'build', 'review'))
    with pytest.raises(ValueError, match='参考来源'):
        source_usage.attach(tid, build_state.read(tid), 'guide', 'writer', False)


def test_deleting_selected_source_keeps_state_readable_and_invalidates_outputs(tmp_path, monkeypatch):
    import asyncio
    import os
    import subprocess
    import sys
    from app.services import files
    task, _ = init_env(tmp_path, monkeypatch); tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    (artifact_store.sources_dir(tid) / 'main.txt').write_text('Selected original')
    page_flow.set_sources(tid, 'main.txt', [])
    before = page_flow.input_versions(tid)
    asyncio.run(files.delete_file('main.txt', tid))
    current = page_flow.input_versions(tid)
    assert all(current[k] != before[k] for k in current)
    assert page_flow.source_basis(tid)['sources'] == {'main.txt': None}
    # Exercise the actual page-state runner, not just the version helper.
    result = subprocess.run([sys.executable, str(Path(__file__).resolve().parents[1] / 'runner.py')],
                            input=json.dumps({'operation': 'state', 'args': {'task_id': tid}, 'data_root': os.environ['DATA_DIR']}),
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr + result.stdout
    assert json.loads(result.stdout)['result']['input_versions'] == current


def test_reference_dependencies_inherit_and_replace_without_losing_used_versions(tmp_path, monkeypatch):
    task, _ = init_env(tmp_path, monkeypatch); tid = task['id']
    artifact_store.ensure_task_skeleton(tid)
    sources = artifact_store.sources_dir(tid); (sources / 'reference.txt').write_text('reference original')
    parsed = artifact_store.work_dir(tid) / 'parse/reference.txt/reference.txt.md'
    parsed.parent.mkdir(parents=True); parsed.write_text('reference original')
    source = source_usage.access(tid, parsed)
    assert source['name'] == 'reference.txt'
    source_usage.record_read(tid, 'parent', source)
    value = build_state.read(tid)
    source_usage.inherit(value, 'child', 'parent', True)
    source_usage.attach(tid, value, 'body/章.docx', 'child', True)
    build_state.save(tid, value)
    assert source_usage.basis(tid, ['body/章.docx'])['body/章.docx']['reference.txt'] == source['sha256']
    (sources / 'reference.txt').unlink()
    assert source_usage.basis(tid, ['body/章.docx'])['body/章.docx']['reference.txt'] is None
    with pytest.raises(ValueError, match='参考来源'):
        source_usage.attach(tid, value, 'body/章.docx', 'child', False)
    source_usage.attach(tid, value, 'body/章.docx', 'new-writer', True)
    assert 'body/章.docx' not in value['source_dependencies']


def test_dispatch_canonical_paths_and_unknown_section_fail_closed(env):
    from app.dispatch_enrich import body_dispatch
    tid = _seed(env)
    result = body_dispatch('写 3.1 项目理解、3.2 总体设计\n保留现有方法与要求。', tid)
    assert result['section_paths'] == ['body/3.1 项目理解与需求分析.docx', 'body/3.2 总体设计方案.docx']
    with pytest.raises(Exception, match='未知或多义'):
        body_dispatch('写 3.1 项目理解、某个完全不存在的章节ZZZ', tid)


def test_multisection_dispatch_preserves_punctuation_inside_full_titles(env):
    from app.dispatch_enrich import body_dispatch
    titles = ['封面', '附件1、法定代表人或单位负责人身份证明复印件', '附件10、供应商相关专利说明', '数据范围、更新频率与数据质量保障']
    content = {'response_documents': [{'name': '响应文件', 'directory': [{'目录名称': title, 'level': 1, 'children': [], '交付形态': '正文编写'} for title in titles]}], 'registry': {}}
    tid = _seed(env, content=content, guide=False, promise=False)
    result = body_dispatch('写 ' + '、'.join(titles), tid)
    assert result['section_paths'] == ['body/' + title + '.docx' for title in titles]
    assert all(path in result['description'] for path in result['section_paths'])
    with pytest.raises(Exception, match='未知或多义'):
        body_dispatch('写 ' + titles[1] + '、根本不存在的节ZZZ', tid)


def test_full_annotation_evidence_keeps_body_lines_and_rejects_printed_commitment(setup):
    tid, value, action_path, record = setup
    source = artifact_store.sources_dir(tid) / '投标.docx'
    old_body = Path(review.documents(tid)[1]['path']).read_bytes()
    doc = Document(source)
    long_text = '请核对尚未确认的事实。' * 35 + '\n尾部关键证据：不能以季度付款推定成本已经回收。'
    doc.add_comment(doc.paragraphs[1].runs[0], long_text, author='合成测试')
    table = doc.add_table(rows=1, cols=1); paragraph = table.cell(0, 0).paragraphs[0]; run = paragraph.add_run('表格锚定原文')
    doc.add_comment(run, '格内完整批注尾部', author='合成测试')
    orphan = doc.add_comment(doc.paragraphs[0].runs[0], '无锚点批注也须保留', author='合成测试')
    for mark in list(doc.element.body.iter()):
        if mark.tag in {qn('w:commentRangeStart'), qn('w:commentRangeEnd'), qn('w:commentReference')} and mark.get(qn('w:id')) == str(orphan.comment_id):
            mark.getparent().remove(mark)
    doc.save(source)
    review.prepare(tid)
    # The new table adds body text; pre-existing paragraphs retain the same prefix/lines.
    primary = next(d for d in review.documents(tid) if d['kind'] == 'bid')
    assert Path(primary['path']).read_bytes().startswith(old_body)
    annotation = primary['annotations']; text = Path(annotation['path']).read_text()
    assert long_text in text and '格内完整批注尾部' in text and '无锚点批注也须保留' in text
    assert '表格锚定原文' in text and '锚点：未定位' in text
    assert len(annotation['comments']) == 3
    prepared = review.prepared_documents(tid)
    assert prepared['annotation_evidence'][0]['text'] == text
    assert 'text' not in primary['annotations']  # no full-text duplication in every state/batch.
    citation = {'document_id': annotation['id'], 'start': 1, 'end': annotation['lines']}
    evidence = review.resolve_evidence(review.evidence_documents(tid), tender_citations(tid), [citation], tender_citations(tid))['bid'][0]
    assert evidence['quote'] == text and evidence['evidence_type'] == 'annotation'
    record['input_version'] = review.basis(tid); action_path.write_text(json.dumps(record))
    batch = review_batches.next_batch(tid, 'review-test')
    row = {'slot': 1, 'tender': tender_citations(tid), 'status': 'responded', 'assessment': '批注不得代替承诺。', 'bid': [citation], 'problems': []}
    with pytest.raises(ValueError, match='Word批注不能代替'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], [row])
    # Mixing one printed paragraph with the annotation must not make the latter
    # eligible as an established response fact.
    mixed = {**row, 'bid': [*value['requirements'][0]['bid'], citation]}
    with pytest.raises(ValueError, match='Word批注不能代替'):
        fixture_candidate(tid, 'review-test', batch['batch_id'], [mixed])
    row.update(status='uncertain', problems=[{'risk_type': 'confirmation', 'title': '资料待确认', 'explanation': '完整批注记录了资料缺口。', 'suggestion': '核对后补齐正文。'}])
    proposed = fixture_candidate(tid, 'review-test', batch['batch_id'], [row])
    assert proposed['comparisons'][0]['bid'][0]['quote'] == text
