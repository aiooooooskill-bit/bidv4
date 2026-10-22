import json,hashlib
from copy import deepcopy
import pytest
from docx import Document
from app import config,db,artifact_store
from app.services import build_state,page_flow
from tests.test_revision_revalidation import env, stale, action, rpc


def old_ready(env):
 tid,root,_=env;p=stale(env);d=Document(p);d.add_paragraph('工程故障注入：【待补】');d.save(p)
 # Explicit fixture of a prior checker accepting this exact document.
 build_state.record_write(tid,'body/实施方案.docx','parent');build_state.record_validation(tid,'body/实施方案.docx')
 assert build_state.snapshot(tid)['body_sections'][0]['status']=='ready'
 return p


def test_failed_current_check_demotes_prior_ready_without_rewriting_original_and_is_idempotent(env):
 p=old_ready(env);before=p.read_bytes();action(env);code,out=rpc(env)
 assert code==0 and out['result'].startswith('[校验未通过]'),out
 assert p.read_bytes()==before;assert build_state.snapshot(env[0])['body_sections'][0]['status']=='draft'
 ledger=build_state.read(env[0]);pending=ledger['pending_sections']['body/实施方案.docx'];assert pending['hash']==hashlib.sha256(before).hexdigest()
 data=(env[1]/'body/.build-state.json').read_bytes();rpc(env);assert (env[1]/'body/.build-state.json').read_bytes()==data

@pytest.mark.parametrize('kind,paths',[('body',['body/实施方案.docx']),('build',[])])
def test_active_body_and_full_build_can_reject_old_ready_validation(env,kind,paths):
 p=old_ready(env);before=p.read_bytes();action(env,kind=kind,section_paths=paths);rpc(env)
 assert build_state.snapshot(env[0])['body_sections'][0]['status']=='draft';assert p.read_bytes()==before

@pytest.mark.parametrize('metadata',[None,{'stopped_at':1},{'outcome':{'state':'completed'}}])
def test_unscoped_stopped_or_closed_read_only_check_does_not_change_registration(env,metadata):
 p=old_ready(env)
 if metadata is not None:action(env,**metadata)
 before=(env[1]/'body/.build-state.json').read_bytes();rpc(env)
 assert (env[1]/'body/.build-state.json').read_bytes()==before


def test_unassigned_child_cannot_demote_registration(env):
 old_ready(env);action(env);before=(env[1]/'body/.build-state.json').read_bytes();rpc(env,'unassigned-child');assert (env[1]/'body/.build-state.json').read_bytes()==before

@pytest.mark.parametrize('race',['stop','file'])
def test_failed_result_cannot_demote_a_changed_action_or_document(env,race):
 from app.services import section_revalidation as s
 p=old_ready(env);action(env);captured=s.capture(env[0],'parent','body/实施方案.docx');before=(env[1]/'body/.build-state.json').read_bytes()
 if race=='stop':action(env,stopped_at=1)
 else:d=Document(p);d.add_paragraph('其他修改');d.save(p)
 with pytest.raises(ValueError):s.reject(env[0],'parent','body/实施方案.docx',captured)
 assert (env[1]/'body/.build-state.json').read_bytes()==before


def test_rejected_section_makes_existing_published_volume_stale(env,monkeypatch):
 p=old_ready(env);tid,root,_=env;doc=page_flow.directory(tid)['content']['response_documents'][0];name=doc['name'];volume=root/('body/整本-'+name+'.docx');volume.write_bytes(p.read_bytes());row={'artifact_id':'old-published','display_name':name,'kind':'tender.volume','content_seq':1}
 meta={'filename':volume.name,'sha256':page_flow.file_hash(volume)}
 monkeypatch.setattr(db,'list_artifact_index',lambda **kw:[row]);monkeypatch.setattr(artifact_store,'read_content_resolved',lambda *a:json.dumps(meta));monkeypatch.setattr(artifact_store,'package_files',lambda *a:[volume])
 value=build_state.read(tid);value.setdefault('volumes',{})[row['artifact_id']]={'basis':build_state.volume_basis(tid,doc),'content_seq':1,'hash':page_flow.file_hash(volume)};build_state.save(tid,value)
 assert build_state.volume_states(tid,build_state.read(tid))[0]['status']=='ready';published=volume.read_bytes();action(env);rpc(env)
 assert build_state.volume_states(tid,build_state.read(tid))[0]['status']=='stale';assert volume.read_bytes()==published and p.read_bytes()==published
