import asyncio
import json
import pytest
from docx import Document
from app import db, company_store as cs, artifact_store, runctx
from app.knowledge import store, fts, ingest
from app.services import bidder, knowledge, company_library as lib, build_state, page_flow, materials
from app.services.errors import BusinessError
from app.tools.company_asset import read_company_asset
from tests.test_company_library import setup
from tests.test_build_state import env, complete_sections, publish_volume


def bind(tid, company, signer='legal'):
    old=bidder.row(tid)
    return asyncio.run(bidder.save_bidder(tid,bidder.BidderBody(company_id=company['id'],signer_id=signer,sign_date='2026-10-12',revision=old['revision'] if old else 0)))


def legal(company):
    return asyncio.run(lib.save_company(lib.CompanyBody(**{**company,'profile':{'legal_name':'验收法人'}})))


def review(company,item,meta,text=None):
    return asyncio.run(lib.review_asset(item['id'],lib.ReviewBody(revision=cs.asset_state(item['id'])['revision'],company_id=company['id'],title=item['title'],business=meta,text=text)))


def test_atomic_detail_cannot_mix_old_fields_with_new_revision(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch)
    old=review(company,item,meta,'原核对文字')
    real_get=db.kb_get_item
    def concurrent_read(kid):
        row=real_get(kid)
        cs.save_asset(kid,row['revision'],company['id'],'并发更新',{**meta,'extra':{'银行账号':{'value':'新账号'}}},'新核对文字')
        return row
    with monkeypatch.context() as patch:
        patch.setattr(db,'kb_get_item',concurrent_read)
        received=asyncio.run(knowledge.get_item(item['id']))
    assert received['revision']==old['revision']
    assert received['content']=='原核对文字'
    with pytest.raises(BusinessError,match='其他操作更新'):
        asyncio.run(lib.review_asset(item['id'],lib.ReviewBody(revision=received['revision'],company_id=company['id'],title=received['title'],business=received['business'],text=received['content'])))
    assert real_get(item['id'])['title']=='并发更新'


def test_corrected_ocr_outline_covers_new_tail_and_new_heading(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch)
    md,outline,_,_=store.kb_parse_paths(item['file_name'])
    md.write_text('# 旧标题\n旧文字')
    outline.write_text(json.dumps([{'标题':'旧标题','level':1,'start_line':1,'end_line':2,'children':[]}]))
    corrected='# 核对标题\n最新文字\n## 新增信息\nOCRTAIL2026'
    review(company,item,meta,corrected)
    content=asyncio.run(knowledge.get_item_content(item['id']))
    assert content['outline'][0]['end_line']==4
    assert content['outline'][0]['children'][0]['标题']=='新增信息'
    hits=db.kb_search_segments(fts.build_match_expr('OCRTAIL2026'),domain='kb',company_id=company['id'])
    assert hits and all(h['line_end']==4 for h in hits)
    assert json.loads(outline.read_text())[0]['标题']=='旧标题'


def test_recursive_library_ancestors_denied_but_project_and_global_materials_allowed(tmp_path,monkeypatch):
    from app.config import workspace_dir
    company,item,meta=setup(tmp_path,monkeypatch);company=legal(company)
    tid=db.create_task('主体范围验证')['id'];artifact_store.ensure_task_skeleton(tid);bind(tid,company)
    review(company,item,meta)
    for ancestor in (workspace_dir(),tmp_path,tmp_path.parent):
        with pytest.raises(BusinessError,match='上级目录'):cs.require_asset_path(tid,ancestor)
    cs.require_asset_path(tid,artifact_store.work_dir(tid))
    cs.require_asset_path(tid,workspace_dir()/'materials')
    assert cs.require_asset_path(tid,store.kb_files_dir()/item['file_name'])['id']==item['id']
    foreign=asyncio.run(lib.save_company(lib.CompanyBody(name='其他公司')))
    review(foreign,item,meta)
    with pytest.raises(BusinessError,match='不属于'):cs.require_asset_path(tid,store.kb_files_dir()/item['file_name'])


@pytest.mark.parametrize('invalid',['representative_missing','legal_name'])
def test_invalid_binding_state_readable_and_reconfirmation_recovers(tmp_path,monkeypatch,invalid):
    company,item,meta=setup(tmp_path,monkeypatch);company=legal(company)
    foreign=asyncio.run(lib.save_company(lib.CompanyBody(name='乙公司',profile={'legal_name':'乙法人'})))
    rep=asyncio.run(lib.save_representative(lib.RepresentativeBody(name='代表')))
    tid=db.create_task('绑定恢复')['id'];artifact_store.ensure_task_skeleton(tid)
    bind(tid,company,rep['id'] if invalid=='representative_missing' else 'legal')
    if invalid=='representative_missing':
        with cs.transaction() as connection:connection.execute('DELETE FROM representatives WHERE id=?',(rep['id'],))
    else:asyncio.run(lib.save_company(lib.CompanyBody(**{**company,'profile':{}})))
    assert bidder.status(tid)['stale']
    assert 'invalid' in bidder.basis(tid)['current']
    assert page_flow.input_versions(tid)['guide']
    assert build_state.snapshot(tid)['body_sections']==[]
    assert not bind(tid,foreign)['stale']


def test_bidder_change_invalidates_published_volume_and_keeps_old_file(env):
    tid,root,_=env
    a=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    b=asyncio.run(lib.save_company(lib.CompanyBody(name='乙公司',profile={'legal_name':'乙法人'})))
    bind(tid,a)
    build_state.record_preparation(tid,'guide');build_state.record_preparation(tid,'promises')
    complete_sections(tid,root);volume=publish_volume(tid,root,'技术标')
    file=artifact_store.package_files(volume['artifact_id'],volume)[0];before=file.read_bytes()
    assert build_state.snapshot(tid)['volume_states'][0]['status']=='ready'
    bind(tid,b)
    state=build_state.snapshot(tid)
    assert state['volume_states'][0]['status']=='stale'
    assert all(s['status']=='stale' for s in state['body_sections'])
    assert file.read_bytes()==before


def test_consumed_certificate_versions_invalidate_only_related_chapters_and_volume(env,monkeypatch):
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    item=db.kb_insert_item(file_name='资质证.png',file_hash='certificate',title='资质证',ext='.png')
    md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True);md.write_text('证号C0001');outline.write_text('[]')
    meta={'doc_type':'qualification_certificate','extra':{'证号':{'value':'C0001'}}}
    ingest.finish(item['id'],meta);review(company,item,meta)
    bind(tid,company);build_state.record_preparation(tid,'guide');build_state.record_preparation(tid,'promises');complete_sections(tid,root)
    runctx.set_run('chapter-session','read-call',tid)
    try:
        assert 'C0001' in read_company_asset(item['id'])
        build_state.record_write(tid,'body/实施方案.docx','chapter-session')
        build_state.record_validation(tid,'body/实施方案.docx')
        publish_volume(tid,root,'技术标')
        unrelated=db.kb_insert_item(file_name='未采用.png',file_hash='other',title='其他',ext='.png');cs.attach_asset(unrelated['id'],company['id'])
        assert not bidder.status(tid)['stale']
        assert all(s['status']=='ready' for s in build_state.snapshot(tid)['body_sections'])
        meta['extra']['证号']['value']='C0003';review(company,item,meta)
        assert bidder.status(tid)['stale']
        current=build_state.snapshot(tid)
        assert [s['status'] for s in current['body_sections']]==['stale','ready']
        assert current['volume_states'][0]['status']=='stale'
        assert current['build_preparation']=={'guide':'ready','promises':'ready'}
        bind(tid,company)
        assert not bidder.status(tid)['stale']
        assert [s['status'] for s in build_state.snapshot(tid)['body_sections']]==['stale','ready']
    finally:runctx.clear_run()


def test_directory_edit_does_not_invalidate_other_ready_chapter(env):
    tid,root,content=env;complete_sections(tid,root)
    next(node for node in content['response_documents'][0]['directory'] if node['目录名称']=='服务保障')['节点概述']='只更新服务保障'
    statuses=build_state.snapshot(tid)['body_sections']
    assert [s['status'] for s in statuses]==['ready','stale']


def test_asset_list_orders_documents_by_best_bm25_hit(tmp_path,monkeypatch):
    company,a,meta=setup(tmp_path,monkeypatch);cs.attach_asset(a['id'],company['id'])
    b=db.kb_insert_item(file_name='次优.png',file_hash='next',title='次优',ext='.png');cs.attach_asset(b['id'],company['id'])
    db.kb_replace_segments(a['id'],[{'body':' '.join(['needle']*20+['filler']*5)},{'body':' '.join(['needle']+['filler']*500)}])
    db.kb_replace_segments(b['id'],[{'body':' '.join(['needle']*5+['filler']*10)}])
    raw=db.kb_search_segments(fts.build_match_expr('needle'),domain='kb',company_id=company['id'])
    assert raw[0]['item_id']==a['id'] and raw[-1]['item_id']==a['id']
    assert [i['id'] for i in asyncio.run(knowledge.list_items(q='needle',company_id=company['id']))['items']]==[a['id'],b['id']]


def test_asset_changed_after_read_cannot_be_attached_as_current_output(env):
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    item=db.kb_insert_item(file_name='资质.png',file_hash='mid-run',title='资质',ext='.png')
    md,outline,_,_=store.kb_parse_paths(item['file_name']);md.parent.mkdir(parents=True);md.write_text('旧证号');outline.write_text('[]')
    meta={'doc_type':'qualification_certificate','extra':{'证号':{'value':'旧证号'}}}
    ingest.finish(item['id'],meta);review(company,item,meta);bind(tid,company)
    runctx.set_run('writing-session','read',tid)
    try:
        assert '旧证号' in read_company_asset(item['id'])
        meta['extra']['证号']['value']='新证号';review(company,item,meta)
        with pytest.raises(ValueError,match='制作期间'):
            build_state.record_write(tid,'body/实施方案.docx','writing-session')
        assert bidder.status(tid)['stale']
        assert build_state.snapshot(tid)['body_sections'][0]['status']=='draft'
        build_state.record_validation(tid,'body/实施方案.docx')
        assert build_state.snapshot(tid)['body_sections'][0]['status']=='draft'
    finally:runctx.clear_run()


def test_search_excerpt_consumption_records_certificate_dependency(env):
    from app.tools.search_knowledge import search_company_assets
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    item=db.kb_insert_item(file_name='许可.png',file_hash='excerpt',title='开户许可',ext='.png')
    meta={'doc_type':'qualification_certificate','extra':{'开户银行':{'value':'EXCERPTBANK'}}}
    ingest.finish(item['id'],meta);review(company,item,meta);bind(tid,company)
    runctx.set_run('search-session','search',tid)
    try:
        assert 'EXCERPTBANK' in search_company_assets('EXCERPTBANK')
        build_state.record_write(tid,'body/实施方案.docx','search-session')
        assert item['id'] in build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets']
    finally:runctx.clear_run()


def test_inserted_original_image_records_its_asset_version(env):
    from PIL import Image
    from app.tools.docx_ops import docx_image_insert
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    item=db.kb_insert_item(file_name='原件.png',file_hash='original-image',title='原件',ext='.png')
    cs.attach_asset(item['id'],company['id']);store.ensure_dirs();Image.new('RGB',(30,30),'white').save(store.kb_files_dir()/item['file_name'])
    bind(tid,company);runctx.set_run('image-session','insert',tid)
    try:
        result=docx_image_insert(dest='body/实施方案.docx',image='knowledge/files/原件.png')
        assert result.startswith('[已插图]'),result
        build_state.record_write(tid,'body/实施方案.docx','image-session')
        assert item['id'] in build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets']
    finally:runctx.clear_run()


def test_corrected_text_before_first_heading_is_also_searchable(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch)
    review(company,item,meta,'OCRPREFIX2026\n\n# 核对后标题\n正文')
    hits=db.kb_search_segments(fts.build_match_expr('OCRPREFIX2026'),domain='kb',company_id=company['id'])
    assert hits and hits[0]['line_start']==1


def test_replacing_chapter_without_old_company_assets_removes_old_dependencies(env):
    from app.services.asset_usage import record_read
    tid,root,_=env
    a=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    b=asyncio.run(lib.save_company(lib.CompanyBody(name='乙公司',profile={'legal_name':'乙法人'})))
    item=db.kb_insert_item(file_name='甲证.png',file_hash='old-a',title='甲证',ext='.png');cs.attach_asset(item['id'],a['id'])
    bind(tid,a);record_read(tid,'old-session',db.kb_get_item(item['id']))
    build_state.record_write(tid,'body/实施方案.docx','old-session');build_state.record_validation(tid,'body/实施方案.docx')
    bind(tid,b)
    doc=Document();doc.add_paragraph('乙公司全新正文，不采用甲公司原件');doc.save(root/'body/实施方案.docx')
    build_state.record_write(tid,'body/实施方案.docx','new-session',replace=True);build_state.record_validation(tid,'body/实施方案.docx')
    cs.attach_asset(item['id'],a['id'])
    assert not bidder.status(tid)['stale']
    assert build_state.snapshot(tid)['body_sections'][0]['status']=='ready'
    assert 'body/实施方案.docx' not in build_state.read(tid)['asset_dependencies']


def test_profile_adoption_rejects_candidate_updated_between_selection_and_save(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch);review(company,item,meta)
    candidate=asyncio.run(lib.candidates(company['id']))['items'][0]
    proposed={**company,'profile':{'bank_account':'旧账号','source-id:bank_account':item['id'],'source-revision:bank_account':str(candidate['revision'])}}
    meta['extra']['银行账号']['value']='新账号';review(company,item,meta)
    with pytest.raises(BusinessError,match='来源资料已更新'):
        asyncio.run(lib.save_company(lib.CompanyBody(**proposed)))
    assert cs.company(company['id'])['profile']==company['profile']
    proposed['profile']['bank_account']='新账号';proposed['profile']['source-revision:bank_account']=str(cs.asset_state(item['id'])['revision'])
    assert asyncio.run(lib.save_company(lib.CompanyBody(**proposed)))['profile']['bank_account']=='新账号'


def test_reconfirmation_cannot_label_old_profile_value_as_new_source_revision(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch);review(company,item,meta)
    company=asyncio.run(lib.save_company(lib.CompanyBody(**{**company,'profile':{'legal_name':'法人','bank_account':'旧账号','source-id:bank_account':item['id'],'source-revision:bank_account':str(cs.asset_state(item['id'])['revision'])}})))
    tid=db.create_task('来源确认')['id'];artifact_store.ensure_task_skeleton(tid);bind(tid,company)
    meta['extra']['银行账号']['value']='新账号';review(company,item,meta)
    with pytest.raises(BusinessError,match='重新核对取值'):bind(tid,company)
    assert bidder.status(tid)['stale'] and 'invalid' in bidder.basis(tid)['current']


def test_manual_text_without_machine_markdown_is_indexed(tmp_path,monkeypatch):
    company,item,meta=setup(tmp_path,monkeypatch);store.kb_parse_paths(item['file_name'])[0].unlink()
    review(company,item,meta,'MANUALWITHOUTOCR2026')
    assert db.kb_search_segments(fts.build_match_expr('MANUALWITHOUTOCR2026'),domain='kb',company_id=company['id'])
    company=legal(company)
    tid=db.create_task('手工全文检索')['id'];artifact_store.ensure_task_skeleton(tid);bind(tid,company)
    from app.tools.search_knowledge import search_company_assets
    runctx.set_run('manual-search','run',tid)
    try:
        result=search_company_assets('MANUALWITHOUTOCR2026')
        assert '摘录：MANUALWITHOUTOCR2026' in result
        assert item['id'] in build_state.read(tid)['asset_reads']['manual-search']['assets']
    finally:runctx.clear_run()


@pytest.mark.parametrize('reread',['none','other','updated'])
def test_partial_revision_retains_old_certificate_after_reconfirmation(env,reread):
    from app.services.asset_usage import record_read
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    a=db.kb_insert_item(file_name='甲证.png',file_hash='partial-a',title='甲证',ext='.png');cs.attach_asset(a['id'],company['id'])
    b=db.kb_insert_item(file_name='乙证.png',file_hash='partial-b',title='乙证',ext='.png');cs.attach_asset(b['id'],company['id'])
    bind(tid,company);record_read(tid,'original',db.kb_get_item(a['id']))
    build_state.record_write(tid,'body/实施方案.docx','original');build_state.record_validation(tid,'body/实施方案.docx')
    old_version=build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets'][a['id']]
    cs.attach_asset(a['id'],company['id']);bind(tid,company)
    if reread!='none':record_read(tid,'partial',db.kb_get_item(b['id'] if reread=='other' else a['id']))
    doc=Document(root/'body/实施方案.docx');doc.add_paragraph('只修改质量复核，旧证件内容仍然保留');doc.save(root/'body/实施方案.docx')
    with pytest.raises(ValueError,match='资料|旧资料版本'):
        build_state.record_write(tid,'body/实施方案.docx','partial',replace=False)
    build_state.record_validation(tid,'body/实施方案.docx')
    assert build_state.snapshot(tid)['body_sections'][0]['status']=='draft'
    assert build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets'][a['id']]==old_version
    # Only replacing the entire chapter can retire the old fact's dependency.
    bind(tid,company);record_read(tid,'replacement',db.kb_get_item(a['id']))
    doc=Document();doc.add_paragraph('完整重写，采用新证件');doc.save(root/'body/实施方案.docx')
    build_state.record_write(tid,'body/实施方案.docx','replacement',replace=True);build_state.record_validation(tid,'body/实施方案.docx')
    assert build_state.snapshot(tid)['body_sections'][0]['status']=='ready'
    assert build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets'][a['id']]['revision']==cs.asset_state(a['id'])['revision']


def test_partial_revision_cannot_keep_previous_company_certificate(env):
    from app.services.asset_usage import record_read
    tid,root,_=env
    a=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    b=asyncio.run(lib.save_company(lib.CompanyBody(name='乙公司',profile={'legal_name':'乙法人'})))
    item=db.kb_insert_item(file_name='甲证.png',file_hash='foreign-partial',title='甲证',ext='.png');cs.attach_asset(item['id'],a['id'])
    bind(tid,a);record_read(tid,'a-session',db.kb_get_item(item['id']));build_state.record_write(tid,'body/实施方案.docx','a-session')
    bind(tid,b)
    with pytest.raises(ValueError,match='资料发生变化'):build_state.record_write(tid,'body/实施方案.docx','b-session')
    build_state.record_validation(tid,'body/实施方案.docx')
    assert build_state.snapshot(tid)['body_sections'][0]['status']=='draft'


def test_delegated_facts_invalidate_child_chapter_and_returned_parent_output(env):
    from app.services.asset_usage import record_read,inherit_reads
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    a=db.kb_insert_item(file_name='甲证.png',file_hash='delegate-a',title='甲证',ext='.png');cs.attach_asset(a['id'],company['id'])
    b=db.kb_insert_item(file_name='乙证.png',file_hash='delegate-b',title='乙证',ext='.png');cs.attach_asset(b['id'],company['id'])
    bind(tid,company);build_state.record_preparation(tid,'guide');build_state.record_preparation(tid,'promises');complete_sections(tid,root)
    record_read(tid,'parent',db.kb_get_item(a['id']));inherit_reads(tid,'child','parent',seed=True)
    # A later parent read is not part of the already dispatched prompt.
    record_read(tid,'parent',db.kb_get_item(b['id']));inherit_reads(tid,'child','parent',seed=True)
    build_state.record_write(tid,'body/实施方案.docx','child',replace=True);build_state.record_validation(tid,'body/实施方案.docx')
    dependencies=build_state.read(tid)['asset_dependencies']['body/实施方案.docx']['assets']
    assert set(dependencies)=={a['id']}
    inherit_reads(tid,'consumer','child')
    build_state.record_write(tid,'body/服务保障.docx','consumer',replace=True);build_state.record_validation(tid,'body/服务保障.docx')
    publish_volume(tid,root,'技术标')
    cs.attach_asset(a['id'],company['id'])
    state=build_state.snapshot(tid)
    assert all(s['status']=='stale' for s in state['body_sections'])
    assert state['volume_states'][0]['status']=='stale'
    assert state['build_preparation']=={'guide':'ready','promises':'ready'}


def test_delegation_rejects_version_changed_before_child_starts(env):
    from app.services.asset_usage import record_read,inherit_reads
    tid,root,_=env
    company=asyncio.run(lib.save_company(lib.CompanyBody(name='甲公司',profile={'legal_name':'甲法人'})))
    item=db.kb_insert_item(file_name='甲证.png',file_hash='stale-delegate',title='甲证',ext='.png');cs.attach_asset(item['id'],company['id'])
    bind(tid,company);record_read(tid,'parent',db.kb_get_item(item['id']));cs.attach_asset(item['id'],company['id'])
    with pytest.raises(ValueError,match='派发期间'):inherit_reads(tid,'child','parent',seed=True)
    assert 'child' not in build_state.read(tid)['asset_reads']
