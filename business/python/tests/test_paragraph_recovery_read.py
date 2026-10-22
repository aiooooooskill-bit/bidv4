import pytest
from docx import Document
from app.tools.docx_ops import docx_section_read
from app.tools.validate_body import validate_body
from tests.test_section_write_atomicity import env


def long_doc(root):
 p=root/'body/实施方案.docx';d=Document()
 for i in range(1,522):d.add_paragraph('真实原稿第'+str(i)+'段')
 d.paragraphs[217].text='原稿完整开头'+('完整核验内容'*220)+'未申请专利注册'+('保留方法'*40)+'原稿完整末尾';d.save(p);return p,d.paragraphs[217].text

def test_point_paragraph_read_keeps_real_index_and_full_accepted_text_without_changes(env):
 p,expected=long_doc(env[1]);before=p.read_bytes();assert '原稿完整末尾' not in docx_section_read('body/实施方案.docx')
 out=docx_section_read('body/实施方案.docx',para=218)
 assert '[P218 完整段落]' in out and expected in out and '截断' not in out
 assert p.read_bytes()==before

@pytest.mark.parametrize('kwargs',[{'para':0},{'para':999},{'para':218,'table':1,'row':1,'col':1}])
def test_invalid_or_mixed_point_reads_reject_without_any_write(env,kwargs):
 p,_=long_doc(env[1]);before=p.read_bytes();assert docx_section_read('body/实施方案.docx',**kwargs).startswith('[读取失败]');assert p.read_bytes()==before

def test_validation_points_to_actual_failed_paragraph_read(env):
 _,root,_=env;p,_=long_doc(root);d=Document(p);d.paragraphs[217].text='工程读取故障注入：【待补】';d.save(p)
 out=validate_body('body/实施方案.docx');assert 'P218：' in out and 'docx_section_read' in out and 'para=218' in out


@pytest.mark.parametrize('bad',[{'para': True}, {'para': '7'}, {'para': 0}, {'table': 'T1','row':1,'col':1}, {'para':7,'extra':True}])
def test_point_read_rpc_rejects_invalid_types_before_access(env,bad):
 from tests.test_section_write_atomicity import request
 _,root,_=env;path=root/'body/实施方案.docx';before=path.read_bytes();ledger=(root/'body/.build-state.json').read_bytes()
 code,result=request(env,'docx_section_read',{'path':'body/实施方案.docx',**bad})
 assert code!=0 and 'error' in result
 assert path.read_bytes()==before and (root/'body/.build-state.json').read_bytes()==ledger
